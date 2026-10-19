import sys
from pptx import Presentation
p=Presentation(sys.argv[1])
for i,s in enumerate(p.slides,1):
    print("=== Slide",i, s.slide_layout.name)
    for sh in s.shapes:
        print(" --",sh.shape_type,sh.name, sh.left,sh.top,sh.width,sh.height)
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                print("   P:", [r.text for r in para.runs])
        if sh.has_table:
            for r in sh.table.rows:
                print("   R:", [c.text for c in r.cells])
        if sh.has_chart:
            c=sh.chart
            print("   chart type", c.chart_type, "cats", list(c.plots[0].categories))
            for pl in c.plots:
                for se in pl.series: print("   series", se.name, list(se.values))
    if s.has_notes_slide: print(" NOTES:", s.notes_slide.notes_text_frame.text)
