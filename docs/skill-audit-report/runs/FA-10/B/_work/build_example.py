"""Build the example deck using the helper code embedded verbatim in SKILL.md."""
import re, sys
from pathlib import Path

R = Path("/home/claude/results/2026-10-13/runs/FA-10/B")
skill = R / "northwind-ppt-template"
md = (skill / "SKILL.md").read_text()
helpers = md.split("## Helpers (copy into your build script)")[1]
code = re.search(r"```python\n(.*?)```", helpers, re.S).group(1)
g = {}
exec(code, g)
ph, set_bullets, add_slide_number, add_table, style_chart, remove_placeholder = (
    g[k] for k in ["ph", "set_bullets", "add_slide_number", "add_table", "style_chart", "remove_placeholder"])

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

prs = Presentation(str(skill / "assets/template.pptx"))
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

# 1 Cover
s = prs.slides.add_slide(prs.slide_layouts[1])
ph(s, 102).text = "Example Co.: Initiating at Overweight"
ph(s, 103).text = "Example Co. (EXCO)  |  Initiation of Coverage  |  2026-10-15  |  Illustrative data"
add_slide_number(s)

# 2 Section
s = prs.slides.add_slide(prs.slide_layouts[2])
ph(s, 102).text = "01"
ph(s, 103).text = "Investment Thesis"
add_slide_number(s)

# 3 Summary
s = prs.slides.add_slide(prs.slide_layouts[3])
ph(s, 102).text = "Margin expansion is under-appreciated"
set_bullets(ph(s, 103), [
    ("Key findings", 0, "heading"),
    ("Revenue grew 12.4% y/y to USD 1,250 million in FY2025E", 0),
    ("Subscription mix reached 61.0% of revenue, up 4.5 pts", 1),
    ("EBITDA margin expands from 18.2% to 22.6% by FY2027E", 0),
    ("Recommendation", 0, "heading"),
    ("Overweight, price target USD 85.00 (DCF, 9.0% WACC, 3.0% terminal growth)", 0),
    ("Implies 24.3% upside to the illustrative price of USD 68.40", 1),
])
add_slide_number(s)

# 4 Two content
s = prs.slides.add_slide(prs.slide_layouts[4])
ph(s, 102).text = "Bull vs bear: risk/reward skewed to the upside"
set_bullets(ph(s, 103), [
    ("Bull case: USD 110.00", 0, "heading"),
    ("Enterprise share gains lift growth to 15.0%", 0),
    ("Operating leverage drives 25.0% EBITDA margin", 0),
    ("Buybacks of USD 300 million p.a.", 0),
])
set_bullets(ph(s, 104), [
    ("Bear case: USD 55.00", 0, "heading"),
    ("Pricing pressure caps growth at 6.0%", 0),
    ("Margins stall at 18.0%", 0),
    ("Multiple de-rates to 9.0x EV/EBITDA", 0),
])
add_slide_number(s)

# 5 Chart
s = prs.slides.add_slide(prs.slide_layouts[5])
ph(s, 102).text = "Revenue and EBITDA (USD millions)"
cd = CategoryChartData()
cd.categories = ["FY2023", "FY2024", "FY2025E", "FY2026E", "FY2027E"]
cd.add_series("Revenue", (980, 1112, 1250, 1390, 1540))
cd.add_series("EBITDA", (170, 202, 245, 293, 348))
gf = ph(s, 103).insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
style_chart(gf.chart)
ph(s, 104).text = "Source: Illustrative data, Northwind Research estimates"
add_slide_number(s)

# 6 Section
s = prs.slides.add_slide(prs.slide_layouts[2])
ph(s, 102).text = "02"
ph(s, 103).text = "Sector Comparables"
add_slide_number(s)

# 7 Comps table
s = prs.slides.add_slide(prs.slide_layouts[3])
ph(s, 102).text = "Sector comps: Example Co. trades at a discount to peers"
remove_placeholder(s, 103)
add_table(s, [
    ["Company", "Ticker", "Mkt cap (USD m)", "EV/EBITDA FY26E", "P/E FY26E", "Rev growth FY26E"],
    ["Example Co.", "EXCO", "8,450", "10.4x", "17.8x", "11.2%"],
    ["Peer Alpha", "ALPH", "15,200", "13.1x", "22.4x", "9.8%"],
    ["Peer Beta", "BETA", "6,730", "12.0x", "20.1x", "8.5%"],
    ["Peer Gamma", "GAMA", "21,900", "14.6x", "24.9x", "12.6%"],
    ["Peer median", "", "15,200", "13.1x", "22.4x", "9.8%"],
], col_widths=[3.13, 1.4, 2.0, 2.0, 1.8, 2.0])
g["add_source"](s, "Source: Illustrative data, Northwind Research estimates")
add_slide_number(s)

# 8 Disclaimer
s = prs.slides.add_slide(prs.slide_layouts[6])
ph(s, 102).text = "Important Disclosures"
set_bullets(ph(s, 103), [
    ("Analyst certification", 0, "heading"),
    ("The research analyst(s) named in this report certify that the views expressed accurately reflect their personal views about the subject securities and issuers.", 0, "plain"),
    ("Illustrative content", 0, "heading"),
    ("All company names, figures and estimates in this example deck are fictional and for template validation only. Not investment advice.", 0, "plain"),
], size=11)
add_slide_number(s)

out = skill / "assets" / "example_presentation.pptx"
out2 = R / "northwind_example_presentation.pptx"
prs.save(str(out2))

print("saved", len(prs.slides), "slides")
