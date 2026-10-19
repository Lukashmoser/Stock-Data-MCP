import re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

OUT = "/home/claude/results/2026-10-13/runs/FA-06/A/Zoom_ZM_LBO_Model.xlsx"
wb = Workbook()
inp = wb.active; inp.title = "Inputs"
lbo = wb.create_sheet("LBO")

DARK = PatternFill("solid", fgColor="1F4E79")
LIGHT = PatternFill("solid", fgColor="D9E1F2")
GREY = PatternFill("solid", fgColor="F2F2F2")
KEY = PatternFill("solid", fgColor="BDD7EE")
thin = Side(style="thin", color="808080")

SAME = re.compile(r"^=\$?[A-Z]{1,3}\$?\d+$")
CROSS = re.compile(r"^='?[A-Za-z &]+'?!\$?[A-Z]{1,3}\$?\d+$")

def put(ws, ref, v, fmt=None, bold=False, fill=None):
    c = ws[ref]; c.value = v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        color = "0000FF"; c.fill = fill or GREY
    elif isinstance(v, (datetime.date,)):
        color = "0000FF"; c.fill = fill or GREY
    elif isinstance(v, str) and v.startswith("="):
        color = "008000" if CROSS.match(v) else ("800080" if SAME.match(v) else "000000")
        if fill: c.fill = fill
    else:
        color = "000000"
        if fill: c.fill = fill
    c.font = Font(name="Arial", size=10, color=color, bold=bold)
    if fmt: c.number_format = fmt
    if not (isinstance(v, str) and not v.startswith("=")):
        c.alignment = Alignment(horizontal="right")
    return c

def label(ws, ref, text, bold=False, indent=0, italic=False):
    c = ws[ref]; c.value = text
    c.font = Font(name="Arial", size=10, bold=bold, italic=italic)
    c.alignment = Alignment(indent=indent)

def header(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text)
    for col in range(1, ncols + 1):
        cc = ws.cell(row=row, column=col); cc.fill = DARK
        cc.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

M = '#,##0.0_);(#,##0.0);"-"_)'
D = '$#,##0.00'
P = '0.0%'
X = '0.0"x"'
X2 = '0.00"x"'
SH = '#,##0.000'

# ---------------- INPUTS ----------------
inp.column_dimensions["A"].width = 52
inp.column_dimensions["B"].width = 14
inp.column_dimensions["C"].width = 12
inp.column_dimensions["D"].width = 70
label(inp, "A1", "Zoom Communications (ZM) - Take-Private LBO: Inputs", bold=True)
label(inp, "A2", "USD millions except per-share. Blue = hard-coded input; black = formula; purple = same-tab link; green = cross-tab link.", italic=True)
for col, t in zip("ABCD", ["Item", "Value", "Units", "Source / note"]):
    c = inp[f"{col}3"]; c.value = t; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)

I = {}  # name -> cell ref
r = 4
def sec(t):
    global r
    r += 1; header(inp, r, t, 4); r += 1
def item(name, lab, v, fmt, unit, src):
    global r
    label(inp, f"A{r}", lab, indent=1)
    put(inp, f"B{r}", v, fmt)
    inp[f"C{r}"].value = unit; inp[f"D{r}"].value = src
    inp[f"D{r}"].font = Font(name="Arial", size=9, italic=True)
    I[name] = f"B{r}"; r += 1

sec("TRANSACTION")
item("price", "Current share price (close 2026-10-15)", 98.51, D, "$/sh", "stockanalysis.com / finance.yahoo.com, close 2026-10-15")
item("prem", "Offer premium to current price", 0.30, P, "%", "User assumption")
item("offer", "Offer price per share", f"={I['price']}*(1+{I['prem']})", D, "$/sh", "Formula")
item("basic", "Basic shares outstanding (Class A + B)", 291.783711, SH, "m sh", "10-Q Q2 FY27 cover: 263.233m A + 28.551m B as of 2026-08-14")
item("dil", "Dilutive securities (RSUs/options, est.)", 6.867104, SH, "m sh", "Est.: Q2 FY27 diluted WASO 299.746m less basic WASO 292.879m (8-K 2026-08-25)")
item("dshares", "Fully diluted shares", f"={I['basic']}+{I['dil']}", SH, "m sh", "Formula")
item("cash", "Cash & cash equivalents (2026-07-31)", 931.992, M, "$m", "8-K Ex.99.1 filed 2026-08-25")
item("mktsec", "Marketable securities (2026-07-31)", 6317.729, M, "$m", "8-K Ex.99.1 filed 2026-08-25")
item("exdebt", "Existing debt", 0.0, M, "$m", "No borrowings on balance sheet (8-K 2026-08-25)")
item("mincash", "Minimum operating cash left in business", 500.0, M, "$m", "Assumption (~10% of LTM revenue)")
item("txfee", "M&A / transaction fees (% of entry EV)", 0.015, P, "%", "Assumption")
item("finfee", "Financing fees (% of debt raised)", 0.0225, P, "%", "Assumption")
item("close", "Assumed closing date", datetime.date(2026, 10, 31), "yyyy-mm-dd", "date", "Assumption")
item("hold", "Holding period", 5, "0", "years", "User assumption")

sec("LTM FINANCIALS (LTM to 2026-07-31 = FY26 + 6M FY27 - 6M FY26)")
rows_ltm = [
    ("rev", "Revenue", 4868.769, 2516.223, 2391.942),
    ("oi", "GAAP income from operations", 1123.633, 624.789, 563.327),
    ("da", "Depreciation & amortization", 132.831, 65.112, 67.479),
    ("sbc", "Stock-based compensation expense", 760.776, 358.737, 390.268),
    ("capexh", "Capex (purchases of PP&E)", 64.961, 43.466, 33.876),
]
label(inp, f"A{r}", "Columns: FY26 (8-K 2026-02-25) | 6M FY27 | 6M FY26 (8-K 2026-08-25) | LTM", italic=True)
for col, t in zip("EFGH", ["FY2026", "6M FY27", "6M FY26", "LTM"]):
    c = inp[f"{col}{r}"]; c.value = t; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)
r += 1
for nm, lab, a, b, c_ in rows_ltm:
    label(inp, f"A{r}", lab, indent=1)
    put(inp, f"E{r}", a, M); put(inp, f"F{r}", b, M); put(inp, f"G{r}", c_, M)
    put(inp, f"H{r}", f"=E{r}+F{r}-G{r}", M)
    put(inp, f"B{r}", f"=H{r}", M)
    inp[f"C{r}"].value = "$m"; inp[f"D{r}"].value = "LTM formula in col H"
    I[nm] = f"B{r}"; r += 1
for col in "EFGH": inp.column_dimensions[col].width = 11
item("ebitda_gaap", "LTM EBITDA (GAAP op. income + D&A)", f"={I['oi']}+{I['da']}", M, "$m", "Formula; SBC treated as expense")
item("ebitda", "LTM Adj. EBITDA (op. income + D&A + SBC)", f"={I['ebitda_gaap']}+{I['sbc']}", M, "$m", "Formula; SBC added back (sponsor/lender convention). Basis for valuation & leverage")
item("margin0", "LTM Adj. EBITDA margin", f"={I['ebitda']}/{I['rev']}", P, "%", "Formula")

sec("OPERATING ASSUMPTIONS (Year 1-5 = columns C..G of the LBO tab)")
item("g1", "Revenue growth Y1", 0.040, P, "%", "Assumption; FY27 guide $5.085-5.095bn implies ~4.5%; Q2 grew 4.9%")
item("g2", "Revenue growth Y2", 0.040, P, "%", "Assumption")
item("g3", "Revenue growth Y3", 0.035, P, "%", "Assumption")
item("g4", "Revenue growth Y4", 0.035, P, "%", "Assumption")
item("g5", "Revenue growth Y5", 0.030, P, "%", "Assumption")
item("m1", "Adj. EBITDA margin Y1", 0.415, P, "%", "Assumption; modest sponsor cost-out from 41.0% LTM")
item("m2", "Adj. EBITDA margin Y2", 0.420, P, "%", "Assumption")
item("m3", "Adj. EBITDA margin Y3", 0.425, P, "%", "Assumption")
item("m4", "Adj. EBITDA margin Y4", 0.430, P, "%", "Assumption")
item("m5", "Adj. EBITDA margin Y5", 0.430, P, "%", "Assumption")
item("sbccash", "Cash cost to replace SBC (% of revenue)", 0.070, P, "%", "Assumption: ~half of LTM SBC (14.6% of rev) replaced with cash comp post take-private")
item("dapct", "D&A (% of revenue)", 0.026, P, "%", "LTM ~2.6%")
item("capexpct", "Capex (% of revenue)", 0.015, P, "%", "LTM ~1.5%")
item("nwcpct", "Net working capital (% of revenue)", -0.20, P, "%", "Approx.: AR $533m less deferred revenue $1,562m = -20.6% of LTM revenue")
item("tax", "Cash tax rate", 0.23, P, "%", "Assumption; 6M FY27 ETR ~22.1%")

sec("FINANCING (market as of 2026-10-13/14)")
item("sofr", "SOFR (2026-10-14)", 0.0388, '0.00%', "%", "FRED series SOFR, obs. 2026-10-14; held flat (no forward curve)")
item("tlspread", "Term Loan B spread over SOFR", 0.0325, '0.00%', "%", "Assumption: typical B+/BB- sponsor software TLB")
item("tlrate", "Term Loan B all-in rate", f"={I['sofr']}+{I['tlspread']}", '0.00%', "%", "Formula")
item("tllev", "Term Loan B size (x LTM Adj. EBITDA)", 4.0, X, "x", "Assumption")
item("tlamort", "TLB mandatory amortization (% of original p.a.)", 0.01, P, "%", "Market standard")
item("sweep", "Excess cash flow swept to TLB", 1.0, P, "%", "Assumption: 100% voluntary prepayment")
item("ust", "7-yr US Treasury yield (2026-10-13)", 0.0511, '0.00%', "%", "FRED series DGS7, obs. 2026-10-13")
item("hyoas", "ICE BofA US High Yield OAS (2026-10-13)", 0.0316, '0.00%', "%", "FRED series BAMLH0A0HYM2, obs. 2026-10-13")
item("nic", "New-issue / LBO concession", 0.0050, '0.00%', "%", "Assumption")
item("ntrate", "Senior notes coupon (fixed)", f"={I['ust']}+{I['hyoas']}+{I['nic']}", '0.00%', "%", "Formula: 7yr UST + HY OAS + concession")
item("ntlev", "Senior notes size (x LTM Adj. EBITDA)", 2.0, X, "x", "Assumption; non-call bullet, outstanding at exit")

# ---------------- LBO sheet ----------------
ws = lbo
ws.column_dimensions["A"].width = 44
for col in range(2, 9): ws.column_dimensions[L(col)].width = 13
label(ws, "A1", "Zoom Communications (ZM) - Take-Private LBO", bold=True)
label(ws, "A2", "USD millions except per-share. Entry assumed 2026-10-31; Year 1 = 12 months to 2027-10-31 ... Year 5 exit 2031-10-31.", italic=True)

def g(n): return f"Inputs!{I[n]}"
# helpful: bring key inputs into LBO via green links in a small block (rows 4-..)
R = {}
row = 4
header(ws, row, "SOURCES & USES", 8); row += 1
def line(name, lab, formula, fmt=M, bold=False, fill=None, col="B", indent=1):
    global row
    label(ws, f"A{row}", lab, bold=bold, indent=indent)
    put(ws, f"{col}{row}", formula, fmt, bold=bold, fill=fill)
    R[name] = f"{col}{row}"; row += 1

label(ws, f"A{row}", "Purchase price", bold=True); row += 1
line("price", "Current share price", f"={g('price')}", D)
line("prem", "Premium", f"={g('prem')}", P)
line("offer", "Offer price per share", f"={g('offer')}", D)
line("dsh", "Fully diluted shares (m)", f"={g('dshares')}", SH)
line("eqp", "Equity purchase price", f"={R['offer']}*{R['dsh']}", M, bold=True)
line("exdebt", "Plus: existing debt", f"={g('exdebt')}")
line("excash", "Less: cash + marketable securities", f"=-({g('cash')}+{g('mktsec')})")
line("ev", "Entry enterprise value", f"={R['eqp']}+{R['exdebt']}+{R['excash']}", M, bold=True)
line("ebitda", "LTM Adj. EBITDA (incl. SBC add-back)", f"={g('ebitda')}")
line("evx", "Entry EV / LTM Adj. EBITDA", f"={R['ev']}/{R['ebitda']}", X, bold=True)
line("ebitdag", "LTM EBITDA (SBC as expense)", f"={g('ebitda_gaap')}")
line("evxg", "Entry EV / LTM EBITDA (SBC expensed)", f"={R['ev']}/{R['ebitdag']}", X)
line("prev", "Implied premium to current market EV", f"={R['ev']}/({R['price']}*{R['dsh']}+{R['exdebt']}+{R['excash']})-1", P)
row += 1
srow = row
for col, t in zip("ABCDEFG", ["Sources", "$m", "% total", "x EBITDA", "Rate", "", ""]):
    c = ws[f"{col}{row}"]; c.value = t if t else None; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)
row += 1
s_start = row
line("tl", "Term Loan B", f"={g('tllev')}*{R['ebitda']}")
line("nt", "Senior Notes", f"={g('ntlev')}*{R['ebitda']}")
line("cashused", "Existing balance sheet cash used", f"={g('cash')}+{g('mktsec')}-{g('mincash')}")
line("spon", "Sponsor equity (plug)", "=0", M, bold=True, fill=KEY)  # placeholder, set later
s_end = row - 1
line("stot", "Total sources", f"=SUM(B{s_start}:B{s_end})", M, bold=True, indent=0)
for rr in range(s_start, row):
    put(ws, f"C{rr}", f"=B{rr}/$B${row-1}", P)
put(ws, f"D{s_start}", f"=B{s_start}/{R['ebitda']}", X); put(ws, f"D{s_start+1}", f"=B{s_start+1}/{R['ebitda']}", X)
put(ws, f"E{s_start}", f"={g('tlrate')}", '0.00%'); put(ws, f"E{s_start+1}", f"={g('ntrate')}", '0.00%')
row += 1
for col, t in zip("ABC", ["Uses", "$m", "% total"]):
    c = ws[f"{col}{row}"]; c.value = t; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)
row += 1
u_start = row
line("u_eq", "Purchase of equity", f"={R['eqp']}")
line("u_ref", "Refinance existing debt", f"={R['exdebt']}")
line("u_tx", "Transaction fees", f"={g('txfee')}*{R['ev']}")
line("u_fin", "Financing fees", f"={g('finfee')}*({R['tl']}+{R['nt']})")
u_end = row - 1
line("utot", "Total uses", f"=SUM(B{u_start}:B{u_end})", M, bold=True, indent=0)
for rr in range(u_start, row):
    put(ws, f"C{rr}", f"=B{rr}/$B${row-1}", P)
# sponsor plug
put(ws, R["spon"], f"={R['utot']}-SUM(B{s_start}:B{s_end-1})", M, bold=True, fill=KEY)
line("chk", "Check: sources - uses (should be 0)", f"={R['stot']}-{R['utot']}", M)
line("totlev", "Total debt / LTM Adj. EBITDA", f"=({R['tl']}+{R['nt']})/{R['ebitda']}", X)
line("eqpct", "Sponsor equity % of total capitalization", f"={R['spon']}/({R['spon']}+{R['tl']}+{R['nt']})", P)

# ---------------- OPERATING MODEL ----------------
row += 1
header(ws, row, "OPERATING MODEL", 8); row += 1
cols = ["B", "C", "D", "E", "F", "G"]  # B = LTM / closing, C..G = Y1..Y5
for col, t in zip(["A"] + cols, ["Period", "LTM 2026-07", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]):
    c = ws[f"{col}{row}"]; c.value = t; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)
    c.alignment = Alignment(horizontal="right" if col != "A" else "left")
row += 1
for col, t in zip(["A"] + cols, ["Period end", "2026-10-31", "2027-10-31", "2028-10-31", "2029-10-31", "2030-10-31", "2031-10-31"]):
    c = ws[f"{col}{row}"]; c.value = t if col == "A" else datetime.date(int(t[:4]), 10, 31)
    if col != "A":
        put(ws, f"{col}{row}", f"=EDATE({g('close')},{12*cols.index(col)})" if col != "B" else f"={g('close')}", "yyyy-mm-dd")
    else:
        label(ws, f"A{row}", "Period end (B = closing)")
row += 1
OP = {}
def oprow(name, lab, fB, fY, fmt=M, bold=False, indent=1):
    """fY: function(col, prevcol, yearidx)->formula"""
    global row
    label(ws, f"A{row}", lab, bold=bold, indent=indent)
    if fB is not None: put(ws, f"B{row}", fB, fmt, bold=bold)
    for i, col in enumerate(cols[1:], start=1):
        put(ws, f"{col}{row}", fY(col, cols[i-1], i), fmt, bold=bold)
    OP[name] = row; row += 1

oprow("g", "Revenue growth", None, lambda c, p, i: f"={g('g%d' % i)}", P)
oprow("rev", "Revenue", f"={g('rev')}", lambda c, p, i: f"={p}{row-1+1}*(1+{c}{row-1})", M, bold=True)
# fix revenue formula properly (row refs)
rr = OP["rev"]
for i, col in enumerate(cols[1:], start=1):
    put(ws, f"{col}{rr}", f"={cols[i-1]}{rr}*(1+{col}{OP['g']})", M, bold=True)
oprow("m", "Adj. EBITDA margin", f"={g('margin0')}", lambda c, p, i: f"={g('m%d' % i)}", P)
oprow("ebitda", "Adj. EBITDA (pre-SBC)", f"={g('ebitda')}", lambda c, p, i: f"={c}{OP['rev']}*{c}{OP['m']}", M, bold=True)
oprow("sbc", "Less: cash replacement of SBC", None, lambda c, p, i: f"=-{c}{OP['rev']}*{g('sbccash')}")
oprow("cebitda", "Cash EBITDA", None, lambda c, p, i: f"={c}{OP['ebitda']}+{c}{OP['sbc']}", M, bold=True)
oprow("da", "Less: D&A", None, lambda c, p, i: f"=-{c}{OP['rev']}*{g('dapct')}")
oprow("ebit", "EBIT", None, lambda c, p, i: f"={c}{OP['cebitda']}+{c}{OP['da']}")
oprow("int", "Less: interest expense", None, lambda c, p, i: "=0")  # filled after debt schedule
oprow("ebt", "Pre-tax income", None, lambda c, p, i: f"={c}{OP['ebit']}+{c}{OP['int']}")
oprow("taxes", "Less: cash taxes", None, lambda c, p, i: f"=-MAX(0,{c}{OP['ebt']})*{g('tax')}")
oprow("ni", "Net income", None, lambda c, p, i: f"={c}{OP['ebt']}+{c}{OP['taxes']}", M, bold=True)
row += 1
label(ws, f"A{row}", "Free cash flow", bold=True); row += 1
oprow("nwc", "Net working capital balance", f"={g('nwcpct')}*B{OP['rev']}", lambda c, p, i: f"={g('nwcpct')}*{c}{OP['rev']}")
oprow("dap", "Plus: D&A", None, lambda c, p, i: f"=-{c}{OP['da']}")
oprow("capex", "Less: capex", None, lambda c, p, i: f"=-{c}{OP['rev']}*{g('capexpct')}")
oprow("dnwc", "Less: increase in NWC", None, lambda c, p, i: f"=-({c}{OP['nwc']}-{p}{OP['nwc']})")
oprow("fcf", "Free cash flow available for debt repayment", None, lambda c, p, i: f"={c}{OP['ni']}+{c}{OP['dap']}+{c}{OP['capex']}+{c}{OP['dnwc']}", M, bold=True)

# ---------------- DEBT SCHEDULE ----------------
row += 1
header(ws, row, "DEBT SCHEDULE", 8); row += 1
for col, t in zip(["A"] + cols, ["", "Closing", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]):
    c = ws[f"{col}{row}"]; c.value = t or None; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)
row += 1
DS = {}
def dsrow(name, lab, fB, fY, fmt=M, bold=False, indent=1):
    global row
    label(ws, f"A{row}", lab, bold=bold, indent=indent)
    if fB is not None: put(ws, f"B{row}", fB, fmt, bold=bold)
    for i, col in enumerate(cols[1:], start=1):
        put(ws, f"{col}{row}", fY(col, cols[i-1], i), fmt, bold=bold)
    DS[name] = row; row += 1

label(ws, f"A{row}", "Cash", bold=True); row += 1
dsrow("cbeg", "Beginning cash", None, lambda c, p, i: f"={p}{row+4}")  # placeholder fixed below
dsrow("cfcf", "Plus: free cash flow", None, lambda c, p, i: f"={c}{OP['fcf']}")
dsrow("cmand", "Less: TLB mandatory amortization", None, lambda c, p, i: "=0")
dsrow("copt", "Less: TLB voluntary prepayment (sweep)", None, lambda c, p, i: "=0")
dsrow("cend", "Ending cash", f"={g('mincash')}", lambda c, p, i: f"=SUM({c}{row-4}:{c}{row-1})", M, bold=True)
for i, col in enumerate(cols[1:], start=1):
    put(ws, f"{col}{DS['cbeg']}", f"={cols[i-1]}{DS['cend']}")
dsrow("avail", "Cash available for sweep (above minimum)", None,
      lambda c, p, i: f"={c}{DS['cbeg']}+{c}{DS['cfcf']}+{c}{DS['cmand']}-{g('mincash')}")
row += 1
label(ws, f"A{row}", "Term Loan B", bold=True); row += 1
dsrow("tbeg", "Beginning balance", None, lambda c, p, i: "=0")
dsrow("tmand", "Mandatory amortization", None, lambda c, p, i: f"=-MIN({c}{row-1},{g('tlamort')}*$B${row+2})")
dsrow("topt", "Voluntary prepayment (cash sweep)", None,
      lambda c, p, i: f"=-MIN({c}{row-2}+{c}{row-1},MAX(0,{c}{DS['avail']})*{g('sweep')})")
dsrow("tend", "Ending balance", f"={R['tl']}", lambda c, p, i: f"=SUM({c}{row-3}:{c}{row-1})", M, bold=True)
for i, col in enumerate(cols[1:], start=1):
    put(ws, f"{col}{DS['tbeg']}", f"={cols[i-1]}{DS['tend']}")
    put(ws, f"{col}{DS['tmand']}", f"=-MIN({col}{DS['tbeg']},{g('tlamort')}*$B${DS['tend']})")
    put(ws, f"{col}{DS['topt']}", f"=-MIN({col}{DS['tbeg']}+{col}{DS['tmand']},MAX(0,{col}{DS['avail']})*{g('sweep')})")
    put(ws, f"{col}{DS['cmand']}", f"={col}{DS['tmand']}")
    put(ws, f"{col}{DS['copt']}", f"={col}{DS['topt']}")
dsrow("tint", "Interest expense (on beginning balance)", None, lambda c, p, i: f"={c}{DS['tbeg']}*{g('tlrate')}")
row += 1
label(ws, f"A{row}", "Senior Notes", bold=True); row += 1
dsrow("nbeg", "Beginning balance", None, lambda c, p, i: "=0")
dsrow("nend", "Ending balance (bullet, non-call)", f"={R['nt']}", lambda c, p, i: f"={c}{row-1}", M, bold=True)
for i, col in enumerate(cols[1:], start=1):
    put(ws, f"{col}{DS['nbeg']}", f"={cols[i-1]}{DS['nend']}")
dsrow("nint", "Interest expense", None, lambda c, p, i: f"={c}{DS['nbeg']}*{g('ntrate')}")
row += 1
dsrow("totint", "Total interest expense", None, lambda c, p, i: f"={c}{DS['tint']}+{c}{DS['nint']}", M, bold=True, indent=0)
dsrow("totdebt", "Total debt", f"=B{DS['tend']}+B{DS['nend']}", lambda c, p, i: f"={c}{DS['tend']}+{c}{DS['nend']}", M, bold=True, indent=0)
dsrow("netdebt", "Net debt", f"=B{DS['totdebt']}-B{DS['cend']}", lambda c, p, i: f"={c}{DS['totdebt']}-{c}{DS['cend']}", M, indent=0)
dsrow("lev", "Total debt / Adj. EBITDA", f"=B{DS['totdebt']}/B{OP['ebitda']}", lambda c, p, i: f"={c}{DS['totdebt']}/{c}{OP['ebitda']}", X, indent=0)
dsrow("cov", "Cash EBITDA / interest", None, lambda c, p, i: f"={c}{OP['cebitda']}/{c}{DS['totint']}", X, indent=0)
# link interest into op model
for col in cols[1:]:
    put(ws, f"{col}{OP['int']}", f"=-{col}{DS['totint']}")

# ---------------- RETURNS ----------------
row += 1
header(ws, row, "RETURNS ANALYSIS", 8); row += 1
line("xmult", "Exit EV / Adj. EBITDA (= entry multiple)", f"={R['evx']}", X)
line("xebitda", "Year 5 Adj. EBITDA", f"=G{OP['ebitda']}")
line("xev", "Exit enterprise value", f"={R['xmult']}*{R['xebitda']}", M, bold=True)
line("xdebt", "Less: total debt at exit", f"=-G{DS['totdebt']}")
line("xcash", "Plus: cash at exit", f"=G{DS['cend']}")
line("xeq", "Exit equity value", f"={R['xev']}+{R['xdebt']}+{R['xcash']}", M, bold=True)
row += 1
for col, t in zip(["A"] + cols, ["Sponsor cash flows", "Closing", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]):
    c = ws[f"{col}{row}"]; c.value = t; c.fill = LIGHT; c.font = Font(name="Arial", size=10, bold=True)
row += 1
label(ws, f"A{row}", "Equity (invested) / proceeds", indent=1)
put(ws, f"B{row}", f"=-{R['spon']}", M)
for col in cols[1:5]: put(ws, f"{col}{row}", 0, M)
put(ws, f"G{row}", f"={R['xeq']}", M)
cfrow = row; row += 1
line("irr", "IRR", f"=IRR(B{cfrow}:G{cfrow})", P, bold=True, fill=KEY)
line("moic", "MOIC", f"={R['xeq']}/{R['spon']}", X2, bold=True, fill=KEY)
line("irrchk", "Check: (MOIC)^(1/hold) - 1", f"={R['moic']}^(1/{g('hold')})-1", P)

# value creation bridge
row += 1
label(ws, f"A{row}", "Value creation bridge", bold=True); row += 1
line("vc_ebitda", "EBITDA growth (at entry multiple)", f"=({R['xebitda']}-{R['ebitda']})*{R['evx']}")
line("vc_mult", "Multiple expansion", f"=({R['xmult']}-{R['evx']})*{R['xebitda']}")
line("vc_delev", "Net debt paydown", f"=(B{DS['netdebt']}-G{DS['netdebt']})")
line("vc_leak", "Less: fees & cash at entry (leakage)", f"={R['xeq']}-{R['spon']}-B{row-3}-B{row-2}-B{row-1}")
line("vc_tot", "Total equity gain", f"=SUM(B{row-4}:B{row-1})", M, bold=True)

# ---------------- SENSITIVITIES ----------------
def sens(title, kind):
    global row
    row += 1
    header(ws, row, title, 8); row += 1
    label(ws, f"A{row}", "Rows: offer premium | Columns: exit EV/Adj. EBITDA", italic=True)
    mults = [-2, -1, 0, 1, 2]
    for j, d in enumerate(mults):
        put(ws, f"{L(2+j)}{row}", f"={R['evx']}+({d})" if d else f"={R['evx']}", X, bold=True, fill=LIGHT)
    hdr = row; row += 1
    for k, p in enumerate([0.20, 0.25, 0.30, 0.35, 0.40]):
        if p == 0.30:
            put(ws, f"A{row}", f"={g('prem')}", P, bold=True, fill=LIGHT)
        else:
            put(ws, f"A{row}", f"={g('prem')}+({round(p-0.30,2)})", P, bold=True, fill=LIGHT)
        for j in range(5):
            colh = f"{L(2+j)}${hdr}"
            eq = f"({g('price')}*(1+$A{row})*{R['dsh']})"
            evp = f"({eq}+{R['exdebt']}+{R['excash']})"
            spon = (f"({eq}+{R['exdebt']}+{g('txfee')}*{evp}+{R['u_fin']}"
                    f"-{R['tl']}-{R['nt']}-{R['cashused']})")
            exiteq = f"({colh}*{R['xebitda']}+{R['xdebt']}+{R['xcash']})"
            f = f"={exiteq}/{spon}" if kind == "moic" else f"=({exiteq}/{spon})^(1/{g('hold')})-1"
            center = (p == 0.30 and j == 2)
            put(ws, f"{L(2+j)}{row}", f, X2 if kind == "moic" else P, bold=center, fill=KEY if center else None)
        row += 1
    note = "Center cell = base case; equals model " + ("MOIC" if kind == "moic" else "IRR") + ". Debt sized on LTM EBITDA, so premium changes only sponsor equity (and fees)."
    label(ws, f"A{row}", note, italic=True); row += 1

sens("SENSITIVITY: IRR (offer premium vs. exit multiple)", "irr")
sens("SENSITIVITY: MOIC (offer premium vs. exit multiple)", "moic")

ws.freeze_panes = "B4"
wb.save(OUT)
print("saved", OUT)
print({k: v for k, v in R.items() if k in ("spon", "irr", "moic", "ev", "evx", "tl", "nt", "chk")})
