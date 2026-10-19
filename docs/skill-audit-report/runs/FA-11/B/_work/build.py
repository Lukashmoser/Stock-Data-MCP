"""Build MSFT 1-year performance & fundamentals deck (5 slides, native charts)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

RUN = "/home/claude/results/2026-10-13/runs/FA-11/B"
OUT = os.path.join(RUN, "out")
os.makedirs(OUT, exist_ok=True)

# ---- Data (stock-data get_price_history, 1y, 1wk; adj_close; fetched 2026-10-16T03:29Z) ----
dates = ["2025-10-13","2025-10-20","2025-10-27","2025-11-03","2025-11-10","2025-11-17","2025-11-24","2025-12-01","2025-12-08","2025-12-15","2025-12-22","2025-12-29","2026-01-05","2026-01-12","2026-01-19","2026-01-26","2026-02-02","2026-02-09","2026-02-16","2026-02-23","2026-03-02","2026-03-09","2026-03-16","2026-03-23","2026-03-30","2026-04-06","2026-04-13","2026-04-20","2026-04-27","2026-05-04","2026-05-11","2026-05-18","2026-05-25","2026-06-01","2026-06-08","2026-06-15","2026-06-22","2026-06-29","2026-07-06","2026-07-13","2026-07-20","2026-07-27","2026-08-03","2026-08-10","2026-08-17","2026-08-24","2026-08-31","2026-09-07","2026-09-14","2026-09-21","2026-09-28","2026-10-05","2026-10-12"]
msft = [509.39,519.34,513.58,492.77,506.02,469.14,488.91,480.11,475.51,482.86,484.64,469.96,476.26,456.96,463.01,427.58,398.61,398.79,395.63,391.15,407.31,393.95,380.33,355.33,371.95,369.37,421.08,422.91,412.77,413.44,420.22,417.78,449.39,415.89,390.01,378.69,372.27,389.76,384.38,393.08,380.98,463.85,499.05,494.47,483.24,513.53,499.7,495.63,493.78,516.17,517.53,535.07,539.44]
qqq = [601.11,614.21,626.13,606.89,606.01,587.31,616.35,622.55,610.75,614.16,621.77,611.04,624.52,619.15,620.61,619.76,607.58,599.88,606.74,605.23,597.71,591.7,580.08,561.38,583.73,609.76,647.46,662.46,672.71,709.71,707.41,716.01,736.73,703.55,719.8,739.04,705.78,711.86,724.75,694.61,683.52,687.27,722.28,730.31,712.7,715.68,718.21,714.14,720.7,744.5,749.58,751.27,728.0]
xlk = [141.78,146.04,149.57,143.34,143.34,135.9,142.38,145.85,142.96,143.86,146.0,143.78,145.62,145.1,144.57,143.36,140.62,139.06,140.37,138.26,136.8,136.31,134.8,129.61,135.67,142.28,153.99,159.84,161.49,175.11,175.84,179.97,190.57,179.88,184.36,190.99,180.9,180.38,185.56,175.39,175.67,175.15,187.75,189.79,183.1,185.47,187.06,187.45,189.38,196.27,199.81,198.78,189.19]
assert len(dates) == len(msft) == len(qqq) == len(xlk) == 53

def tr(s):
    return (s[-1] / s[0] - 1) * 100

def idx(s):
    return [round(v / s[0] * 100, 2) for v in s]

tr_m, tr_q, tr_x = tr(msft), tr(qqq), tr(xlk)

# ---- Fundamentals (stock-data get_financials, SEC XBRL) ----
qtrs = ["FY25 Q1","FY25 Q2","FY25 Q3","FY25 Q4","FY26 Q1","FY26 Q2","FY26 Q3","FY26 Q4"]
qend = ["2024-09-30","2024-12-31","2025-03-31","2025-06-30","2025-09-30","2025-12-31","2026-03-31","2026-06-30"]
rev = [65585,69632,70066,76441,77673,81273,82886,90007]
opi = [30552,31653,32000,34323,37961,38275,38398,40603]
opm = [o / r * 100 for o, r in zip(opi, rev)]

fys = ["FY2023","FY2024","FY2025","FY2026"]
ocf = [87582,118548,136162,182935]
capex = [28107,44477,64551,115948]
fcf = [o - c for o, c in zip(ocf, capex)]

# ---- Styling ----
NAVY = RGBColor(0x0B, 0x2A, 0x4A)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
ORANGE = RGBColor(0xE0, 0x7A, 0x1F)
GREY = RGBColor(0x7F, 0x8C, 0x99)
DARK = RGBColor(0x26, 0x32, 0x3D)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_text(slide, text, x, y, w, h, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def header(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    add_text(slide, title, 0.5, 0.25, 12.3, 0.8, size=22, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, subtitle, 0.5, 0.95, 12.3, 0.5, size=13, color=GREY)

def footer(slide, text):
    add_text(slide, text, 0.5, 6.85, 12.3, 0.5, size=9, color=GREY)

def style_series(series, color, width=2.25, smooth=False):
    fmt = series.format
    fmt.line.color.rgb = color
    fmt.line.width = Pt(width)
    series.smooth = smooth

# ---- Slide 1: Title ----
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
add_text(s, "Microsoft: 1-Year Stock Performance and Fundamentals", 0.8, 2.4, 11.7, 1.4,
         size=40, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
add_text(s, "MSFT vs. QQQ vs. XLK  |  Revenue & operating margin  |  Capex & free cash flow", 0.8, 3.9, 11.7, 0.6,
         size=18, color=RGBColor(0xC9, 0xD6, 0xE3))
add_text(s, "Market data as of 2026-10-15 close  |  Financials through FY2026 Q4 (quarter ended 2026-06-30)", 0.8, 4.6, 11.7, 0.5,
         size=14, color=RGBColor(0xC9, 0xD6, 0xE3))

# ---- Slide 2: Total return line chart ----
s = prs.slides.add_slide(BLANK)
header(s, f"MSFT +{tr_m:.1f}% over 1 year, lagging QQQ (+{tr_q:.1f}%) and XLK (+{tr_x:.1f}%)",
       "1-year total return, weekly, indexed to 100 at first weekly close (dividend-adjusted)")
cd = CategoryChartData()
cd.categories = dates
cd.add_series(f"MSFT ({tr_m:+.1f}%)", idx(msft))
cd.add_series(f"QQQ ({tr_q:+.1f}%)", idx(qqq))
cd.add_series(f"XLK ({tr_x:+.1f}%)", idx(xlk))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3), cd)
ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
for ser, col in zip(ch.plots[0].series, [BLUE, ORANGE, GREY]):
    style_series(ser, col, 2.75 if col == BLUE else 2.0)
ca = ch.category_axis
ca.tick_labels.font.size = Pt(10)
ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
# show every 4th weekly label
cax = ca._element
for tag, val in (("c:tickLblSkip", "4"), ("c:tickMarkSkip", "4")):
    el = etree.SubElement(cax, qn(tag)); el.set("val", val)
va = ch.value_axis
va.tick_labels.font.size = Pt(10)
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
va.minimum_scale = 60; va.maximum_scale = 150
va.tick_labels.number_format = '0'; va.tick_labels.number_format_is_linked = False
footer(s, "Source: Yahoo Finance via stock-data get_price_history (period 1y, interval 1wk, adj_close), fetched 2026-10-16. "
          "53 weekly points per series, week of 2025-10-13 to week of 2026-10-12 (last point = 2026-10-15 close). "
          "Total return = last / first adjusted close - 1.")

# ---- Slide 3: Revenue + operating margin combo ----
s = prs.slides.add_slide(BLANK)
header(s, f"FY26 Q4 revenue ${rev[-1]:,}M (+{(rev[-1]/rev[3]-1)*100:.1f}% YoY); op. margin {opm[-1]:.1f}%, off {max(opm):.1f}% peak",
       "Quarterly revenue (USD millions, bars) and operating margin (%, line, right axis), last 8 fiscal quarters")
cd = CategoryChartData()
cd.categories = qtrs
cd.add_series("Revenue (USD M)", rev)
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3), cd)
ch = gf.chart
bar_ser = ch.plots[0].series[0]
bar_ser.format.fill.solid(); bar_ser.format.fill.fore_color.rgb = BLUE
pl = ch.plots[0]; pl.gap_width = 60
pl.has_data_labels = True
pl.data_labels.number_format = '#,##0'; pl.data_labels.number_format_is_linked = False
pl.data_labels.font.size = Pt(10)
from pptx.enum.chart import XL_LABEL_POSITION
pl.data_labels.position = XL_LABEL_POSITION.INSIDE_END
pl.data_labels.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
va = ch.value_axis; va.minimum_scale = 0; va.maximum_scale = 140000; va.major_unit = 20000
va.tick_labels.number_format = '#,##0'; va.tick_labels.number_format_is_linked = False
va.tick_labels.font.size = Pt(10)
va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
ch.category_axis.tick_labels.font.size = Pt(11)

# Inject a line chart on a secondary axis into the same plotArea (native combo chart).
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
def E(tag, **attrs):
    el = etree.Element(qn(tag))
    for k, v in attrs.items():
        el.set(k, v)
    return el
def SE(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, v)
    return el

cs = ch._chartSpace
plotArea = cs.find(".//" + qn("c:plotArea"))
barChart = plotArea.find(qn("c:barChart"))
line = E("c:lineChart")
SE(line, "c:grouping", val="standard")
SE(line, "c:varyColors", val="0")
ser = SE(line, "c:ser")
SE(ser, "c:idx", val="1"); SE(ser, "c:order", val="1")
tx = SE(ser, "c:tx"); sr = SE(tx, "c:strRef"); SE(sr, "c:f").text = "Sheet1!$C$1"
sc = SE(sr, "c:strCache"); SE(sc, "c:ptCount", val="1"); pt = SE(sc, "c:pt", idx="0"); SE(pt, "c:v").text = "Operating margin (%)"
spPr = SE(ser, "c:spPr")
ln = SE(spPr, "a:ln", w=str(int(Pt(3)))); sf = SE(ln, "a:solidFill"); SE(sf, "a:srgbClr", val="E07A1F")
mk = SE(ser, "c:marker"); SE(mk, "c:symbol", val="circle"); SE(mk, "c:size", val="7")
mksp = SE(mk, "c:spPr"); f2 = SE(mksp, "a:solidFill"); SE(f2, "a:srgbClr", val="E07A1F")
dl = SE(ser, "c:dLbls")
SE(dl, "c:numFmt", formatCode='0.0"%"', sourceLinked="0")
dsp = SE(dl, "c:spPr"); SE(dsp, "a:noFill")
txPr = SE(dl, "c:txPr"); bp = SE(txPr, "a:bodyPr"); SE(txPr, "a:lstStyle"); p = SE(txPr, "a:p")
pPr = SE(p, "a:pPr"); dr = SE(pPr, "a:defRPr", sz="1000", b="1"); f3 = SE(dr, "a:solidFill"); SE(f3, "a:srgbClr", val="E07A1F")
SE(p, "a:endParaRPr", lang="en-US")
SE(dl, "c:dLblPos", val="t")
for t in ("showLegendKey", "showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize"):
    SE(dl, "c:" + t, val="1" if t == "showVal" else "0")
cat = SE(ser, "c:cat"); cr = SE(cat, "c:strRef"); SE(cr, "c:f").text = f"Sheet1!$A$2:$A${len(qtrs)+1}"
cc = SE(cr, "c:strCache"); SE(cc, "c:ptCount", val=str(len(qtrs)))
for i, q in enumerate(qtrs):
    pt = SE(cc, "c:pt", idx=str(i)); SE(pt, "c:v").text = q
val = SE(ser, "c:val"); nr = SE(val, "c:numRef"); SE(nr, "c:f").text = f"Sheet1!$C$2:$C${len(qtrs)+1}"
nc = SE(nr, "c:numCache"); SE(nc, "c:formatCode").text = "0.0"; SE(nc, "c:ptCount", val=str(len(qtrs)))
for i, m in enumerate(opm):
    pt = SE(nc, "c:pt", idx=str(i)); SE(pt, "c:v").text = f"{m:.4f}"
SE(ser, "c:smooth", val="0")
SE(line, "c:marker", val="1")
SE(line, "c:axId", val="50010"); SE(line, "c:axId", val="50020")
barChart.addnext(line)
# secondary axes: hidden cat axis + right-hand value axis
last_ax = plotArea.findall(qn("c:valAx"))[-1]
catAx2 = E("c:catAx")
SE(catAx2, "c:axId", val="50010"); sco = SE(catAx2, "c:scaling"); SE(sco, "c:orientation", val="minMax")
SE(catAx2, "c:delete", val="1"); SE(catAx2, "c:axPos", val="b"); SE(catAx2, "c:majorTickMark", val="none")
SE(catAx2, "c:minorTickMark", val="none"); SE(catAx2, "c:tickLblPos", val="nextTo")
SE(catAx2, "c:crossAx", val="50020"); SE(catAx2, "c:crosses", val="autoZero")
SE(catAx2, "c:auto", val="1"); SE(catAx2, "c:lblAlgn", val="ctr"); SE(catAx2, "c:lblOffset", val="100"); SE(catAx2, "c:noMultiLvlLbl", val="0")
valAx2 = E("c:valAx")
SE(valAx2, "c:axId", val="50020"); sco = SE(valAx2, "c:scaling"); SE(sco, "c:orientation", val="minMax")
SE(sco, "c:max", val="50"); SE(sco, "c:min", val="20")
SE(valAx2, "c:delete", val="0"); SE(valAx2, "c:axPos", val="r")
SE(valAx2, "c:numFmt", formatCode='0"%"', sourceLinked="0")
SE(valAx2, "c:majorTickMark", val="out"); SE(valAx2, "c:minorTickMark", val="none"); SE(valAx2, "c:tickLblPos", val="nextTo")
tp = SE(valAx2, "c:txPr"); SE(tp, "a:bodyPr"); SE(tp, "a:lstStyle"); pp = SE(tp, "a:p"); ppr = SE(pp, "a:pPr"); SE(ppr, "a:defRPr", sz="1000"); SE(pp, "a:endParaRPr", lang="en-US")
SE(valAx2, "c:crossAx", val="50010"); SE(valAx2, "c:crosses", val="max"); SE(valAx2, "c:crossBetween", val="between")
last_ax.addnext(catAx2); catAx2.addnext(valAx2)
# legend on
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
# write the margin column into the embedded workbook so "Edit Data" stays consistent
from openpyxl import load_workbook
import io
wb_part = gf.chart.part.chart_workbook.xlsx_part
wb = load_workbook(io.BytesIO(wb_part.blob))
ws = wb.worksheets[0]
ws["C1"] = "Operating margin (%)"
for i, m in enumerate(opm):
    ws.cell(row=i + 2, column=3, value=round(m, 4))
buf = io.BytesIO(); wb.save(buf); wb_part._blob = buf.getvalue()
footer(s, "Source: Microsoft 10-K/10-Q XBRL via stock-data get_financials (quarterly income). Fiscal year ends 30 June. "
          "Q4 values derived as FY less 9M (latest 10-K filed 2026-07-29). Operating margin = operating income / revenue.")

# ---- Slide 4: Capex & FCF ----
s = prs.slides.add_slide(BLANK)
header(s, f"Capex rose {capex[-1]/capex[0]:.1f}x since FY2023 to ${capex[-1]:,}M; FCF fell to ${fcf[-1]:,}M in FY2026",
       "Capital expenditures vs. free cash flow by fiscal year (USD millions); FCF = operating cash flow - capex (PP&E additions)")
cd = CategoryChartData()
cd.categories = fys
cd.add_series("Capex (USD M)", capex)
cd.add_series("Free cash flow (USD M)", fcf)
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(8.6), Inches(5.3), cd)
ch = gf.chart
pl = ch.plots[0]; pl.gap_width = 70; pl.overlap = -10
for ser_, col in zip(pl.series, [ORANGE, BLUE]):
    ser_.format.fill.solid(); ser_.format.fill.fore_color.rgb = col
pl.has_data_labels = True
pl.data_labels.number_format = '#,##0'; pl.data_labels.number_format_is_linked = False
pl.data_labels.font.size = Pt(11)
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
va = ch.value_axis; va.minimum_scale = 0
va.tick_labels.number_format = '#,##0'; va.tick_labels.number_format_is_linked = False
va.tick_labels.font.size = Pt(10)
va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
ch.category_axis.tick_labels.font.size = Pt(12)
# side table
rows = [["FY", "OCF", "Capex", "FCF", "Capex/OCF"]] + [
    [fy, f"{o:,}", f"{c:,}", f"{f:,}", f"{c/o*100:.1f}%"] for fy, o, c, f in zip(fys, ocf, capex, fcf)]
tbl = s.shapes.add_table(len(rows), 5, Inches(9.3), Inches(2.0), Inches(3.6), Inches(2.4)).table
for r, row in enumerate(rows):
    for c_, v in enumerate(row):
        cell = tbl.cell(r, c_); cell.text = v
        para = cell.text_frame.paragraphs[0]; para.alignment = PP_ALIGN.RIGHT if c_ else PP_ALIGN.LEFT
        for run in para.runs:
            run.font.size = Pt(10); run.font.bold = (r == 0)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else DARK
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if r == 0 else RGBColor(0xF2, 0xF4, 0xF7)
add_text(s, "USD millions", 9.3, 1.6, 3.6, 0.4, size=10, color=GREY)
footer(s, "Source: Microsoft 10-K XBRL via stock-data get_financials (annual cash flow); FY2026 10-K filed 2026-07-29. "
          "Capex = PaymentsToAcquirePropertyPlantAndEquipment; excludes finance-lease additions.")

# ---- Slide 5: Key takeaways ----
s = prs.slides.add_slide(BLANK)
header(s, "Takeaways: growth accelerated, but the stock lagged as capex squeezed FCF")
yoy_q4 = (rev[-1] / rev[3] - 1) * 100
fy26_rev = sum(rev[4:]); fy25_rev = sum(rev[:4])
fy26_opm = sum(opi[4:]) / fy26_rev * 100; fy25_opm = sum(opi[:4]) / fy25_rev * 100
bullets = [
    f"1. Stock: MSFT total return {tr_m:+.1f}% (2025-10-17 to 2026-10-15) vs. QQQ {tr_q:+.1f}% and XLK {tr_x:+.1f}%, "
    f"underperforming by {tr_q-tr_m:.1f} and {tr_x-tr_m:.1f} pts. Trough on 2026-03-27 ({min(idx(msft))-100:.1f}% from start); "
    f"+{(msft[-1]/min(msft)-1)*100:.1f}% since, led by the post-FY26 Q4 earnings rally (late July 2026).",
    f"2. Revenue: ${rev[0]:,}M (FY25 Q1) to ${rev[-1]:,}M (FY26 Q4); FY26 Q4 up {yoy_q4:.1f}% YoY. "
    f"FY2026 revenue ${fy26_rev:,}M vs. ${fy25_rev:,}M in FY2025 (+{(fy26_rev/fy25_rev-1)*100:.1f}%).",
    f"3. Margins: full-year operating margin {fy25_opm:.1f}% (FY25) to {fy26_opm:.1f}% (FY26), but the quarterly trend faded from "
    f"{max(opm):.1f}% (FY26 Q1) to {opm[-1]:.1f}% (FY26 Q4) as AI infrastructure costs ramped.",
    f"4. Capex: ${capex[0]:,}M (FY23) to ${capex[-1]:,}M (FY26), +{(capex[-1]/capex[-2]-1)*100:.1f}% in FY26 alone; "
    f"capex absorbed {capex[-1]/ocf[-1]*100:.1f}% of operating cash flow (FY23: {capex[0]/ocf[0]*100:.1f}%).",
    f"5. Free cash flow: ${fcf[1]:,}M peak (FY24) to ${fcf[-1]:,}M (FY26) even as OCF grew to ${ocf[-1]:,}M; "
    f"the AI build-out is funded internally at the cost of near-term FCF.",
]
add_text(s, bullets, 0.6, 1.2, 12.1, 5.4, size=17, color=DARK)
for p in s.shapes[-1].text_frame.paragraphs:
    p.space_after = Pt(14)
footer(s, "Sources: stock-data get_price_history / get_financials (Yahoo Finance; SEC XBRL). Market data as of 2026-10-15 close; financials through quarter ended 2026-06-30.")

path = os.path.join(OUT, "MSFT_1Y_Performance_Fundamentals.pptx")
prs.save(path)
print(path)
print(f"TR MSFT {tr_m:.2f} QQQ {tr_q:.2f} XLK {tr_x:.2f}")
print("OPM", [round(x, 1) for x in opm])
print("FCF", fcf)
print("min idx msft", min(idx(msft)), dates[idx(msft).index(min(idx(msft)))])
print("FY25/26 rev", fy25_rev, fy26_rev, round(fy25_opm, 1), round(fy26_opm, 1), round(yoy_q4, 1))
