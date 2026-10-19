import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import ColorScaleRule

ASOF="2026-10-14"
BLUE=Font(name="Arial",size=10,color="0000FF"); BLK=Font(name="Arial",size=10,color="000000")
GRN=Font(name="Arial",size=10,color="008000"); BOLD=Font(name="Arial",size=10,bold=True)
HDR_FILL=PatternFill("solid",fgColor="1F4E79"); HDR_FONT=Font(name="Arial",size=10,bold=True,color="FFFFFF")
SUB_FILL=PatternFill("solid",fgColor="D9E1F2"); IN_FILL=PatternFill("solid",fgColor="F2F2F2")
OUT_FILL=PatternFill("solid",fgColor="BDD7EE")
thin=Side(style="thin"); med=Side(style="medium"); thick=Side(style="thick")
PCT="0.0%"; MUL='0.0"x"'; NUM='#,##0;(#,##0);"-"'; USD='$#,##0.00'; BETA="0.00"

def box(ws,r1,c1,r2,c2,side=med):
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            cell=ws.cell(r,c); b=cell.border
            cell.border=Border(left=side if c==c1 else b.left,right=side if c==c2 else b.right,
                               top=side if r==r1 else b.top,bottom=side if r==r2 else b.bottom)
def hdr(ws,r,text,c1=1,c2=14):
    ws.cell(r,c1,text)
    for c in range(c1,c2+1):
        ws.cell(r,c).fill=HDR_FILL; ws.cell(r,c).font=HDR_FONT
def sub(ws,r,vals,c1=1):
    for i,v in enumerate(vals):
        c=ws.cell(r,c1+i,v); c.fill=SUB_FILL; c.font=BOLD; c.alignment=Alignment(horizontal="center")
def inp(ws,ref,val,src,fmt=None):
    c=ws[ref]; c.value=val; c.font=BLUE; c.fill=IN_FILL
    c.comment=Comment(src,"Analyst"); 
    if fmt: c.number_format=fmt
    return c
def f(ws,ref,formula,fmt=None,font=BLK):
    c=ws[ref]; c.value=formula; c.font=font
    if fmt: c.number_format=fmt
    return c

QSRC="Source: stock-data get_quote (Yahoo Finance), fetched 2026-10-15T02:06Z (US close 2026-10-14)"
# ---------------- Peer data ----------------
peers=[
 # name,ticker,ltm_end,rev,rev_prior,gp,ebit,da,ni,cash,debt,price,mcap(M),src_note, da_note
 ("Alphabet","GOOGL","2026-06-30",445867,371399,271517,147628,25086,244205,242474,100164,347.64,4251614.7,
  "stock-data get_financials quarterly (SEC XBRL) Q3-25..Q2-26","ESTIMATE: H1-26 depreciation 13,586 (10-Q filed 2026-07, goog-20260630.htm) + H2-25 est. 11,500 (H1-25 was 9,485)"),
 ("Amazon","AMZN","2026-06-30",775680,670038,393810,93712,75200,135281,122988,132224,254.08,2740586.6,
  "stock-data get_financials quarterly (SEC XBRL) Q3-25..Q2-26","stock-data get_financials quarterly cashflow, sum Q3-25..Q2-26"),
 ("Meta Platforms","META","2026-06-30",228247,178804,186587,86927,22729,68098,90260,83664,737.54,1878887.7,
  "stock-data get_financials quarterly (SEC XBRL) Q3-25..Q2-26","stock-data get_financials quarterly cashflow, sum Q3-25..Q2-26"),
 ("Apple","AAPL","2026-06-27",466823,408625,227123,154859,13100,128930,62399,82347,342.96,5005219.8,
  "stock-data get_financials quarterly (SEC XBRL) Q4-FY25..Q3-FY26","stock-data get_financials quarterly cashflow, sum Q4-FY25..Q3-FY26"),
 ("Oracle","ORCL","2026-08-31",71777,59018,None,23056,10881,18920,37077,125337,130.55,395762.5,
  "stock-data get_financials quarterly (SEC XBRL) Q2-FY26..Q1-FY27; debt from 10-Q for 2026-08-31 (orcl-20260831.htm)","FY26 D&A 9,294 (8-K ex99.1, FY26 results) + Q1-FY27 3,358 - Q1-FY26 1,771 (10-Q orcl-20260831.htm)"),
]
msft=("Microsoft (target)","MSFT","2026-06-30",331839,281724,225465,155237,39000,133749,76843,40294,533.69,3962939.4,
  "stock-data get_financials annual FY2026 (10-K filed 2026-07-29)","FY2026 D&A ~39,000: stockanalysis.com / TIKR (rounded, not verified to 10-K)")

def build_comps(ws):
    ws.sheet_view.showGridLines=False
    ws["A1"]="TECHNOLOGY MEGA-CAP - COMPARABLE COMPANY ANALYSIS"; ws["A1"].font=Font(name="Arial",size=12,bold=True)
    ws["A2"]="Microsoft (MSFT) vs. Alphabet (GOOGL) • Amazon (AMZN) • Meta (META) • Apple (AAPL) • Oracle (ORCL)"
    ws["A3"]=f"As of {ASOF} (prices) | LTM to latest reported quarter | USD millions except per-share and ratios"
    hdr(ws,5,"OPERATING METRICS (LTM)",1,14)
    sub(ws,6,["Company","Ticker","LTM end","Revenue","Revenue prior LTM","Rev growth","Gross profit","Gross margin","EBIT","EBIT margin","D&A","EBITDA","EBITDA margin","Net income"])
    rows=[msft]+peers
    for i,p in enumerate(rows):
        r=7+i if i==0 else 8+i  # MSFT row 7, blank row 8, peers 9-13
        ws.cell(r,1,p[0]).font=BOLD if i==0 else BLK; ws.cell(r,2,p[1]); ws.cell(r,3,p[2])
        for col,val,fmt in [(4,p[3],NUM),(5,p[4],NUM),(7,p[5],NUM),(9,p[6],NUM),(11,p[7],NUM),(14,p[8],NUM)]:
            if val is None:
                ws.cell(r,col,"n/a").font=BLK; ws.cell(r,col).comment=Comment("Not reported in XBRL standardized data (Oracle does not report cost of revenue/gross profit)","Analyst")
            else:
                inp(ws,f"{L(col)}{r}",val,("Source: "+p[13]) if col!=11 else ("Source: "+p[14]),fmt)
        f(ws,f"F{r}",f"=D{r}/E{r}-1",PCT)
        f(ws,f"H{r}",f'=IF(ISNUMBER(G{r}),G{r}/D{r},"n/a")',PCT)
        f(ws,f"J{r}",f"=I{r}/D{r}",PCT); f(ws,f"L{r}",f"=I{r}+K{r}",NUM); f(ws,f"M{r}",f"=L{r}/D{r}",PCT)
    stats=[("Maximum","MAX"),("75th Percentile","Q3"),("Median","MEDIAN"),("25th Percentile","Q1"),("Minimum","MIN")]
    def statf(fn,rng):
        return {"MAX":f"=MAX({rng})","Q3":f"=QUARTILE({rng},3)","MEDIAN":f"=MEDIAN({rng})","Q1":f"=QUARTILE({rng},1)","MIN":f"=MIN({rng})"}[fn]
    for j,(lab,fn) in enumerate(stats):
        r=15+j; ws.cell(r,1,f"Peer {lab}").font=BOLD
        for c in range(1,15): ws.cell(r,c).fill=IN_FILL
        for col in ["F","H","J","M"]:
            f(ws,f"{col}{r}",statf(fn,f"{col}9:{col}13"),PCT)
    box(ws,5,1,19,14)
    # valuation
    hdr(ws,22,"VALUATION MULTIPLES",1,14)
    sub(ws,23,["Company","Ticker","Share price","Diluted shares (M)","Market cap","Cash + ST inv.","Total debt","Enterprise value","EV / Revenue","EV / EBITDA","EV / EBIT","P / E","FCF note",""])
    for i,p in enumerate(rows):
        r=24 if i==0 else 25+i   # MSFT 24, blank 25, peers 26-30
        o=7 if i==0 else 8+i
        ws.cell(r,1,f"=A{o}"); ws.cell(r,2,f"=B{o}")
        inp(ws,f"C{r}",p[11],QSRC,USD)
        inp(ws,f"D{r}",round(p[12]/p[11],1),"Source: implied = market_cap / price from stock-data get_quote 2026-10-14 (basic shares basis; quote market cap)",'#,##0.0')
        f(ws,f"E{r}",f"=C{r}*D{r}",NUM)
        inp(ws,f"F{r}",p[9],"Source: "+p[13]+" - cash & equivalents + current marketable securities (excludes long-term investments)",NUM)
        inp(ws,f"G{r}",p[10],"Source: "+p[13]+" - current + non-current debt (excludes lease liabilities)",NUM)
        f(ws,f"H{r}",f"=E{r}-F{r}+G{r}",NUM)
        f(ws,f"I{r}",f"=H{r}/D{o}",MUL); f(ws,f"J{r}",f"=H{r}/L{o}",MUL); f(ws,f"K{r}",f"=H{r}/I{o}",MUL); f(ws,f"L{r}",f"=E{r}/N{o}",MUL)
    ws["M27"]="P/E depressed by equity-investment gains"; ws["M26"]="P/E distorted by investment gains"; ws["M28"]="Q3-25 one-off tax charge"
    for j,(lab,fn) in enumerate(stats):
        r=32+j; ws.cell(r,1,f"Peer {lab}").font=BOLD
        for c in range(1,15): ws.cell(r,c).fill=IN_FILL
        for col in ["I","J","K","L"]:
            f(ws,f"{col}{r}",statf(fn,f"{col}26:{col}30"),MUL)
    for c in ["J34"]: ws[c].fill=OUT_FILL; ws[c].font=Font(name="Arial",size=10,bold=True)
    box(ws,22,1,36,14)
    # MSFT premium
    ws["A38"]="MSFT EV/EBITDA premium/(discount) to peer median"; f(ws,"J38","=J24/J34-1",PCT)
    ws["A39"]="MSFT P/E premium/(discount) to peer median"; f(ws,"L39","=L24/L34-1",PCT)
    ws["A40"]="Peer median LTM revenue growth"; f(ws,"F40","=F17",PCT)
    ws["A41"]="Peer median EBITDA margin"; f(ws,"M41","=M17",PCT)
    notes=["NOTES & METHODOLOGY",
     "Data: SEC XBRL standardized financials via stock-data MCP (get_financials); prices/market caps via stock-data get_quote (Yahoo, close 2026-10-14).",
     "LTM = sum of last four reported quarters; prior LTM = the four quarters before. Fiscal year ends differ (AAPL Sep, ORCL May, MSFT Jun, others Dec).",
     "EBITDA = Operating income (EBIT) + D&A from the cash flow statement. GOOGL LTM D&A is an estimate (see cell comment).",
     "EV = market cap - cash & current marketable securities + total debt (operating and finance lease liabilities excluded for all companies).",
     "P/E for GOOGL, AMZN is depressed by large unrealized equity-investment gains in 2026; META LTM includes a one-off tax charge. Use EV/EBITDA and EV/EBIT as the primary multiples.",
     "Statistics are calculated over the five peers only (MSFT excluded)."]
    for i,t in enumerate(notes):
        ws.cell(44+i,1,t).font=BOLD if i==0 else BLK
    ws.column_dimensions["A"].width=34
    for c in range(2,15): ws.column_dimensions[L(c)].width=15
    ws.freeze_panes="B7"

# ---------------- Comps workbook ----------------
wbC=Workbook(); wsC=wbC.active; wsC.title="Comps"; build_comps(wsC)
wbC.save("../MSFT_Comps_2026-10-14.xlsx")

# ---------------- DCF workbook ----------------
wb=Workbook(); ws=wb.active; ws.title="DCF"; ws.sheet_view.showGridLines=False
wsW=wb.create_sheet("WACC"); wsB=wb.create_sheet("Beta"); wsCp=wb.create_sheet("Comps"); build_comps(wsCp)
wsW.sheet_view.showGridLines=False; wsB.sheet_view.showGridLines=False
YRS=[f"FY{y}E" for y in range(2027,2037)]; HIST=["FY2024A","FY2025A","FY2026A"]
PC=[L(c) for c in range(5,15)]  # E..N projection columns
ws["A1"]="Microsoft Corporation (MSFT) - DCF Model"; ws["A1"].font=Font(name="Arial",size=13,bold=True)
ws["A2"]=f"Ticker: MSFT | Valuation date: {ASOF} | Fiscal year end: June 30 | USD millions except per-share"
ws["A4"]="Case selector (1=Bear, 2=Base, 3=Bull)"; ws["A4"].font=BOLD
inp(ws,"B4",2,"Assumption: case selector. 1=Bear, 2=Base, 3=Bull. Drives the 'Selected case' block and all projections.")
f(ws,"C4",'=CHOOSE(B4,"Bear case","Base case","Bull case")'); ws["C4"].font=BOLD

hdr(ws,6,"MARKET DATA & KEY INPUTS")
items=[(7,"Current share price ($)",533.69,QSRC.replace("get_quote","get_quote MSFT"),USD),
 (8,"Diluted shares (M)",7453,"Source: MSFT 10-K FY2026 (filed 2026-07-29) via stock-data get_financials - FY2026 weighted-average diluted shares 7,453M (cover page 2026-07-23: 7,425.5M basic outstanding)",'#,##0'),
 (10,"Cash + short-term investments ($M)",76843,"Source: MSFT 10-K FY2026 balance sheet 2026-06-30 via stock-data get_financials: cash 20,935 + short-term investments 55,908",NUM),
 (11,"Total debt ($M)",40294,"Source: MSFT 10-K FY2026 balance sheet 2026-06-30 via stock-data get_financials: current portion LTD 9,227 + LTD 31,067. Excludes finance/operating lease liabilities.",NUM),
 (13,"Tax rate (cash/operating)",0.19,"Assumption: 19.0%. FY2026 effective rate 19.4% (32,185/165,934), FY2025 17.6%; stock-data get_financials.",PCT),
 (14,"Change in NWC (% of change in revenue)",0.05,"Assumption: 5% of incremental revenue (MSFT runs negative working capital via unearned revenue; conservative small use of cash).",PCT),
 (17,"FY2026A diluted EPS ($)",17.95,"Source: MSFT 10-K FY2026 via stock-data get_financials (eps_diluted FY2026)",USD)]
for r,lab,val,src,fmt in items:
    ws.cell(r,1,lab); inp(ws,f"B{r}",val,src,fmt)
ws["A9"]="Market capitalisation ($M)"; f(ws,"B9","=B7*B8",NUM)
ws["A12"]="Net debt / (net cash) ($M)"; f(ws,"B12","=B11-B10",NUM)
ws["A15"]="Base WACC (from WACC sheet)"; f(ws,"B15","=WACC!B29",PCT,GRN)
ws["A16"]="Peer median EV/LTM EBITDA (from Comps sheet)"; f(ws,"B16","=Comps!J34",MUL,GRN)
box(ws,6,1,17,4,thick)

hdr(ws,19,"SCENARIO ASSUMPTIONS")
sc={
 "BEAR":dict(g=[.12,.10,.09,.08,.07,.06,.05,.045,.04,.035],m=[.44,.435,.43,.425,.42,.415,.41,.41,.41,.41],
   da=[.14,.15,.155,.16,.16,.16,.155,.15,.145,.14],cx=[.38,.34,.29,.26,.23,.21,.195,.18,.17,.165],wadj=0.005,tg=0.025),
 "BASE":dict(g=[.16,.15,.13,.12,.11,.10,.08,.07,.06,.05],m=[.46,.465,.47,.47,.47,.47,.465,.46,.46,.46],
   da=[.14,.15,.155,.16,.16,.16,.155,.15,.145,.14],cx=[.36,.32,.27,.24,.22,.20,.185,.17,.16,.155],wadj=0.0,tg=0.03),
 "BULL":dict(g=[.19,.18,.16,.15,.13,.12,.10,.09,.07,.06],m=[.47,.475,.48,.485,.49,.49,.49,.49,.49,.49],
   da=[.14,.15,.155,.16,.16,.16,.155,.15,.145,.14],cx=[.34,.29,.24,.21,.19,.175,.16,.15,.145,.14],wadj=-0.005,tg=0.035)}
why={"g":"Assumption ({c}): revenue growth path. FY2026A growth 17.8%, FY2025A 14.9% (10-K). Peer median LTM growth ~20% (Comps). Fades toward terminal growth by FY2036.",
     "m":"Assumption ({c}): EBIT margin. FY2026A 46.8%, FY2025A 45.6%. Peer median EBIT margin lower; MSFT mix supports mid-40s with D&A headwind from AI capex.",
     "da":"Assumption ({c}): D&A % revenue. FY2026A ~11.8% (39,000 est.); rises as FY2025-26 AI capex is depreciated, partly offset by 15->25yr useful-life change from FY2027 (TIKR, 2026-09-29).",
     "cx":"Assumption ({c}): cash capex % revenue. FY2026A 34.9% (115,948, 10-K); Q1 FY2027 spend guided >$50B incl. leases (TIKR 2026-09-29). Normalises toward ~15% as AI build-out matures.",
     "wadj":"Assumption ({c}): adjustment added to base WACC from WACC sheet.","tg":"Assumption ({c}): perpetuity growth rate after FY2036."}
blocks={"BEAR":20,"BASE":29,"BULL":38}
labels=[("g","Revenue growth (%)"),("m","EBIT margin (%)"),("da","D&A (% of revenue)"),("cx","CapEx (% of revenue)"),("wadj","WACC adjustment vs. base (pp)"),("tg","Terminal growth rate (%)")]
for name,r0 in blocks.items():
    hdr(ws,r0,f"{name} CASE ASSUMPTIONS"); sub(ws,r0+1,["Assumption","","",""]+YRS)
    for k,(key,lab) in enumerate(labels):
        r=r0+2+k; ws.cell(r,1,lab)
        if key in ("wadj","tg"):
            inp(ws,f"E{r}",sc[name][key],why[key].format(c=name.title()),PCT)
        else:
            for j,col in enumerate(PC):
                inp(ws,f"{col}{r}",sc[name][key][j],why[key].format(c=name.title()),PCT)
    box(ws,r0,1,r0+7,14,thin)
# Selected
hdr(ws,47,"SELECTED CASE (feeds model)"); sub(ws,48,["Assumption","","",""]+YRS)
for k,(key,lab) in enumerate(labels):
    r=49+k; ws.cell(r,1,lab)
    cols=PC if key not in ("wadj","tg") else ["E"]
    for col in cols:
        f(ws,f"{col}{r}",f"=CHOOSE($B$4,{col}{22+k},{col}{31+k},{col}{40+k})",PCT)
ws["A53"]="Selected WACC (base + adjustment)"; f(ws,"E53","=$B$15+CHOOSE($B$4,E26,E35,E44)",PCT)
ws["A54"]="Terminal growth rate (%)"
box(ws,47,1,54,14,thick)

# Projections
hdr(ws,56,"HISTORICAL & PROJECTED FREE CASH FLOW ($M)"); sub(ws,57,["Fiscal year"]+HIST+YRS)
hist={58:[245122,281724,331839],60:[109433,128528,155237],64:[22287,34153,39000],66:[44477,64551,115948],70:[118548,136162,182935]}
hsrc={58:"Source: MSFT 10-K FY2024/FY2025/FY2026 via stock-data get_financials (annual income statement)",
      60:"Source: MSFT 10-K FY2024/FY2025/FY2026 via stock-data get_financials (operating income)",
      64:"D&A ('Depreciation, amortization, and other'): FY2024 22,287 and FY2025 34,153 = model knowledge (10-K cash flow statement); FY2026 ~39,000 = stockanalysis.com / TIKR article 2026-09-29 (rounded, not verified). stock-data tool did not report this line.",
      66:"Source: MSFT 10-K via stock-data get_financials (PaymentsToAcquirePropertyPlantAndEquipment; excludes finance leases)",
      70:"Source: MSFT 10-K via stock-data get_financials (net cash from operations)"}
names={58:"Revenue",59:"  % growth",60:"EBIT (operating income)",61:"  % margin",62:"(-) Taxes on EBIT",63:"NOPAT",64:"(+) D&A",65:"  % of revenue",
 66:"(-) CapEx",67:"  % of revenue",68:"(-) Increase in NWC",69:"Unlevered free cash flow",70:"Memo: operating cash flow (reported)",71:"Memo: reported FCF (OCF - CapEx)",72:"EBITDA (EBIT + D&A)"}
for r,n in names.items(): ws.cell(r,1,n)
for r,vals in hist.items():
    for j,col in enumerate(["B","C","D"]): inp(ws,f"{col}{r}",vals[j],hsrc[r],NUM)
f(ws,"B59",'=B58/245122*0+B58/211915-1',PCT); ws["B59"].comment=Comment("FY2023 revenue 211,915 per MSFT 10-K via stock-data get_financials","Analyst")
f(ws,"B59","=B58/211915-1",PCT)
for col,prev in [("C","B"),("D","C")]: f(ws,f"{col}59",f"={col}58/{prev}58-1",PCT)
for col in ["B","C","D"]:
    f(ws,f"{col}61",f"={col}60/{col}58",PCT); f(ws,f"{col}65",f"={col}64/{col}58",PCT); f(ws,f"{col}67",f"={col}66/{col}58",PCT)
    f(ws,f"{col}71",f"={col}70-{col}66",NUM); f(ws,f"{col}72",f"={col}60+{col}64",NUM)
for j,col in enumerate(PC):
    prev="D" if j==0 else PC[j-1]
    f(ws,f"{col}58",f"={prev}58*(1+{col}49)",NUM); f(ws,f"{col}59",f"={col}58/{prev}58-1",PCT)
    f(ws,f"{col}60",f"={col}58*{col}50",NUM); f(ws,f"{col}61",f"={col}60/{col}58",PCT)
    f(ws,f"{col}62",f"={col}60*$B$13",NUM); f(ws,f"{col}63",f"={col}60-{col}62",NUM)
    f(ws,f"{col}64",f"={col}58*{col}51",NUM); f(ws,f"{col}65",f"={col}64/{col}58",PCT)
    f(ws,f"{col}66",f"={col}58*{col}52",NUM); f(ws,f"{col}67",f"={col}66/{col}58",PCT)
    f(ws,f"{col}68",f"=({col}58-{prev}58)*$B$14",NUM)
    f(ws,f"{col}69",f"={col}63+{col}64-{col}66-{col}68",NUM); ws[f"{col}69"].font=BOLD
    f(ws,f"{col}72",f"={col}60+{col}64",NUM)
box(ws,56,1,72,14,thick)

hdr(ws,74,"DISCOUNTING (mid-year convention)"); sub(ws,75,["",""," ",""]+YRS)
ws["A76"]="Discount period (years)"; ws["A77"]="Discount factor"; ws["A78"]="PV of unlevered FCF"
for j,col in enumerate(PC):
    f(ws,f"{col}76","=0.5" if j==0 else f"={PC[j-1]}76+1","0.0")
    f(ws,f"{col}77",f"=1/(1+$E$53)^{col}76","0.0000"); f(ws,f"{col}78",f"={col}69*{col}77",NUM)
ws["E76"].font=BLUE; ws["E76"].comment=Comment("Assumption: mid-year convention from start of FY2027 (2026-07-01); the ~3.5 months elapsed to the valuation date are ignored (slightly conservative).","Analyst")
box(ws,74,1,78,14,thick)

hdr(ws,80,"TERMINAL VALUE",1,4)
tv=[(81,"Terminal-year FCF x (1+g)","=N69*(1+E54)",NUM),(82,"Terminal value - perpetuity growth","=B81/(E53-E54)",NUM),
    (83,"PV of terminal value","=B82*N77",NUM),(84,"Implied terminal EV / EBITDA (FY2036)","=B82/N72",MUL),
    (85,"Cross-check: TV at peer median EV/EBITDA","=N72*B16",NUM),(86,"Cross-check: PV of exit-multiple TV","=B85*N77",NUM)]
for r,lab,fm,fmt in tv: ws.cell(r,1,lab); f(ws,f"B{r}",fm,fmt)
box(ws,80,1,86,4,thick)

hdr(ws,88,"VALUATION SUMMARY",1,4)
vs=[(89,"Sum of PV of FCF (FY2027-FY2036)","=SUM(E78:N78)",NUM),(90,"PV of terminal value","=B83",NUM),(91,"Enterprise value","=B89+B90",NUM),
 (92,"Terminal value as % of EV","=B90/B91",PCT),(93,"(-) Net debt / (+) net cash","=-B12",NUM),(94,"Equity value","=B91+B93",NUM),
 (95,"Diluted shares (M)","=B8",'#,##0'),(96,"IMPLIED VALUE PER SHARE ($)","=B94/B95",USD),(97,"Current share price ($)","=B7",USD),
 (98,"Implied upside / (downside)","=B96/B97-1",PCT),(99,"DCF-implied EV / FY2026A EBITDA","=B91/D72",MUL),(100,"Peer median EV / LTM EBITDA","=B16",MUL),
 (101,"DCF multiple premium / (discount) to peers","=B99/B100-1",PCT),(102,"DCF-implied P/E on FY2026A EPS","=B96/B17",MUL),
 (103,"Cross-check: value/share using exit multiple TV ($)","=(B89+B86-B12)/B8",USD),(104,"Market-implied EV / FY2026A EBITDA","=(B9+B12)/D72",MUL)]
for r,lab,fm,fmt in vs: ws.cell(r,1,lab); f(ws,f"B{r}",fm,fmt)
for r in (91,94,96,98):
    ws[f"A{r}"].font=BOLD; ws[f"B{r}"].fill=OUT_FILL; ws[f"B{r}"].font=BOLD
box(ws,88,1,104,4,thick)

# Scenario summary engine (all three cases live side by side)
hdr(ws,106,"SCENARIO SUMMARY - all three cases calculated simultaneously"); sub(ws,107,["Line","","",""]+YRS)
eng={}
r=108
for name,r0 in blocks.items():
    g,m,da,cx,wa,tg=[r0+2+k for k in range(6)]
    R,FCF,DF=r,r+1,r+2
    ws.cell(R,1,f"{name.title()}: revenue"); ws.cell(FCF,1,f"{name.title()}: unlevered FCF"); ws.cell(DF,1,f"{name.title()}: discount factor")
    for j,col in enumerate(PC):
        prev="$D$58" if j==0 else f"{PC[j-1]}{R}"
        f(ws,f"{col}{R}",f"={prev}*(1+{col}{g})",NUM)
        f(ws,f"{col}{FCF}",f"={col}{R}*{col}{m}*(1-$B$13)+{col}{R}*{col}{da}-{col}{R}*{col}{cx}-({col}{R}-{prev})*$B$14",NUM)
        f(ws,f"{col}{DF}",f"=1/(1+$B$15+$E${wa})^{col}$76","0.0000")
    eng[name]=(R,FCF,DF,wa,tg); r+=3
sub(ws,118,["Output","Bear","Base","Bull"])
outs=["Enterprise value ($M)","PV of terminal value ($M)","TV % of EV","Equity value ($M)","Value per share ($)","Upside / (downside)","WACC","Terminal growth","Implied EV / FY2026A EBITDA"]
for i,o in enumerate(outs): ws.cell(119+i,1,o)
for k,name in enumerate(["BEAR","BASE","BULL"]):
    c=L(2+k); R,FCF,DF,wa,tg=eng[name]
    f(ws,f"{c}120",f"=N{FCF}*(1+$E${tg})/($B$15+$E${wa}-$E${tg})*N{DF}",NUM)
    f(ws,f"{c}119",f"=SUMPRODUCT(E{FCF}:N{FCF},E{DF}:N{DF})+{c}120",NUM)
    f(ws,f"{c}121",f"={c}120/{c}119",PCT); f(ws,f"{c}122",f"={c}119-$B$12",NUM)
    f(ws,f"{c}123",f"={c}122/$B$8",USD); f(ws,f"{c}124",f"={c}123/$B$7-1",PCT)
    f(ws,f"{c}125",f"=$B$15+$E${wa}",PCT); f(ws,f"{c}126",f"=$E${tg}",PCT); f(ws,f"{c}127",f"={c}119/$D$72",MUL)
    ws[f"{c}123"].fill=OUT_FILL; ws[f"{c}123"].font=BOLD
ws["A128"]="Check: selected-case value = scenario engine value (should be 0)"; f(ws,"B128","=ROUND(B96-INDEX(B123:D123,1,B4),6)","0.000000")
box(ws,106,1,128,14,thick)

# ---------------- Sensitivity tables ----------------
def axis(center,step): return [f"={center}-{2*step}",f"={center}-{step}",f"={center}",f"={center}+{step}",f"={center}+{2*step}"]
def color(ws,rng): ws.conditional_formatting.add(rng,ColorScaleRule(start_type="min",start_color="F8696B",mid_type="percentile",mid_value=50,mid_color="FFFFFF",end_type="max",end_color="63BE7B"))
# T1 WACC x g
hdr(ws,131,"SENSITIVITY 1: IMPLIED VALUE PER SHARE ($) - WACC (rows) x TERMINAL GROWTH (columns), selected case",1,7)
ws["A132"]="WACC \\ Terminal g"; ws["A132"].font=BOLD
for j,v in enumerate(axis("$E$54",0.005)): f(ws,f"{L(2+j)}132",v,PCT); ws[f"{L(2+j)}132"].fill=SUB_FILL; ws[f"{L(2+j)}132"].font=BOLD
for i,v in enumerate(axis("$E$53",0.005)):
    r=133+i; f(ws,f"A{r}",v,PCT); ws[f"A{r}"].fill=SUB_FILL; ws[f"A{r}"].font=BOLD
    for j in range(5):
        c=L(2+j)
        f(ws,f"{c}{r}",f"=(SUMPRODUCT($E$69:$N$69/(1+$A{r})^$E$76:$N$76)+$N$69*(1+{c}$132)/($A{r}-{c}$132)/(1+$A{r})^$N$76-$B$12)/$B$8",USD)
ws["D135"].fill=OUT_FILL; ws["D135"].font=BOLD; color(ws,"B133:F137"); box(ws,131,1,137,7,thick)
ws["A138"]="Check: centre cell - model value (should be 0)"; f(ws,"B138","=ROUND(D135-B96,6)","0.000000")

# T2 growth shift x margin shift - helper engine
hdr(ws,140,"SENSITIVITY 2: IMPLIED VALUE PER SHARE ($) - REVENUE GROWTH SHIFT (rows, pp every year) x EBIT MARGIN SHIFT (columns, pp every year)",1,7)
ws["A141"]="Growth shift \\ Margin shift"; ws["A141"].font=BOLD
dg=[-0.04,-0.02,0,0.02,0.04]; dm=[-0.04,-0.02,0,0.02,0.04]
for j,v in enumerate(dm):
    c=ws.cell(141,2+j,v); c.font=BLUE; c.fill=SUB_FILL; c.number_format='+0.0%;-0.0%;0.0%'
ws["B141"].comment=Comment("Assumption: margin shift added to every projection year's selected-case EBIT margin.","Analyst")
helper0=160
hdr(ws,158,"Sensitivity 2 helper engine: revenue path and FCF (at selected-case margins) for each growth shift",1,14)
for i,v in enumerate(dg):
    r=142+i; c=ws.cell(r,1,v); c.font=BLUE; c.fill=SUB_FILL; c.number_format='+0.0%;-0.0%;0.0%'
    Rr=helper0+2*i; Fr=Rr+1
    ws.cell(Rr,1,f"Revenue, growth shift row {i+1}"); ws.cell(Fr,1,f"FCF at base margin, growth shift row {i+1}")
    for j,col in enumerate(PC):
        prev="$D$58" if j==0 else f"{PC[j-1]}{Rr}"
        f(ws,f"{col}{Rr}",f"={prev}*(1+{col}$49+$A${r})",NUM)
        f(ws,f"{col}{Fr}",f"={col}{Rr}*{col}$50*(1-$B$13)+{col}{Rr}*{col}$51-{col}{Rr}*{col}$52-({col}{Rr}-{prev})*$B$14",NUM)
    for j in range(5):
        c=L(2+j)
        f(ws,f"{c}{r}",f"=(SUMPRODUCT(($E${Fr}:$N${Fr}+$E${Rr}:$N${Rr}*{c}$141*(1-$B$13))*$E$77:$N$77)+($N${Fr}+$N${Rr}*{c}$141*(1-$B$13))*(1+$E$54)/($E$53-$E$54)*$N$77-$B$12)/$B$8",USD)
ws["A142"].comment=Comment("Assumption: growth shift added to every projection year's selected-case revenue growth.","Analyst")
ws["D144"].fill=OUT_FILL; ws["D144"].font=BOLD; color(ws,"B142:F146"); box(ws,140,1,146,7,thick)
ws["A147"]="Check: centre cell - model value (should be 0)"; f(ws,"B147","=ROUND(D144-B96,6)","0.000000")

# T3 Beta x Rf
hdr(ws,149,"SENSITIVITY 3: IMPLIED VALUE PER SHARE ($) - BETA (rows) x RISK-FREE RATE (columns), selected case",1,7)
ws["A150"]="Beta \\ Risk-free"; ws["A150"].font=BOLD
for j,v in enumerate(axis("WACC!$B$4",0.0025)): f(ws,f"{L(2+j)}150",v,PCT,GRN); ws[f"{L(2+j)}150"].fill=SUB_FILL
for i,v in enumerate(axis("WACC!$B$5",0.1)):
    r=151+i; f(ws,f"A{r}",v,"0.00",GRN); ws[f"A{r}"].fill=SUB_FILL
    for j in range(5):
        c=L(2+j)
        w=f"(WACC!$B$27*({c}$150+$A{r}*WACC!$B$6)+WACC!$B$28*({c}$150+WACC!$B$11)*(1-$B$13)+CHOOSE($B$4,$E$26,$E$35,$E$44))"
        f(ws,f"{c}{r}",f"=(SUMPRODUCT($E$69:$N$69/(1+{w})^$E$76:$N$76)+$N$69*(1+$E$54)/({w}-$E$54)/(1+{w})^$N$76-$B$12)/$B$8",USD)
ws["D153"].fill=OUT_FILL; ws["D153"].font=BOLD; color(ws,"B151:F155"); box(ws,149,1,155,7,thick)
ws["A156"]="Check: centre cell - model value (should be 0)"; f(ws,"B156","=ROUND(D153-B96,6)","0.000000")
ws.column_dimensions["A"].width=48
for c in range(2,15): ws.column_dimensions[L(c)].width=13
ws.freeze_panes="B5"

# ---------------- Beta sheet ----------------
d=json.load(open("beta_data.json"))
wsB["A1"]="MSFT 5-year beta vs. S&P 500 - monthly total-return (MSFT adj. close) vs. index price returns"; wsB["A1"].font=BOLD
sub(wsB,3,["Month-end","MSFT adj. close","S&P 500 close","MSFT return","S&P 500 return"])
for i,(mo,a,b) in enumerate(zip(d["months"],d["msft"],d["spx"])):
    r=4+i; wsB.cell(r,1,mo)
    inp(wsB,f"B{r}",a,"Source: stock-data get_price_history MSFT 10y 1mo (Yahoo adj_close), fetched 2026-10-15" if i==0 else "Source: stock-data get_price_history MSFT monthly adj_close",'0.00')
    inp(wsB,f"C{r}",b,"Source: stock-data get_price_history ^GSPC 10y 1mo (close), fetched 2026-10-15" if i==0 else "Source: stock-data get_price_history ^GSPC monthly close",'#,##0.00')
    if i>0:
        f(wsB,f"D{r}",f"=B{r}/B{r-1}-1",PCT); f(wsB,f"E{r}",f"=C{r}/C{r-1}-1",PCT)
last=4+len(d["months"])-1
wsB["G4"]="Beta (SLOPE, raw)"; f(wsB,"H4",f"=SLOPE(D5:D{last},E5:E{last})","0.000"); wsB["H4"].fill=OUT_FILL; wsB["H4"].font=BOLD
wsB["G5"]="Observations (monthly returns)"; f(wsB,"H5",f"=COUNT(D5:D{last})","0")
wsB["G6"]="R-squared"; f(wsB,"H6",f"=RSQ(D5:D{last},E5:E{last})","0.000")
wsB["G7"]="Blume-adjusted beta (memo)"; f(wsB,"H7","=0.67*H4+0.33","0.000")
wsB["G8"]="Window: 2021-09 to 2026-09 month-end closes (60 complete months; partial Oct-2026 excluded)"
wsB.column_dimensions["A"].width=12
for c in "BCDE": wsB.column_dimensions[c].width=15
wsB.column_dimensions["G"].width=34

# ---------------- WACC sheet ----------------
wsW["A1"]="MSFT - Weighted Average Cost of Capital"; wsW["A1"].font=Font(name="Arial",size=12,bold=True)
hdr(wsW,3,"COST OF EQUITY (CAPM)",1,4)
wsW["A4"]="Risk-free rate (10Y UST, 2026-10-13)"; inp(wsW,"B4",0.0521,"Source: FRED DGS10 via stock-data get_macro_series, latest observation 2026-10-13 = 5.21%",PCT)
wsW["A5"]="Beta (5Y monthly vs S&P 500, raw)"; f(wsW,"B5","=Beta!H4","0.00",GRN)
wsW["A6"]="Equity risk premium"; inp(wsW,"B6",0.05,"Assumption: 5.0% mature-market ERP (skill default range 5.0-6.0%; Damodaran implied ERP has run ~4-4.5% - model knowledge).",PCT)
wsW["A7"]="Cost of equity"; f(wsW,"B7","=B4+B5*B6",PCT)
box(wsW,3,1,7,4,thick)
hdr(wsW,9,"COST OF DEBT",1,4)
wsW["A10"]="Credit rating"; inp(wsW,"B10","AAA","Source: model knowledge - MSFT rated AAA (S&P) / Aaa (Moody's)")
wsW["A11"]="Spread over 10Y UST"; inp(wsW,"B11",0.006,"Assumption: ~60bp AAA corporate spread over 10Y UST (model knowledge; not verified against live MSFT bond yields).",PCT)
wsW["A12"]="Pre-tax cost of debt"; f(wsW,"B12","=B4+B11",PCT)
wsW["A13"]="Tax rate"; f(wsW,"B13","=DCF!B13",PCT,GRN)
wsW["A14"]="After-tax cost of debt"; f(wsW,"B14","=B12*(1-B13)",PCT)
wsW["A15"]="Memo: FY2026 interest expense / avg. debt"; inp(wsW,"C15",3051,"Source: MSFT 10-K FY2026 via stock-data get_financials (interest expense; includes finance-lease interest)",NUM)
f(wsW,"B15","=C15/AVERAGE(43151,DCF!B11)",PCT); wsW["B15"].comment=Comment("FY2025 total debt 43,151 = 2,999 + 40,152 (10-K FY2025 via stock-data). Inflated by finance-lease interest, so not used.","Analyst")
box(wsW,9,1,15,4,thick)
hdr(wsW,17,"CAPITAL STRUCTURE (market values)",1,4)
cs=[(18,"Share price ($)","=DCF!B7",USD),(19,"Diluted shares (M)","=DCF!B8",'#,##0'),(20,"Market capitalisation ($M)","=B18*B19",NUM),
    (21,"Total debt ($M)","=DCF!B11",NUM),(22,"Cash + ST investments ($M)","=DCF!B10",NUM),(23,"Net debt / (net cash) ($M)","=B21-B22",NUM),(24,"Total capital (E + D) ($M)","=B20+B21",NUM)]
for r,lab,fm,fmt in cs:
    wsW.cell(r,1,lab); f(wsW,f"B{r}",fm,fmt,GRN if "DCF!" in fm else BLK)
box(wsW,17,1,24,4,thick)
hdr(wsW,26,"WACC",1,4); sub(wsW,26,["WACC","Weight","Cost","Contribution"])
wsW["A27"]="Equity"; f(wsW,"B27","=B20/B24",PCT); f(wsW,"C27","=B7",PCT); f(wsW,"D27","=B27*C27","0.00%")
wsW["A28"]="Debt (gross)"; f(wsW,"B28","=B21/B24",PCT); f(wsW,"C28","=B14",PCT); f(wsW,"D28","=B28*C28","0.00%")
wsW["A29"]="WEIGHTED AVERAGE COST OF CAPITAL"; f(wsW,"B29","=D27+D28",PCT); wsW["B29"].fill=OUT_FILL; wsW["B29"].font=BOLD; wsW["A29"].font=BOLD
wsW["A31"]="Gross debt weights used (MSFT is net cash); lease liabilities excluded from debt and from the equity bridge for consistency."
box(wsW,26,1,29,4,thick)
wsW.column_dimensions["A"].width=42
for c in "BCD": wsW.column_dimensions[c].width=14
wb.save("../MSFT_DCF_Model_2026-10-14.xlsx")
print("ok")
