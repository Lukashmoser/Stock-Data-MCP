import sys,subprocess,os,shutil,json
from openpyxl import load_workbook
src=os.path.abspath(sys.argv[1]); out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"recalc_out"); os.makedirs(out,exist_ok=True)
subprocess.run(["soffice","--headless","-env:UserInstallation=file:///home/claude/results/2026-10-13/runs/FA-04/B/_work/lo_profile","--convert-to","xlsx","--outdir",out,src],check=True,capture_output=True,timeout=120)
calc=os.path.join(out,os.path.basename(src))
wb=load_workbook(calc,data_only=True); errs={}; n=0
wbf=load_workbook(src)
for ws in wbf:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith("="): n+=1
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith("#"): errs.setdefault(c.value,[]).append(f"{ws.title}!{c.coordinate}")
print(json.dumps({"status":"success" if not errs else "errors_found","total_formulas":n,"total_errors":sum(len(v) for v in errs.values()),"error_summary":errs})[:2000])
print(calc)
