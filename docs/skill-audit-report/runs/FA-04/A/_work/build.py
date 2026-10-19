import json
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

OUT='/home/claude/results/2026-10-13/runs/FA-04/A/'
ASOF='2026-10-14'
FN='Calibri'
DARK=PatternFill('solid',fgColor='1F4E79'); LIGHT=PatternFill('solid',fgColor='D9E1F2')
GREY=PatternFill('solid',fgColor='F2F2F2'); MED=PatternFill('solid',fgColor='BDD7EE')
BLUE=Font(name=FN,color='0000FF',size=10); BLACK=Font(name=FN,color='000000',size=10)
GREEN=Font(name=FN,color='008000',size=10); BOLD=Font(name=FN,bold=True,size=10)
WH=Font(name=FN,bold=True,color='FFFFFF',size=11)
thin=Side(style='thin'); med=Side(style='medium'); thick=Side(style='thick')
PCT='0.0%;(0.0%);"-"'; MIL='#,##0;(#,##0);"-"'; PS='$#,##0.00'; MULT='0.0"x"'
SRC_PR='Source: Microsoft FY26 Q4 earnings release (8-K exhibit 99.1), 2026-07-29, financial statements, https://microsoft.gcs-web.com/node/35101/html'
SRC_SA=lambda t:f'Source: stockanalysis.com /stocks/{t}/statistics/, accessed {ASOF}; close {ASOF}; LTM through latest reported quarter'

def inp(ws,ref,val,fmt=None,comment=None):
    c=ws[ref]; c.value=val; c.font=BLUE; c.fill=GREY
    if fmt: c.number_format=fmt
    if comment: c.comment=Comment(comment,'Analyst',width=320,height=110)
    c.alignment=Alignment(horizontal='center'); return c
def fml(ws,ref,val,fmt=None,font=None,bold=False):
    c=ws[ref]; c.value=val; c.font=font or (Font(name=FN,bold=True,size=10) if bold else BLACK)
    if fmt: c.number_format=fmt
    c.alignment=Alignment(horizontal='center'); return c
def lab(ws,ref,val,bold=False,indent=0):
    c=ws[ref]; c.value=val; c.font=BOLD if bold else BLACK; c.alignment=Alignment(indent=indent)
def section(ws,row,text,c1=1,c2=15):
    ws.cell(row=row,column=c1,value=text)
    for c in range(c1,c2+1):
        x=ws.cell(row=row,column=c); x.fill=DARK; x.font=WH
    ws.merge_cells(start_row=row,start_column=c1,end_row=row,end_column=c2)
def colhdr(ws,row,vals,c1=1):
    for i,v in enumerate(vals):
        x=ws.cell(row=row,column=c1+i,value=v); x.fill=LIGHT; x.font=BOLD; x.alignment=Alignment(horizontal='center',wrap_text=True)
def box(ws,r1,r2,c1,c2,side=med):
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            cell=ws.cell(row=r,column=c); b=cell.border
            cell.border=Border(left=side if c==c1 else b.left,right=side if c==c2 else b.right,
                               top=side if r==r1 else b.top,bottom=side if r==r2 else b.bottom)
def output(ws,ref):
    ws[ref].fill=MED; ws[ref].font=Font(name=FN,bold=True,size=10)

# ---------------- COMPS ----------------
PEERS=[('Alphabet','GOOGL',445870,271520,147630,173160,244120,53270,347.64,12230,120790,242470),
       ('Amazon','AMZN',775680,393810,93710,168910,135280,-11630,254.08,10790,251640,122990),
       ('Meta Platforms','META',228250,186590,86930,109660,68100,40980,737.54,2550,112320,90260),
       ('Apple','AAPL',466820,227120,154860,167960,128930,136680,342.96,14590,84340,146520),
       ('Oracle','ORCL',71780,45910,24600,34430,18740,-28720,130.55,3030,169140,37080)]
MSFT=('Microsoft (target)','MSFT',331839,225465,155237,193771,133749,66987,533.69,7430,128810,76843)
def build_comps(ws):
    ws.title='Comps'
    ws['A1']='BIG TECH - COMPARABLE COMPANY ANALYSIS (peer set for MSFT DCF)'; ws['A1'].font=Font(name=FN,bold=True,size=14)
    ws['A2']='Alphabet (GOOGL) • Amazon (AMZN) • Meta Platforms (META) • Apple (AAPL) • Oracle (ORCL) | Target: Microsoft (MSFT)'
    ws['A3']=f'Prices as of {ASOF} close | LTM financials through each company\'s latest reported quarter | USD millions except per-share amounts and ratios'
    section(ws,5,'OPERATING METRICS (LTM)',1,12)
    colhdr(ws,6,['Company','Ticker','Revenue','Gross Profit','Gross Margin','Operating Income','Operating Margin','EBITDA','EBITDA Margin','Net Income','Free Cash Flow','FCF Margin'])
    def oprow(r,d,src):
        lab(ws,f'A{r}',d[0]); ws[f'B{r}']=d[1]; ws[f'B{r}'].alignment=Alignment(horizontal='center')
        for col,v in zip('CDFHJK',d[2:8]): inp(ws,f'{col}{r}',v,MIL,src)
        fml(ws,f'E{r}',f'=D{r}/C{r}',PCT); fml(ws,f'G{r}',f'=F{r}/C{r}',PCT)
        fml(ws,f'I{r}',f'=H{r}/C{r}',PCT); fml(ws,f'L{r}',f'=K{r}/C{r}',PCT)
    for i,p in enumerate(PEERS): oprow(7+i,p,SRC_SA(p[1].lower()))
    stats=[('Maximum','MAX({})'),('75th Percentile','QUARTILE({},3)'),('Median','MEDIAN({})'),('25th Percentile','QUARTILE({},1)'),('Minimum','MIN({})')]
    for j,(n,f) in enumerate(stats):
        r=13+j; lab(ws,f'A{r}',n,bold=True)
        for c in range(1,13): ws.cell(row=r,column=c).fill=GREY
        for col in 'EGIL': fml(ws,f'{col}{r}','='+f.format(f'{col}7:{col}11'),PCT)
    oprow(19,MSFT,SRC_PR+' (EBITDA = operating income 155,237 + D&A 38,534)')
    lab(ws,'A19','Microsoft (target, excluded from stats)',bold=True)
    box(ws,5,19,1,12)
    section(ws,21,'VALUATION MULTIPLES',1,12)
    colhdr(ws,22,['Company','Ticker','Share Price ($)','Shares Out (M)','Market Cap','Total Debt','Cash & ST Inv.','Enterprise Value','EV / Revenue','EV / EBITDA','P / E','FCF Yield'])
    def valrow(r,orow,d,src):
        lab(ws,f'A{r}',d[0]); ws[f'B{r}']=d[1]; ws[f'B{r}'].alignment=Alignment(horizontal='center')
        inp(ws,f'C{r}',d[8],PS,src); inp(ws,f'D{r}',d[9],MIL,src); inp(ws,f'F{r}',d[10],MIL,src+'; total debt incl. lease liabilities'); inp(ws,f'G{r}',d[11],MIL,src)
        fml(ws,f'E{r}',f'=C{r}*D{r}',MIL); fml(ws,f'H{r}',f'=E{r}+F{r}-G{r}',MIL)
        fml(ws,f'I{r}',f'=H{r}/C{orow}',MULT); fml(ws,f'J{r}',f'=H{r}/H{orow}',MULT)
        fml(ws,f'K{r}',f'=E{r}/J{orow}',MULT); fml(ws,f'L{r}',f'=K{orow}/E{r}',PCT)
    for i,p in enumerate(PEERS): valrow(23+i,7+i,p,SRC_SA(p[1].lower()))
    for j,(n,f) in enumerate(stats):
        r=29+j; lab(ws,f'A{r}',n,bold=True)
        for c in range(1,13): ws.cell(row=r,column=c).fill=GREY
        for col in 'IJKL': fml(ws,f'{col}{r}','='+f.format(f'{col}23:{col}27'),PCT if col=='L' else MULT)
    valrow(35,19,MSFT,'Source: stockanalysis.com /stocks/msft/statistics/ (price close 2026-10-14, 7.43bn shares, total debt incl. leases 128.81bn); cash+STI from FY26 Q4 release (20,935 + 55,908)')
    lab(ws,'A35','Microsoft (target, excluded from stats)',bold=True)
    for c in 'IJ': output(ws,f'{c}31')
    box(ws,21,35,1,12)
    notes=['NOTES & METHODOLOGY',
      'Data: peer LTM financials, prices, shares, debt and cash from stockanalysis.com statistics pages accessed 2026-10-14 (close 2026-10-14). MSFT financials from FY26 (June 30, 2026) earnings release.',
      'EBITDA = operating income + D&A (provider definition for peers). FCF = operating cash flow - capex. EV = market cap + total debt (incl. lease liabilities) - cash & short-term investments.',
      'Statistics exclude MSFT. Peer GOOGL/AMZN net income includes large unrealised gains on AI-lab investments, which depresses their P/E; EV/EBITDA is the primary multiple.',
      'Peer fiscal periods differ (AAPL Sep FYE, ORCL May FYE, MSFT Jun FYE); LTM basis used for all. AMZN and ORCL FCF negative due to AI capex.']
    for i,t in enumerate(notes):
        ws[f'A{38+i}']=t; ws[f'A{38+i}'].font=BOLD if i==0 else BLACK
    ws.column_dimensions['A'].width=34
    for c in range(2,13): ws.column_dimensions[CL(c)].width=14
    ws.row_dimensions[6].height=30; ws.row_dimensions[22].height=30
    ws.freeze_panes='C7'

wb=Workbook(); build_comps(wb.active); wb.save(OUT+'MSFT_Comps_2026-10-14.xlsx')

# ---------------- DCF workbook ----------------
wb=Workbook(); dcf=wb.active; dcf.title='DCF'
wacc=wb.create_sheet('WACC'); beta=wb.create_sheet('Beta'); scen=wb.create_sheet('Scenarios'); comps=wb.create_sheet('Comps'); build_comps(comps)

# Beta sheet
pr=json.load(open(OUT+'_work/prices.json'))
beta['A1']='MSFT 5-YEAR MONTHLY BETA vs S&P 500'; beta['A1'].font=Font(name=FN,bold=True,size=14)
beta['A2']='Month-end adjusted closes Sep-2021 to Sep-2026 (61 prices -> 60 monthly returns). Oct-2026 partial month excluded.'
colhdr(beta,4,['Month','MSFT Adj Close','S&P 500 Close','MSFT Return','S&P 500 Return'])
for i,((d,m),(_,s)) in enumerate(zip(pr['msft'],pr['spx'])):
    r=5+i; beta[f'A{r}']=d; beta[f'A{r}'].alignment=Alignment(horizontal='center')
    inp(beta,f'B{r}',m,'#,##0.00','Source: Yahoo Finance MSFT monthly history (adjusted close), accessed 2026-10-14' if i==0 else None)
    inp(beta,f'C{r}',s,'#,##0.00','Source: Yahoo Finance ^GSPC monthly history (price index), accessed 2026-10-14' if i==0 else None)
    if i>0:
        fml(beta,f'D{r}',f'=B{r}/B{r-1}-1',PCT); fml(beta,f'E{r}',f'=C{r}/C{r-1}-1',PCT)
last=5+len(pr['msft'])-1
lab(beta,'G4','Regression output',bold=True)
lab(beta,'G5','Beta (SLOPE)'); fml(beta,'H5',f'=SLOPE(D6:D{last},E6:E{last})','0.00'); output(beta,'H5')
lab(beta,'G6','Observations'); fml(beta,'H6',f'=COUNT(D6:D{last})','0')
lab(beta,'G7','Correlation'); fml(beta,'H7',f'=CORREL(D6:D{last},E6:E{last})','0.00')
lab(beta,'G8','R-squared'); fml(beta,'H8','=H7^2','0.00')
lab(beta,'G9','Blume-adjusted beta (memo)'); fml(beta,'H9','=0.67*H5+0.33','0.00')
box(beta,4,last,1,5); box(beta,4,9,7,8)
beta.column_dimensions['A'].width=12; beta.column_dimensions['G'].width=26
for c in 'BCDEH': beta.column_dimensions[c].width=15

# WACC sheet
w=wacc
w['A1']='MICROSOFT - WEIGHTED AVERAGE COST OF CAPITAL'; w['A1'].font=Font(name=FN,bold=True,size=14)
w['A2']=f'As of {ASOF} | USD millions'
section(w,3,'COST OF EQUITY (CAPM)',1,4)
lab(w,'A4','Risk-free rate (10Y UST par yield)'); inp(w,'B4',0.0519,'0.00%','Source: U.S. Treasury Daily Par Yield Curve, 10 Yr = 5.19% on 2026-10-14, https://home.treasury.gov/resource-center/data-chart-center/interest-rates/TextView?type=daily_treasury_yield_curve&field_tdr_date_value_month=202610 (matches ycharts)')
lab(w,'A5','Risk-free rate date'); inp(w,'B5','2026-10-14',None,'Date of 10Y yield observation')
lab(w,'A6','Beta (5Y monthly vs S&P 500)'); fml(w,'B6','=Beta!H5','0.00',GREEN)
lab(w,'A7','Beta observations (months)'); fml(w,'B7','=Beta!H6','0',GREEN)
lab(w,'A8','Equity risk premium'); inp(w,'B8',0.05,'0.0%','Assumption: 5.0% mature-market ERP (mid-point of 4.5-5.5% practitioner range; Damodaran implied ERP ~4.3-4.6% in 2025-26). Not verified for Oct-2026.')
lab(w,'A9','Cost of equity',True); fml(w,'B9','=B4+B6*B8','0.00%',bold=True)
box(w,3,9,1,4)
section(w,11,'COST OF DEBT',1,4)
lab(w,'A12','Credit rating (S&P / Moody\'s)'); inp(w,'B12','AAA / Aaa',None,'Source: model knowledge - MSFT rated AAA (S&P) / Aaa (Moody\'s)')
lab(w,'A13','Spread over 10Y UST'); inp(w,'B13',0.004,'0.00%','Assumption: ~40bp long-dated AAA industrial spread; MSFT bonds not quoted directly')
lab(w,'A14','Pre-tax cost of debt'); fml(w,'B14','=B4+B13','0.00%')
lab(w,'A15','Tax rate'); inp(w,'B15',0.19,'0.0%','Assumption: FY26 effective rate 19.4% (32,185 / 165,934, FY26 Q4 release); 19.0% normalised')
lab(w,'A16','After-tax cost of debt',True); fml(w,'B16','=B14*(1-B15)','0.00%',bold=True)
box(w,11,16,1,4)
section(w,18,'CAPITAL STRUCTURE (market values)',1,4)
lab(w,'A19','Share price ($)'); fml(w,'B19','=DCF!B8',PS,GREEN)
lab(w,'A20','Diluted shares (M)'); fml(w,'B20','=DCF!B9',MIL,GREEN)
lab(w,'A21','Market capitalisation'); fml(w,'B21','=B19*B20',MIL)
lab(w,'A22','Total debt incl. lease liabilities'); inp(w,'B22',128810,MIL,'Source: stockanalysis.com MSFT statistics (total debt 128.81bn incl. finance & operating lease liabilities, as of 2026-06-30), accessed 2026-10-14')
lab(w,'A23','  memo: financial debt (bonds)',indent=1); inp(w,'B23',40294,MIL,SRC_PR+': current portion of LTD 9,227 + long-term debt 31,067')
lab(w,'A24','  memo: lease liabilities (derived)',indent=1); fml(w,'B24','=B22-B23',MIL)
lab(w,'A25','Cash + short-term investments'); inp(w,'B25',76843,MIL,SRC_PR+': cash 20,935 + short-term investments 55,908 at 2026-06-30')
lab(w,'A26','Net debt / (net cash)'); fml(w,'B26','=B22-B25',MIL)
lab(w,'A27','Enterprise value (market)'); fml(w,'B27','=B21+B26',MIL)
box(w,18,27,1,4)
section(w,29,'WACC CALCULATION',1,4)
colhdr(w,30,['Component','Weight','Cost','Contribution'])
lab(w,'A31','Equity'); fml(w,'B31','=B21/(B21+B22)',PCT); fml(w,'C31','=B9','0.00%'); fml(w,'D31','=B31*C31','0.00%')
lab(w,'A32','Debt (gross)'); fml(w,'B32','=B22/(B21+B22)',PCT); fml(w,'C32','=B16','0.00%'); fml(w,'D32','=B32*C32','0.00%')
lab(w,'A33','WACC',True); fml(w,'D33','=D31+D32','0.00%'); output(w,'D33'); output(w,'A33')
box(w,29,33,1,4)
w['A35']='Note: weights use gross debt (incl. leases) and market equity. Scenario WACC adjustments (Bear +0.5pp / Bull -0.5pp) are applied on the DCF sheet.'
w.column_dimensions['A'].width=36
for c in 'BCD': w.column_dimensions[c].width=15

# DCF sheet
d=dcf
YRS_H=['FY2024A','FY2025A','FY2026A']; YRS_P=[f'FY{y}E' for y in range(2027,2037)]
P=[CL(c) for c in range(6,16)]  # F..O
H=['C','D','E']
d['A1']='MICROSOFT CORPORATION (MSFT) - DCF MODEL'; d['A1'].font=Font(name=FN,bold=True,size=14)
d['A2']=f'Ticker: MSFT | Valuation date: {ASOF} | Fiscal year end: June 30 | USD millions except per share'
lab(d,'A4','Case selector (1=Bear, 2=Base, 3=Bull)',True); inp(d,'B4',2,'0','Input: 1 = Bear, 2 = Base, 3 = Bull')
lab(d,'A5','Active case'); fml(d,'B5','=CHOOSE(B4,"Bear","Base","Bull")',None,bold=True)
section(d,7,'MARKET DATA & CAPITAL STRUCTURE',1,4)
lab(d,'A8','Current share price ($)'); inp(d,'B8',533.69,PS,'Source: stockanalysis.com / Yahoo Finance, MSFT close 2026-10-14')
lab(d,'A9','Diluted shares (M)'); inp(d,'B9',7453,MIL,SRC_PR+': FY26 diluted weighted average shares 7,453M')
lab(d,'A10','Market cap'); fml(d,'B10','=B8*B9',MIL)
lab(d,'A11','Cash + short-term investments'); fml(d,'B11','=WACC!B25',MIL,GREEN)
lab(d,'A12','Total debt incl. leases'); fml(d,'B12','=WACC!B22',MIL,GREEN)
lab(d,'A13','Net debt / (net cash)'); fml(d,'B13','=B12-B11',MIL)
lab(d,'A14','WACC (active case)',True); fml(d,'B14','=WACC!D33+B59','0.00%',bold=True)
lab(d,'A15','Terminal growth (active case)',True); fml(d,'B15','=B60','0.0%',bold=True)
box(d,7,15,1,4)

ASSUMP={'Bear':dict(g=[.12,.10,.09,.08,.07,.06,.05,.045,.04,.035],m=[.45,.445,.44,.44,.44,.435,.435,.43,.43,.43],
                    da=[.125,.135,.14,.145,.145,.14,.135,.13,.13,.13],cx=[.35,.33,.29,.25,.22,.19,.17,.16,.15,.145],nwc=.08,t=.19,adj=.005,tg=.025),
        'Base':dict(g=[.16,.15,.14,.13,.12,.11,.10,.08,.06,.045],m=[.465,.465,.47,.475,.48,.48,.48,.48,.48,.48],
                    da=[.125,.135,.14,.14,.14,.135,.13,.125,.12,.12],cx=[.33,.30,.26,.22,.19,.17,.155,.145,.135,.13],nwc=.08,t=.19,adj=0.0,tg=.03),
        'Bull':dict(g=[.19,.18,.17,.16,.15,.13,.12,.10,.08,.06],m=[.475,.48,.485,.49,.495,.50,.50,.50,.50,.50],
                    da=[.125,.13,.135,.135,.135,.13,.125,.12,.115,.115],cx=[.31,.27,.23,.19,.16,.145,.135,.125,.12,.115],nwc=.08,t=.19,adj=-.005,tg=.035)}
NOTE={'Bear':'Bear: AI monetisation disappoints / Azure decelerates to high-single digits by FY30; margins compress on depreciation of AI capex; capex stays elevated longer.',
      'Base':'Base: FY27 growth ~16% (FY26 +17.8%, Azure >35%), decelerating to 4.5% by FY36; EBIT margin ~46.5-48% (FY26 46.8%); capex/revenue normalises from 35% (FY26) to 13% as AI build-out matures; D&A rises to 14% then fades.',
      'Bull':'Bull: Copilot/Azure AI sustain high-teens growth through FY31; operating leverage lifts EBIT margin to 50%; capex intensity falls faster.'}
starts={'Bear':18,'Base':29,'Bull':40}
labels=['Revenue growth','EBIT margin','D&A % of revenue','CapEx % of revenue','Change in NWC % of change in revenue','Tax rate']
for case,s in starts.items():
    a=ASSUMP[case]; section(d,s,f'{case.upper()} CASE ASSUMPTIONS',1,15)
    colhdr(d,s+1,['Assumption','','','','']+YRS_P)
    rows=[a['g'],a['m'],a['da'],a['cx'],[a['nwc']]*10,[a['t']]*10]
    for k,(lb,vals) in enumerate(zip(labels,rows)):
        r=s+2+k; lab(d,f'A{r}',lb)
        for j,col in enumerate(P):
            inp(d,f'{col}{r}',vals[j],PCT,(NOTE[case] if j==0 else None) if k<4 else ('Assumption: MSFT working capital use ~10% of revenue growth in FY26 (-4,895 / +50,115); 8% assumed' if k==4 and j==0 else ('Assumption: 19% normalised tax rate (FY26 effective 19.4%)' if k==5 and j==0 else None)))
    lab(d,f'A{s+8}','WACC adjustment vs. CAPM WACC'); inp(d,f'B{s+8}',a['adj'],'+0.0%;-0.0%;0.0%','Assumption: scenario risk adjustment to CAPM WACC (Bear +0.5pp, Bull -0.5pp)')
    lab(d,f'A{s+9}','Terminal growth rate'); inp(d,f'B{s+9}',a['tg'],PCT,'Assumption: Bear 2.5% / Base 3.0% / Bull 3.5%; below 10Y UST (5.19%) and nominal GDP')
    box(d,s,s+9,1,15)
section(d,51,'SELECTED CASE (consolidation - drives projections)',1,15)
colhdr(d,52,['Assumption','','','','']+YRS_P)
for k,lb in enumerate(labels):
    r=53+k; lab(d,f'A{r}',lb)
    for col in P: fml(d,f'{col}{r}',f'=CHOOSE($B$4,{col}{20+k},{col}{31+k},{col}{42+k})',PCT)
lab(d,'A59','WACC adjustment'); fml(d,'B59','=CHOOSE($B$4,B26,B37,B48)','+0.0%;-0.0%;0.0%')
lab(d,'A60','Terminal growth'); fml(d,'B60','=CHOOSE($B$4,B27,B38,B49)',PCT)
box(d,51,60,1,15)

section(d,62,'HISTORICAL & PROJECTED FINANCIALS / UNLEVERED FREE CASH FLOW',1,15)
colhdr(d,63,['USD millions','','']+YRS_H+YRS_P[:0] if False else ['USD millions','']+YRS_H+YRS_P)
# hist inputs
hist={'rev':[245122,281724,331839],'ebit':[109433,128528,155237],'da':[22287,29433,38534],'cx':[44477,64551,115948],'cfo':[118548,136162,182935],'ni':[88136,101832,133749]}
hsrc=['Source: Microsoft FY2024 10-K (filed 2024-07-30); model knowledge, cross-checked to stockanalysis.com',
      'Source: Microsoft FY2025 10-K (filed 2025-07-30) / FY26 Q4 release comparatives',
      SRC_PR]
rowmap={'rev':64,'ebit':66,'da':70,'cx':73,'cfo':77}
for key,r in rowmap.items():
    for j,col in enumerate(H): inp(d,f'{col}{r}',hist[key][j],MIL,hsrc[j]+(' (D&A = "depreciation, amortization, and other")' if key=='da' else '')+(' ; FY24 D&A 22,287 from model knowledge' if key=='da' and j==0 else ''))
lab(d,'A64','Revenue',True); lab(d,'A65','  % growth'); lab(d,'A66','EBIT (operating income)',True); lab(d,'A67','  % margin')
lab(d,'A68','(-) Taxes on EBIT'); lab(d,'A69','NOPAT',True); lab(d,'A70','(+) D&A'); lab(d,'A71','  % of revenue')
lab(d,'A72','EBITDA'); lab(d,'A73','(-) Capital expenditures'); lab(d,'A74','  % of revenue'); lab(d,'A75','(-) Increase in NWC')
lab(d,'A76','Unlevered free cash flow',True); lab(d,'A77','Memo: operating cash flow (reported)'); lab(d,'A78','Memo: reported FCF (CFO - capex)')
for j,col in enumerate(H):
    if j>0: fml(d,f'{col}65',f'={col}64/{H[j-1]}64-1',PCT)
    fml(d,f'{col}67',f'={col}66/{col}64',PCT); fml(d,f'{col}71',f'={col}70/{col}64',PCT)
    fml(d,f'{col}72',f'={col}66+{col}70',MIL); fml(d,f'{col}74',f'={col}73/{col}64',PCT)
    fml(d,f'{col}78',f'={col}77-{col}73',MIL)
prev='E'
for j,col in enumerate(P):
    fml(d,f'{col}64',f'={prev}64*(1+{col}53)',MIL); fml(d,f'{col}65',f'={col}64/{prev}64-1',PCT)
    fml(d,f'{col}66',f'={col}64*{col}54',MIL); fml(d,f'{col}67',f'={col}66/{col}64',PCT)
    fml(d,f'{col}68',f'={col}66*{col}58',MIL); fml(d,f'{col}69',f'={col}66-{col}68',MIL)
    fml(d,f'{col}70',f'={col}64*{col}55',MIL); fml(d,f'{col}71',f'={col}70/{col}64',PCT)
    fml(d,f'{col}72',f'={col}66+{col}70',MIL)
    fml(d,f'{col}73',f'={col}64*{col}56',MIL); fml(d,f'{col}74',f'={col}73/{col}64',PCT)
    fml(d,f'{col}75',f'=({col}64-{prev}64)*{col}57',MIL)
    fml(d,f'{col}76',f'={col}69+{col}70-{col}73-{col}75',MIL,bold=True)
    prev=col
box(d,62,78,1,15)

section(d,80,'DISCOUNTING (mid-year convention)',1,15)
colhdr(d,81,['','','','','']+YRS_P)
lab(d,'A82','Discount period (years)'); inp(d,'F82',0.5,'0.0','Assumption: mid-year convention from start of FY27 (2026-07-01); no stub adjustment for the 3.5 months elapsed to the valuation date')
for j,col in enumerate(P[1:]): fml(d,f'{col}82',f'={P[j]}82+1','0.0')
lab(d,'A83','Discount factor')
lab(d,'A84','PV of unlevered FCF',True)
for col in P:
    fml(d,f'{col}83',f'=1/(1+$B$14)^{col}82','0.0000'); fml(d,f'{col}84',f'={col}76*{col}83',MIL)
box(d,80,84,1,15)

section(d,86,'TERMINAL VALUE & VALUATION SUMMARY',1,4)
items=[('Terminal-year FCF (FY2037)','=O76*(1+B15)',MIL),('Terminal value (Gordon growth)','=B87/(B14-B15)',MIL),
 ('PV of terminal value','=B88*O83',MIL),('Sum of PV of FCF (FY27-FY36)','=SUM(F84:O84)',MIL),
 ('Enterprise value','=B89+B90',MIL),('(-) Total debt incl. leases','=B12',MIL),('(+) Cash + short-term investments','=B11',MIL),
 ('Equity value','=B91-B92+B93',MIL),('Diluted shares (M)','=B9',MIL),('IMPLIED VALUE PER SHARE ($)','=B94/B95',PS),
 ('Current share price ($)','=B8',PS),('Implied upside / (downside)','=B96/B97-1',PCT),('Terminal value as % of EV','=B89/B91',PCT)]
for i,(lb,f,fmt) in enumerate(items):
    r=87+i; lab(d,f'A{r}',lb, lb.isupper() or lb.startswith('Enterprise') or lb.startswith('Equity')); fml(d,f'B{r}',f,fmt)
output(d,'B96'); output(d,'A96'); output(d,'B98')
box(d,86,99,1,4)

section(d,101,'COMPS CROSS-CHECK',1,4)
colhdr(d,102,['Metric','MSFT (DCF)','Peer median','Comment'])
cc=[('EV / LTM EBITDA (FY26A)','=B91/E72','=Comps!J31','DCF EV vs peer LTM median'),
    ('EV / NTM EBITDA (FY27E)','=B91/F72',None,'Forward basis'),
    ('Implied terminal exit multiple (TV / FY36 EBITDA)','=B88/O72','=Comps!J31','Gordon TV expressed as EBITDA multiple'),
    ('P / E (FY26A net income)','=B94/E106','=Comps!K31','Peer P/E distorted by AI-stake gains'),
    ('EV / LTM revenue','=B91/E64','=Comps!I31',''),]
for i,(lb,f1,f2,cm) in enumerate(cc):
    r=103+i; lab(d,f'A{r}',lb); fml(d,f'B{r}',f1,MULT)
    if f2: fml(d,f'C{r}',f2,MULT,GREEN)
    d[f'D{r}']=cm
lab(d,'A108','Memo: FY26A net income'); 
# put NI in E106? use dedicated cell
d['B106']='=B94/B108'
inp(d,'B108',133749,MIL,SRC_PR+': FY26 net income')
lab(d,'A109','Exit-multiple cross-check: value/share at peer median EV/EBITDA on FY36 EBITDA ($)')
fml(d,'B109','=(B90+O72*Comps!J31/(1+B14)^(O82+0.5)-B12+B11)/B9',PS)
lab(d,'A110','Market-implied EV / LTM EBITDA (current)'); fml(d,'B110','=(B10+B13)/E72',MULT)
box(d,101,110,1,4)

section(d,112,'SCENARIO SUMMARY (all cases computed simultaneously on Scenarios sheet)',1,6)
colhdr(d,113,['Case','Value / share ($)','Upside / (downside)','WACC','Terminal growth','TV % of EV'])
for i,cn in enumerate(['Bear','Base','Bull']):
    r=114+i; lab(d,f'A{r}',cn,True); col=['C','D','E'][i]
    fml(d,f'B{r}',f'=Scenarios!{col}30',PS,GREEN); fml(d,f'C{r}',f'=B{r}/$B$8-1',PCT)
    fml(d,f'D{r}',f'=Scenarios!{col}24','0.00%',GREEN); fml(d,f'E{r}',f'=Scenarios!{col}25',PCT,GREEN); fml(d,f'F{r}',f'=Scenarios!{col}31',PCT,GREEN)
output(d,'B115')
lab(d,'A117','Check: active case value = Scenarios sheet (should be 0)'); fml(d,'B117','=ROUND(B96-INDEX(B114:B116,B4),6)','0.000000')
box(d,112,117,1,6)

# ---- Sensitivity 1: WACC x g
section(d,119,'SENSITIVITY 1: VALUE PER SHARE ($) - WACC (rows) x TERMINAL GROWTH (columns)',1,7)
d['A120']='WACC \\ g'; d['A120'].font=BOLD; d['A120'].fill=LIGHT
for j,off in enumerate([-0.01,-0.005,0,0.005,0.01]):
    c=CL(2+j); fml(d,f'{c}120',f'=$B$15{off:+}' if off else '=$B$15',PCT,bold=True); d[f'{c}120'].fill=LIGHT
for i,off in enumerate([-0.01,-0.005,0,0.005,0.01]):
    r=121+i; fml(d,f'A{r}',f'=$B$14{off:+}' if off else '=$B$14','0.00%',bold=True); d[f'A{r}'].fill=LIGHT
    for j in range(5):
        c=CL(2+j)
        fml(d,f'{c}{r}',f'=(SUMPRODUCT($F$76:$O$76/(1+$A{r})^$F$82:$O$82)+$O$76*(1+{c}$120)/($A{r}-{c}$120)/(1+$A{r})^$O$82-$B$12+$B$11)/$B$9',PS)
output(d,'D123'); box(d,119,125,1,7)

# ---- Sensitivity 2: growth shift x margin shift
section(d,127,'SENSITIVITY 2: VALUE PER SHARE ($) - REVENUE GROWTH SHIFT (rows, pp added to every year) x EBIT MARGIN SHIFT (columns, pp added to every year)',1,7)
d['A128']='Growth Δ \\ Margin Δ'; d['A128'].font=BOLD; d['A128'].fill=LIGHT
DEL=[-0.04,-0.02,0,0.02,0.04]; DELG=[-0.03,-0.015,0,0.015,0.03]
for j,v in enumerate(DEL):
    c=CL(2+j); inp(d,f'{c}128',v,'+0.0%;-0.0%;0.0%','Sensitivity axis: pp shift applied to EBIT margin in every projection year' if j==0 else None); d[f'{c}128'].font=Font(name=FN,bold=True,color='0000FF',size=10)
# helper block rows 137-147
lab(d,'A136','Helper: revenue & FCF (at selected-case margin) for each growth shift',True)
colhdr(d,137,['Growth shift','Row type','','','']+YRS_P)
for i,v in enumerate(DELG):
    r=129+i
    inp(d,f'A{r}',v,'+0.0%;-0.0%;0.0%','Sensitivity axis: pp shift applied to revenue growth in every projection year' if i==0 else None); d[f'A{r}'].font=Font(name=FN,bold=True,color='0000FF',size=10)
    rr=138+2*i; fr=rr+1
    fml(d,f'A{rr}',f'=A{r}','+0.0%;-0.0%;0.0%'); d[f'B{rr}']='Revenue'; fml(d,f'A{fr}',f'=A{r}','+0.0%;-0.0%;0.0%'); d[f'B{fr}']='UFCF'
    prevc='E'
    for col in P:
        prevref='$E$64' if prevc=='E' else f'{prevc}{rr}'
        fml(d,f'{col}{rr}',f'={prevref}*(1+{col}$53+$A{rr})',MIL)
        fml(d,f'{col}{fr}',f'={col}{rr}*{col}$54*(1-{col}$58)+{col}{rr}*{col}$55-{col}{rr}*{col}$56-({col}{rr}-{prevref})*{col}$57',MIL)
        prevc=col
    for j in range(5):
        c=CL(2+j)
        fml(d,f'{c}{r}',f'=(SUMPRODUCT(($F${fr}:$O${fr}+$F${rr}:$O${rr}*{c}$128*(1-$F$58:$O$58))*$F$83:$O$83)+($O${fr}+$O${rr}*{c}$128*(1-$O$58))*(1+$B$15)/($B$14-$B$15)*$O$83-$B$12+$B$11)/$B$9',PS)
output(d,'D131'); box(d,127,133,1,7); box(d,136,147,1,15,thin)

# ---- Sensitivity 3: beta x rf
section(d,149,'SENSITIVITY 3: VALUE PER SHARE ($) - BETA (rows) x RISK-FREE RATE (columns); cost of debt moves with Rf',1,7)
d['A150']='Beta \\ Rf'; d['A150'].font=BOLD; d['A150'].fill=LIGHT
for j,off in enumerate([-0.005,-0.0025,0,0.0025,0.005]):
    c=CL(2+j); fml(d,f'{c}150',f'=WACC!$B$4{off:+}' if off else '=WACC!$B$4','0.00%',bold=True); d[f'{c}150'].fill=LIGHT
for i,off in enumerate([-0.2,-0.1,0,0.1,0.2]):
    r=151+i; fml(d,f'A{r}',f'=WACC!$B$6{off:+}' if off else '=WACC!$B$6','0.00',bold=True); d[f'A{r}'].fill=LIGHT
    for j in range(5):
        c=CL(2+j)
        W=f'(({c}$150+$A{r}*WACC!$B$8)*WACC!$B$31+({c}$150+WACC!$B$13)*(1-WACC!$B$15)*WACC!$B$32+$B$59)'
        fml(d,f'{c}{r}',f'=(SUMPRODUCT($F$76:$O$76/(1+{W})^$F$82:$O$82)+$O$76*(1+$B$15)/({W}-$B$15)/(1+{W})^$O$82-$B$12+$B$11)/$B$9',PS)
output(d,'D153'); box(d,149,155,1,7)

d.column_dimensions['A'].width=46
for c in range(2,16): d.column_dimensions[CL(c)].width=13
d.freeze_panes='B4'

# Scenarios sheet: compact full DCF per case
s=scen
s['A1']='SCENARIO ENGINE - full DCF recomputed for each case (reads DCF scenario blocks)'; s['A1'].font=Font(name=FN,bold=True,size=13)
colhdr(s,3,['Case','','Bear','Base','Bull'])
base={'C':20,'D':31,'E':42}; adj={'C':26,'D':37,'E':48}; tgr={'C':27,'D':38,'E':49}
# per case: rows block of years; to keep compact, put yearly rows in separate blocks below
lab(s,'A24','WACC'); lab(s,'A25','Terminal growth'); lab(s,'A26','Sum PV FCF'); lab(s,'A27','PV terminal value'); lab(s,'A28','Enterprise value'); lab(s,'A29','Equity value'); lab(s,'A30','Value per share ($)',True); lab(s,'A31','TV % of EV')
blk={'C':35,'D':45,'E':55}
for col in 'CDE':
    b=blk[col]; br=base[col]
    section(s,b,f'{ {"C":"BEAR","D":"BASE","E":"BULL"}[col]} CASE CASH FLOWS',1,13)
    colhdr(s,b+1,['Item','','FY2026A']+YRS_P,1)
    # columns D..M for years; C = FY26A
    Y=[CL(c) for c in range(4,14)]
    lab(s,f'A{b+2}','Revenue'); fml(s,f'C{b+2}','=DCF!E64',MIL,GREEN)
    lab(s,f'A{b+3}','EBIT'); lab(s,f'A{b+4}','NOPAT'); lab(s,f'A{b+5}','UFCF'); lab(s,f'A{b+6}','Discount factor'); lab(s,f'A{b+7}','PV UFCF'); lab(s,f'A{b+8}','Period')
    pv='C'
    for k,yc in enumerate(Y):
        dc=P[k]  # DCF assumption column
        fml(s,f'{yc}{b+2}',f'={pv}{b+2}*(1+DCF!{dc}{br})',MIL)
        fml(s,f'{yc}{b+3}',f'={yc}{b+2}*DCF!{dc}{br+1}',MIL)
        fml(s,f'{yc}{b+4}',f'={yc}{b+3}*(1-DCF!{dc}{br+5})',MIL)
        fml(s,f'{yc}{b+5}',f'={yc}{b+4}+{yc}{b+2}*DCF!{dc}{br+2}-{yc}{b+2}*DCF!{dc}{br+3}-({yc}{b+2}-{pv}{b+2})*DCF!{dc}{br+4}',MIL)
        fml(s,f'{yc}{b+8}',f'=DCF!{dc}82','0.0',GREEN)
        fml(s,f'{yc}{b+6}',f'=1/(1+{col}$24)^{yc}{b+8}','0.0000')
        fml(s,f'{yc}{b+7}',f'={yc}{b+5}*{yc}{b+6}',MIL)
        pv=yc
    fml(s,f'{col}24',f'=WACC!D33+DCF!B{adj[col]}','0.00%',GREEN)
    fml(s,f'{col}25',f'=DCF!B{tgr[col]}',PCT,GREEN)
    fml(s,f'{col}26',f'=SUM(D{b+7}:M{b+7})',MIL)
    fml(s,f'{col}27',f'=M{b+5}*(1+{col}25)/({col}24-{col}25)*M{b+6}',MIL)
    fml(s,f'{col}28',f'={col}26+{col}27',MIL)
    fml(s,f'{col}29',f'={col}28-DCF!B12+DCF!B11',MIL)
    fml(s,f'{col}30',f'={col}29/DCF!B9',PS); output(s,f'{col}30')
    fml(s,f'{col}31',f'={col}27/{col}28',PCT)
box(s,3,31,1,5)
s.column_dimensions['A'].width=22
for c in range(2,14): s.column_dimensions[CL(c)].width=12
wb.move_sheet('Comps',offset=0)
wb.active=0
wb.save(OUT+'MSFT_DCF_Model_2026-10-14.xlsx')
print('ok')
