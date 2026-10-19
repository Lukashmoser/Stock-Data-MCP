import calendar, json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, ScatterChart, Reference, Series
from openpyxl.workbook.defined_name import DefinedName

# MSFT monthly adj close (Yahoo Finance via stock-data get_price_history, 1mo bars; adj_close of each monthly bar = last trading day close of month, split/dividend adjusted)
msft = {"2021-09":270.57,"2021-10":318.26,"2021-11":317.86,"2021-12":323.37,"2022-01":299.0,"2022-02":287.88,"2022-03":297.05,"2022-04":267.38,"2022-05":262.55,"2022-06":248.02,"2022-07":271.11,"2022-08":253.04,"2022-09":225.39,"2022-10":224.65,"2022-11":247.61,"2022-12":232.74,"2023-01":240.5,"2023-02":242.67,"2023-03":280.49,"2023-04":298.94,"2023-05":320.2,"2023-06":332.04,"2023-07":327.54,"2023-08":320.26,"2023-09":308.52,"2023-10":330.37,"2023-11":370.99,"2023-12":368.18,"2024-01":389.27,"2024-02":405.74,"2024-03":412.69,"2024-04":381.9,"2024-05":407.94,"2024-06":439.21,"2024-07":411.1,"2024-08":410.65,"2024-09":423.61,"2024-10":400.03,"2024-11":417.71,"2024-12":415.78,"2025-01":409.42,"2025-02":392.38,"2025-03":371.03,"2025-04":390.67,"2025-05":455.85,"2025-06":492.54,"2025-07":528.28,"2025-08":502.56,"2025-09":513.72,"2025-10":513.58,"2025-11":488.91,"2025-12":480.57,"2026-01":427.58,"2026-02":391.15,"2026-03":368.68,"2026-04":406.13,"2026-05":449.39,"2026-06":372.32,"2026-07":463.85,"2026-08":507.29,"2026-09":512.9}
# FRED DGS10, monthly end-of-period (last available daily observation in month)
y10 = {"2021-09":1.52,"2021-10":1.55,"2021-11":1.43,"2021-12":1.52,"2022-01":1.79,"2022-02":1.83,"2022-03":2.32,"2022-04":2.89,"2022-05":2.85,"2022-06":2.98,"2022-07":2.67,"2022-08":3.15,"2022-09":3.83,"2022-10":4.1,"2022-11":3.68,"2022-12":3.88,"2023-01":3.52,"2023-02":3.92,"2023-03":3.48,"2023-04":3.44,"2023-05":3.64,"2023-06":3.81,"2023-07":3.97,"2023-08":4.09,"2023-09":4.59,"2023-10":4.88,"2023-11":4.37,"2023-12":3.88,"2024-01":3.99,"2024-02":4.25,"2024-03":4.2,"2024-04":4.69,"2024-05":4.51,"2024-06":4.36,"2024-07":4.09,"2024-08":3.91,"2024-09":3.81,"2024-10":4.28,"2024-11":4.18,"2024-12":4.58,"2025-01":4.58,"2025-02":4.24,"2025-03":4.23,"2025-04":4.17,"2025-05":4.41,"2025-06":4.24,"2025-07":4.37,"2025-08":4.23,"2025-09":4.16,"2025-10":4.11,"2025-11":4.02,"2025-12":4.18,"2026-01":4.26,"2026-02":3.97,"2026-03":4.3,"2026-04":4.4,"2026-05":4.45,"2026-06":4.44,"2026-07":4.75,"2026-08":4.75,"2026-09":5.29}
months = sorted(msft)
assert months == sorted(y10) and len(months) == 61

BLUE = Font(color="0000FF"); BLACK = Font(color="000000"); GREEN = Font(color="008000")
BOLD = Font(bold=True); HDR = PatternFill("solid", fgColor="D9E1F2")
FIRST, LAST = 6, 6 + len(months) - 1   # data rows on Inputs and Calc

wb = Workbook()
inp = wb.active; inp.title = "Inputs"
inp["A1"] = "MSFT vs US 10-Year Treasury - monthly inputs"; inp["A1"].font = Font(bold=True, size=13)
inp["A2"] = "Blue = hardcoded input. Month-end dates are calendar month-ends; values are the last trading-day observation in each month."
inp["A3"] = "Retrieved 2026-10-15. Window: 2021-09-30 to 2026-09-30 (61 month-ends -> 60 monthly changes). Partial Oct-2026 excluded."
hdrs = ["Month-end", "MSFT adj. close (USD/share)", "10Y Treasury yield (%)"]
srcs = ["Source: calendar month-end (constructed)",
        "Source: Yahoo Finance via stock-data get_price_history (MSFT, interval 1mo, adj_close = split & dividend adjusted)",
        "Source: FRED series DGS10 (Board of Governors / U.S. Treasury), monthly end-of-period aggregation, via stock-data get_macro_series; FRED last_updated 2026-10-13"]
for c, (h, s) in enumerate(zip(hdrs, srcs), start=1):
    inp.cell(4, c, h).font = BOLD; inp.cell(4, c).fill = HDR
    inp.cell(5, c, s).font = Font(italic=True, size=8)
    inp.cell(5, c).alignment = Alignment(wrap_text=True, vertical="top")
inp.row_dimensions[5].height = 60
for i, m in enumerate(months):
    r = FIRST + i
    yy, mm = map(int, m.split("-"))
    d = date(yy, mm, calendar.monthrange(yy, mm)[1])
    inp.cell(r, 1, d).number_format = "yyyy-mm-dd"; inp.cell(r, 1).font = BLUE
    inp.cell(r, 2, msft[m]).number_format = "0.00"; inp.cell(r, 2).font = BLUE
    inp.cell(r, 3, y10[m]).number_format = "0.00"; inp.cell(r, 3).font = BLUE
inp.column_dimensions["A"].width = 14; inp.column_dimensions["B"].width = 34; inp.column_dimensions["C"].width = 40

calc = wb.create_sheet("Calc")
calc["A1"] = "Calculations (all cells below row 5 are live formulas)"; calc["A1"].font = Font(bold=True, size=13)
calc["A2"] = "Return = adj close / prior adj close - 1. Yield change in percentage points. Rolling corr = CORREL over trailing 12 monthly pairs."
calc["A3"] = "Full-period correlation:"; calc["A3"].font = BOLD
calc["C3"] = f"=CORREL(D{FIRST+1}:D{LAST},E{FIRST+1}:E{LAST})"; calc["C3"].number_format = "0.000"
calc["D3"] = "Obs:"; calc["E3"] = f"=COUNT(D{FIRST+1}:D{LAST})"
ch = ["Month-end", "MSFT adj. close", "10Y yield (%)", "MSFT monthly return", "Change in 10Y yield (pp)", "12M rolling correlation"]
for c, h in enumerate(ch, start=1):
    calc.cell(5, c, h).font = BOLD; calc.cell(5, c).fill = HDR
for r in range(FIRST, LAST + 1):
    calc.cell(r, 1, f"=Inputs!A{r}").number_format = "yyyy-mm-dd"
    calc.cell(r, 2, f"=Inputs!B{r}").number_format = "0.00"
    calc.cell(r, 3, f"=Inputs!C{r}").number_format = "0.00"
    for c in (1, 2, 3): calc.cell(r, c).font = GREEN
    if r == FIRST:
        calc.cell(r, 4, '=NA()'); calc.cell(r, 5, '=NA()'); calc.cell(r, 6, '=NA()')
    else:
        calc.cell(r, 4, f"=B{r}/B{r-1}-1").number_format = "0.0%"
        calc.cell(r, 5, f"=C{r}-C{r-1}").number_format = "0.00"
        if r - 11 >= FIRST + 1:
            calc.cell(r, 6, f"=CORREL(D{r-11}:D{r},E{r-11}:E{r})").number_format = "0.000"
        else:
            calc.cell(r, 6, "=NA()")
for col, w in zip("ABCDEF", (14, 16, 14, 20, 24, 24)):
    calc.column_dimensions[col].width = w

FIRST_ROLL = FIRST + 12
summ = wb.create_sheet("Summary")
summ["A1"] = "Summary - MSFT monthly returns vs change in 10Y Treasury yield"; summ["A1"].font = Font(bold=True, size=13)
rows = [
    ("First month", f"=Calc!A{FIRST}", "yyyy-mm-dd"),
    ("Last month", f"=Calc!A{LAST}", "yyyy-mm-dd"),
    ("Monthly price rows", f"=COUNT(Inputs!B{FIRST}:B{LAST})", "0"),
    ("Monthly return observations", "=Calc!E3", "0"),
    ("First MSFT adj. close (USD)", f"=Calc!B{FIRST}", "0.00"),
    ("Last MSFT adj. close (USD)", f"=Calc!B{LAST}", "0.00"),
    ("First 10Y yield (%)", f"=Calc!C{FIRST}", "0.00"),
    ("Last 10Y yield (%)", f"=Calc!C{LAST}", "0.00"),
    ("Full-period correlation", "=Calc!C3", "0.000"),
    ("Latest 12M rolling correlation", f"=Calc!F{LAST}", "0.000"),
    ("Min 12M rolling correlation", f"=MIN(Calc!F{FIRST_ROLL}:F{LAST})", "0.000"),
    ("Max 12M rolling correlation", f"=MAX(Calc!F{FIRST_ROLL}:F{LAST})", "0.000"),
]
for i, (k, f, nf) in enumerate(rows, start=3):
    summ.cell(i, 1, k)
    summ.cell(i, 2, f).number_format = nf
summ.column_dimensions["A"].width = 32; summ.column_dimensions["B"].width = 14

lc = LineChart()
lc.title = "MSFT return vs 10Y yield change: 12-month rolling correlation"
lc.y_axis.title = "Correlation"; lc.x_axis.title = "Month-end"
lc.y_axis.scaling.min = -1; lc.y_axis.scaling.max = 1
lc.x_axis.number_format = "yyyy-mm"; lc.height = 9; lc.width = 22
data = Reference(calc, min_col=6, min_row=FIRST_ROLL, max_row=LAST)
s = Series(data, title="12M rolling correlation")
lc.series.append(s)
lc.set_categories(Reference(calc, min_col=1, min_row=FIRST_ROLL, max_row=LAST))
lc.legend.position = "b"
lc.x_axis.delete = False; lc.y_axis.delete = False
summ.add_chart(lc, "D3")

sc = ScatterChart()
sc.title = "Monthly: MSFT return vs change in 10Y yield"
sc.style = 13; sc.height = 9; sc.width = 22
sc.x_axis.title = "Change in 10Y yield (pp)"; sc.y_axis.title = "MSFT monthly return"
xs = Reference(calc, min_col=5, min_row=FIRST + 1, max_row=LAST)
ys = Reference(calc, min_col=4, min_row=FIRST + 1, max_row=LAST)
ss = Series(ys, xs, title="Months")
ss.marker.symbol = "circle"; ss.graphicalProperties.line.noFill = True
sc.series.append(ss); sc.legend = None
sc.x_axis.delete = False; sc.y_axis.delete = False
summ.add_chart(sc, "D22")

chk = wb.create_sheet("Checks")
chk["A1"] = "Checks"; chk["A1"].font = BOLD
checks = [
    ("Price and yield row counts match", f"=COUNT(Inputs!B{FIRST}:B{LAST})=COUNT(Inputs!C{FIRST}:C{LAST})"),
    ("No blank inputs", f"=COUNTBLANK(Inputs!A{FIRST}:C{LAST})=0"),
    ("All prices positive", f"=MIN(Inputs!B{FIRST}:B{LAST})>0"),
    ("Dates strictly increasing", f"=SUMPRODUCT(--(Inputs!A{FIRST+1}:A{LAST}>Inputs!A{FIRST}:A{LAST-1}))={LAST-FIRST}"),
    ("Correlations within [-1,1]", f"=AND(MAX(Calc!F{FIRST_ROLL}:F{LAST})<=1,MIN(Calc!F{FIRST_ROLL}:F{LAST})>=-1,ABS(Calc!C3)<=1)"),
]
for i, (k, f) in enumerate(checks, start=3):
    chk.cell(i, 1, k); chk.cell(i, 2, f)
chk.cell(3 + len(checks), 1, "All checks pass").font = BOLD
chk.cell(3 + len(checks), 2, f"=AND(B3:B{2+len(checks)})").font = BOLD
chk.column_dimensions["A"].width = 36

for name, ref in [("FullPeriodCorr", "Calc!$C$3"), ("LatestRollingCorr", f"Calc!$F${LAST}")]:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

wb.calculation.fullCalcOnLoad = True
wb.save("/home/claude/results/2026-10-13/runs/FA-12/B/out/MSFT_10Y_correlation.xlsx")
print("saved", FIRST, LAST)
