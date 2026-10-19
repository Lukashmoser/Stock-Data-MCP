import json
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

W = "/home/claude/results/2026-10-13/runs/FA-02/A/"
d = json.load(open(W + "_work/calc.json"))

NAVY = RGBColor(0x14, 0x2B, 0x4C); GREY = RGBColor(0x55, 0x5F, 0x6B); LIGHT = RGBColor(0xEE, 0xF1, 0xF5)
ACC = RGBColor(0x1F, 0x6F, 0xB2)

# ---- charts ----
ads = d["ads_pool"]
names = list(ads.keys()); vals = [ads[n] / 1000 for n in names]
growth = {"Alphabet": d["GOOGL_ads"]["growth"], "Meta": None, "ByteDance (est.)": d["BD"]["growth"],
          "Amazon": d["AMZN_ads"]["growth"], "Reddit": d["RDDT"]["growth"]}
meta_ad_prior = 196175 - 87955 + 87955  # placeholder not used
fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
cols = ["#8A96A6", "#1F6FB2", "#8A96A6", "#8A96A6", "#8A96A6"]
b = ax.barh(names[::-1], vals[::-1], color=cols[::-1])
for rect, v in zip(b, vals[::-1]):
    ax.text(rect.get_width() + 4, rect.get_y() + rect.get_height() / 2, f"${v:,.1f}B", va="center", fontsize=9, color="#333")
ax.set_xlabel("TTM advertising revenue, USD billions (ByteDance = 2025 est., 60% ad share)", fontsize=8)
ax.spines[["top", "right"]].set_visible(False); ax.tick_params(labelsize=9); ax.set_xlim(0, 370)
plt.tight_layout(); plt.savefig(W + "_work/ad_scale.png"); plt.close()

fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
labels = ["Reddit", "Meta", "Amazon ads", "ByteDance (2025E)", "Alphabet ads"]
g = [d["RDDT"]["growth"], d["META"]["growth"], d["AMZN_ads"]["growth"], d["BD"]["growth"], d["GOOGL_ads"]["growth"]]
bars = ax.bar(labels, g, color=["#8A96A6", "#1F6FB2", "#8A96A6", "#8A96A6", "#8A96A6"])
for rect, v in zip(bars, g):
    ax.text(rect.get_x() + rect.get_width() / 2, v + 1, f"{v:.1f}%", ha="center", fontsize=9)
ax.set_ylabel("Revenue growth, % (TTM vs prior TTM)", fontsize=8)
ax.spines[["top", "right"]].set_visible(False); ax.tick_params(labelsize=8)
plt.tight_layout(); plt.savefig(W + "_work/growth.png"); plt.close()

# ---- deck ----
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def title(s, t, sub=None):
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9)).text_frame
    tb.word_wrap = True
    p = tb.paragraphs[0]; p.text = t; p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = NAVY
    if sub:
        p2 = tb.add_paragraph(); p2.text = sub; p2.font.size = Pt(13); p2.font.color.rgb = GREY

def foot(s, t):
    tb = s.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.5)).text_frame
    tb.word_wrap = True
    p = tb.paragraphs[0]; p.text = t; p.font.size = Pt(9); p.font.color.rgb = GREY

def bullets(s, items, x, y, w, h, size=14):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bold = it.startswith("**")
        p.text = it.strip("*"); p.font.size = Pt(size if not bold else size + 1); p.font.bold = bold
        p.font.color.rgb = NAVY if bold else RGBColor(0x22, 0x22, 0x22); p.space_after = Pt(6)

def table(s, rows, x, y, w, h, colw=None, size=11):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size); p.font.bold = (i == 0)
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else RGBColor(0x22, 0x22, 0x22)
            c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
    if colw:
        for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)

M, R = d["META"], d["RDDT"]

# 1 Title
s = prs.slides.add_slide(blank)
title(s, "Meta in Digital Advertising: Competitive Landscape",
      "Meta vs Alphabet (Google ads), Amazon (ads), Reddit and ByteDance/TikTok | Data through 2026-06-30 quarter; as of 2026-10-14")
bullets(s, [
    "**Bottom line",
    f"Meta is the #2 ad platform globally (TTM ad revenue USD {M['ad_ttm']:,}M) and the fastest-growing at scale: TTM revenue +{M['growth']:.1f}%, Q2 2026 +{M['q2_growth']:.1f}%, versus Google advertising +{d['GOOGL_ads']['growth']:.1f}%.",
    "AI ranking (Advantage+ run-rate >USD 75B; Generative Recommender) is converting into both more impressions (+14%) and higher prices (+12%) in Q2 2026 - a rare combination.",
    "Main threats: ByteDance/TikTok for attention and commerce video; Amazon for closed-loop, purchase-verified measurement; Google's AI Mode for intent capture. Reddit is a fast-growing niche, not a scale threat.",
    "Key risk is cost, not demand: R&D at 31.4% of TTM revenue and 2026 capex guidance of USD 130-145B compress margins and FCF while the payoff is still proving out.",
], 0.5, 1.6, 12.3, 5)
foot(s, "Sources: company 10-Q/8-K filings (SEC EDGAR), press releases, Reuters via whbl.com, Sacra, eMarketer. TTM = Q3 2025 to Q2 2026.")

# 2 Market map
s = prs.slides.add_slide(blank)
title(s, "Market scale: five ad platforms, three business models", "TTM advertising revenue (USD billions)")
s.shapes.add_picture(W + "_work/ad_scale.png", Inches(0.4), Inches(1.4), width=Inches(7.2))
bullets(s, [
    "**Model archetypes",
    "Intent (Alphabet): search query = explicit demand; highest-margin format, now being re-plumbed for AI answers.",
    "Discovery / social graph (Meta, TikTok, Reddit): demand creation in feeds; AI ranking replaces user-declared targeting.",
    "Commerce / closed loop (Amazon, TikTok Shop): ad and purchase on the same platform - best attribution.",
    "Meta is ~0.7x Alphabet's ad revenue and ~2.9x Amazon's; Reddit is ~1% of Meta.",
], 7.8, 1.5, 5.1, 5, size=13)
foot(s, "Alphabet = Google advertising line (Search & other + YouTube ads + Network). Amazon = 'Advertising services' segment line. Meta = advertising revenue. Reddit = total revenue (~all ads). ByteDance = Sacra 2025 revenue est. USD 186B x ~60% ad share.")

# 3 Positioning
s = prs.slides.add_slide(blank)
title(s, "Competitive positioning")
rows = [["", "Meta", "Alphabet (Google ads)", "Amazon ads", "ByteDance / TikTok", "Reddit"],
        ["Core asset", "3.6B daily users across FB, IG, WhatsApp, Messenger, Threads", "Search intent + YouTube; Network", "Shopper intent + purchase data; Prime Video", "Interest graph, short video, live commerce (Douyin/TikTok Shop)", "Topic communities; high-intent discussion; search referral traffic"],
        ["Funnel position", "Full-funnel, strongest mid/lower (performance, app, CTW)", "Lower funnel (search) + upper (YouTube CTV)", "Lowest funnel (sponsored products) + CTV/DSP", "Upper/mid funnel moving to lower via Shop", "Mid funnel; consideration, research"],
        ["Advertiser base", "Millions of SMBs + large brands; global", "Broadest; agencies and SMB", "Endemic brands/sellers; growing non-endemic", "D2C, apps, e-commerce sellers, brands", "Active advertisers +70% y/y; SMB push"],
        ["Moat", "Scale of first-party engagement data + AI ranking + compute", "Query intent, Android/Chrome distribution", "Closed-loop sales attribution", "Engagement intensity; recommendation engine", "Unique UGC corpus (also licensed to AI labs)"],
        ["Key vulnerability", "Capex/opex intensity; regulation (EU, youth safety, legal charges)", "AI answers reduce clicks; antitrust remedies", "Ad load vs shopper experience", "US JV governance (ByteDance 19.9%); geopolitics", "Dependence on Google search traffic; scale"]]
table(s, rows, 0.4, 1.3, 12.5, 5.2, [1.6, 2.3, 2.2, 2.1, 2.2, 2.1], size=10)
foot(s, "Sources: Meta Q2 2026 call (DAP 3.6B); Reddit Q2 2026 call; eMarketer 2025-12-19 on TikTok US JV structure; analyst synthesis.")

# 4 Ad formats & audiences
s = prs.slides.add_slide(blank)
title(s, "Ad formats and audiences")
rows = [["", "Primary formats", "Audience / reach", "Recent product moves"],
        ["Meta", "Feed, Stories, Reels, click-to-message (WhatsApp/Messenger), Threads ads, Advantage+ automated campaigns", "3.6B DAP (Q2 2026); broad global demographics", "Generative Recommender; Advantage+ >USD 75B run-rate; WhatsApp/Threads monetization ramp"],
        ["Alphabet", "Search text/shopping, Performance Max, AI Max, YouTube (Shorts, CTV), Display Network", "Search near-universal; AI Mode >1B MAU (company claim)", "AI Max GA (Apr 2026); ads in AI Overviews/AI Mode"],
        ["Amazon", "Sponsored Products/Brands/Display, Amazon DSP, Prime Video ads, Twitch", "Prime members; shoppers at point of purchase", "Ads Agent expanded to 11 more countries; claims 6% lower CPA"],
        ["ByteDance", "In-feed video, TopView, Spark Ads, Smart+ (automated), Live/Shop ads", "TikTok >1B users ex-China; Douyin in China; young skew", "US JV closed 2026-01-22; Symphony gen-AI creative"],
        ["Reddit", "Promoted posts, conversation ads, Dynamic Product Ads, shopping listing ads (alpha)", "DAUq 130.3M, WAUq 514.6M (Q2 2026)", "Reddit Max automation; DPA revenue >2x y/y; dual attribution beta"]]
table(s, rows, 0.4, 1.3, 12.5, 5.0, [1.4, 4.3, 3.2, 3.6], size=10)
foot(s, "Sources: company calls and releases Q2 2026; TikTok format names from model knowledge.")

# 5 Financials
s = prs.slides.add_slide(blank)
title(s, "Financial comparison", "USD millions; TTM = Q3 2025-Q2 2026 unless noted")
rows = [["Metric", "Meta", "Reddit", "ByteDance", "Alphabet (Google ads)", "Amazon (ads services)"],
        ["Revenue / ad revenue", f"{M['ttm_rev']:,}", f"{R['ttm_rev']:,.1f}", "~186,000 (FY2025E)", f"{d['GOOGL_ads']['ttm']:,}", f"~{d['AMZN_ads']['ttm']:,}"],
        ["Growth % (TTM y/y)", f"{M['growth']:.1f}%", f"{R['growth']:.1f}%", "~20.0% (FY2025E)", f"{d['GOOGL_ads']['growth']:.1f}%", f"{d['AMZN_ads']['growth']:.1f}%"],
        ["Latest quarter growth (Q2 2026)", f"{M['q2_growth']:.1f}%", f"{R['q2_growth']:.1f}%", "n/d", f"{d['GOOGL_ads']['q2_growth']:.1f}%", f"{d['AMZN_ads']['q2_growth']:.1f}%"],
        ["Gross margin %", f"{M['gm']:.1f}%", f"{R['gm']:.1f}%", "n/d", "n/d (segment)", "n/d (segment)"],
        ["R&D % revenue", f"{M['rnd']:.1f}%", f"{R['rnd']:.1f}%", "n/d", "n/d (segment)", "n/d (segment)"],
        ["S&M % revenue", f"{M['sm']:.1f}%", f"{R['sm']:.1f}%", "n/d", "n/d (segment)", "n/d (segment)"],
        ["SG&A % revenue", f"{M['sga']:.1f}%", f"{R['sga']:.1f}%", "n/d", "n/d", "n/d"],
        ["Market cap / valuation", f"~{M['mcap']:,.0f} (2026-09-25)", f"{R['mcap']:,} (2026-10-12)", "550,000 (Feb 2026 deal)", "Parent-level only", "Parent-level only"]]
table(s, rows, 0.4, 1.4, 12.5, 4.6, [2.6, 2.1, 1.9, 2.0, 2.0, 1.9], size=11)
foot(s, "Meta, Reddit: 10-K FY2025 / 10-Q Q2 2026, TTM derived. Alphabet: Q3 2025 per storyboard18, Q4 2025 release, Q2 2026 10-Q. Amazon: rounded quarterly figures from press coverage. ByteDance: Sacra estimate; Reuters 2026-02-25 valuation. Market caps: Google Finance.")

# 6 Growth chart + commentary
s = prs.slides.add_slide(blank)
title(s, "Growth and efficiency: Meta grows like a challenger at incumbent scale")
s.shapes.add_picture(W + "_work/growth.png", Inches(0.4), Inches(1.4), width=Inches(7.2))
bullets(s, [
    "**Read-across",
    f"Meta added ~USD {(M['ttm_rev']-M['prior'])/1000:,.0f}B of TTM revenue y/y - more than Google advertising (~USD {(d['GOOGL_ads']['ttm']-d['GOOGL_ads']['prior'])/1000:,.0f}B) and Amazon ads (~USD {(d['AMZN_ads']['ttm']-d['AMZN_ads']['prior'])/1000:,.0f}B).",
    "Cost structure diverges: Meta spends 31.4% of revenue on R&D but only 5.5% on S&M - the product sells itself through self-serve auctions. Reddit spends 23.0% on S&M to build its advertiser base.",
    "Reddit has the best gross margin (91.4%) but is ~1/80th of Meta's revenue.",
    "Meta's Q2 2026 operating income fell y/y (USD 18,775M vs 20,441M) as R&D +67% and legal/severance charges hit.",
], 7.8, 1.5, 5.1, 5, size=13)
foot(s, "Growth basis differs: Meta and Reddit total revenue; Alphabet/Amazon ad lines; ByteDance FY2025 vs FY2024 estimate (Sacra).")

# 7 AI
s = prs.slides.add_slide(blank)
title(s, "AI's effect on ad targeting and measurement")
rows = [["", "Targeting / ranking", "Creative", "Measurement", "Net effect"],
        ["Meta", "Generative Recommender (LLM-based); ranking gains: FB ad clicks +8.3%, conversions +15.7%", "Advantage+ creative gen-AI; >USD 75B Advantage+ run-rate", "Modeled conversions, Conversions API; post-ATT signal loss largely rebuilt", "Strongest beneficiary: price/ad +12% and impressions +14% in Q2 2026"],
        ["Alphabet", "AI Max broad-match + query expansion; PMax", "Asset generation in PMax", "Modeled attribution; Google Analytics", "Defensive: AI Mode may cut clicks, offset by higher-value queries; Search & other +16.8% in Q2 2026"],
        ["Amazon", "Shopper signals + Ads Agent", "Gen-AI image/video tools", "Closed-loop sales attribution (AMC)", "Structural measurement advantage; ad growth 26.2% in Q2 2026"],
        ["ByteDance", "Interest-graph recommender (best-in-class)", "Symphony gen-AI avatars/video", "Shop closed loop; US JV retrains algorithm", "Strong, but US algorithm retraining is a wildcard"],
        ["Reddit", "Reddit Max automation; contextual from conversations", "Limited", "Dual attribution beta; third-party ROAS studies", "Gains from automation (Max revenue +150% q/q); risk from AI search reducing referral traffic"]]
table(s, rows, 0.4, 1.3, 12.5, 4.6, [1.3, 3.0, 2.4, 2.8, 3.0], size=10)
bullets(s, ["Implication: AI shifts advantage from declared targeting to whoever owns the most first-party behaviour data plus the compute to model it - favouring Meta, Google and ByteDance; measurement edge shifts to closed-loop players (Amazon, TikTok Shop)."], 0.4, 6.05, 12.5, 0.8, size=12)
foot(s, "Sources: Meta Q2 2026 call (nasdaq.com); digitalapplied.com on Alphabet Q2 2026; ecomcrew.com on Amazon Q2 2026; Reddit Q2 2026 call. Some product names from model knowledge.")

# 8 Strategic synthesis
s = prs.slides.add_slide(blank)
title(s, "Strategic synthesis")
bullets(s, [
    "**Where Meta wins",
    "Scale + AI flywheel: more engagement data -> better ranking -> higher ROAS -> higher prices without losing volume.",
    "SMB self-serve depth and automation (Advantage+) lowers the cost of acquiring advertisers; S&M only 5.5% of revenue.",
    "Untapped surfaces: WhatsApp business messaging, Threads, Meta AI assistant monetization.",
    "**Where Meta is exposed",
    "Attention: TikTok remains the benchmark for short video; US JV closed 2026-01-22 removes the ban overhang that briefly helped Reels.",
    "Measurement: Amazon and TikTok Shop prove sales directly; Meta relies on modeled conversions.",
    "Cost: 2026 capex guidance USD 130-145B and expense guidance USD 165-169B - margin compression is the main debate.",
    "Regulation/legal: Q2 2026 USD 2.4B legal charge; EU less-personalised-ads remedies.",
    "**Watch-list (next 2-4 quarters)",
    "Q3 2026 revenue vs USD 61-64B guide; price-per-ad vs impressions mix; Advantage+ share; TikTok US post-retraining performance; Google AI Mode ad load.",
], 0.5, 1.3, 12.3, 5.5, size=13)
foot(s, "Analyst synthesis based on the sources cited on prior slides.")

# 9 Sources
s = prs.slides.add_slide(blank)
title(s, "Sources and notes")
bullets(s, [
    "Meta 8-K Q4/FY2025 release (2026-01-28); Meta 10-Q Q2 2026 (sec.gov/.../meta-20260630.htm); Q2 2026 call highlights (nasdaq.com).",
    "Reddit Q4/FY2025 release (businesswire, 2026-02-05); Reddit Q2 2026 8-K release (2026-07-30); Q2 2026 call (webull.com).",
    "Alphabet 10-Q Q2 2026; Q4 2025 earnings release; Q3 2025 figures via storyboard18.com.",
    "Amazon advertising: ppc.land and ecomcrew.com (Q2 2026); storyboard18.com (Q3 2025). Q3 2024-Q2 2025 comparatives from model knowledge.",
    "ByteDance: Sacra (revenue estimates), Reuters via whbl.com 2026-02-25 (USD 550B valuation), augment.market 2026-06-23 (USD 600B secondary, gray-market quotes), eMarketer 2025-12-19 (TikTok US JV).",
    "Prices: Google Finance (META close 2026-09-25 snapshot; RDDT close 2026-10-12).",
    "Reddit H1 2024 revenue (USD 524.2M) and Meta H1 2024 revenue (USD 75,526M) used for prior-TTM growth are from model knowledge.",
], 0.5, 1.3, 12.3, 5.5, size=12)

prs.save(W + "Meta_Digital_Ads_Competitive_Analysis.pptx")
print("ok")
