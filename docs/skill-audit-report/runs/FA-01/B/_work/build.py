import openpyxl
from openpyxl.styles import Font
from openpyxl.comments import Comment

SRC = '/home/claude/results/2026-10-13/runs/FA-01/B/3S_template.xlsx'
OUT = '/home/claude/results/2026-10-13/runs/FA-01/B/MSFT_3S_model.xlsx'
wb = openpyxl.load_workbook(SRC)

BLUE, BLACK, GREEN = '0000FF', '000000', '008000'
H = ['B', 'C', 'D', 'E']          # FY2023A..FY2026A
P = ['F', 'G', 'H', 'I', 'J']     # FY2027E..FY2031E
ALL = H + P
LABELS = ['FY2023A', 'FY2024A', 'FY2025A', 'FY2026A', 'FY2027E', 'FY2028E', 'FY2029E', 'FY2030E', 'FY2031E']
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%'
EPS = '0.00'


def prev(col):
    return chr(ord(col) - 1)


def put(ws, ref, val, kind=None, fmt=NUM):
    c = ws[ref]
    c.value = val
    if kind is None:
        if isinstance(val, str) and val.startswith('='):
            kind = 'link' if '!' in val else 'f'
        else:
            kind = 'in'
    color = {'in': BLUE, 'f': BLACK, 'link': GREEN}[kind]
    c.font = Font(name=c.font.name, size=c.font.sz, bold=c.font.b, color=color)
    if fmt:
        c.number_format = fmt


def row(ws, r, vals, cols, fmt=NUM):
    for col, v in zip(cols, vals):
        if v is None:
            continue
        put(ws, f'{col}{r}', v, fmt=fmt)


# relabel headers
for ws in wb.worksheets[1:]:
    for col, lab in zip(ALL, LABELS):
        if ws[f'{col}4'].value:
            ws[f'{col}4'].value = lab
    if ws['A2'].value and '[currency]' in str(ws['A2'].value):
        ws['A2'].value = 'Units: USD millions unless stated (per-share in USD; shares in millions)'

# ---------------- Cover ----------------
cv = wb['Cover']
cv['B3'] = 'Microsoft Corporation'
cv['B4'] = 'MSFT'
cv['B5'] = 'June 30'
cv['B6'] = 'USD millions (per-share values in USD)'
cv['B7'] = '2026-10-14'
cv['B8'] = 'Claude (financial-analysis:3-statement-model)'
cv['B9'] = ('Form 10-K FY2026 (filed 2026-07-29), 10-K FY2025 (filed 2025-07-30), 10-K FY2024 (filed 2024-07-30) '
            'via SEC EDGAR XBRL R-pages; stock-data get_financials cross-check')
cv['A15'] = 'Columns: FY2023A to FY2026A = historical actuals (FYE June 30); FY2027E to FY2031E = projections'
cv['A18'] = 'Mapping notes'
notes = [
    'IS: D&A (depreciation + intangible amortization per 10-K notes) is shown on its own line and removed from cost of revenue, so EBIT = reported operating income.',
    'IS: SG&A = Sales and marketing + General and administrative. Other income/(expense) is split into interest income, interest expense and the residual.',
    'BS: Other current assets as reported; Other long-term assets = operating lease ROU + equity & other investments + other long-term assets.',
    'BS: Accrued & other current liabilities = accrued compensation + short-term income taxes + short-term unearned revenue + other current liabilities.',
    'BS: Deferred taxes & other LT liabilities = LT income taxes + LT unearned revenue + deferred income taxes + operating lease liabilities + other LT liabilities.',
    'BS: Microsoft retires repurchased shares (no treasury stock); projected buybacks are charged to retained earnings.',
    'CFS: Reported "Depreciation, amortization, and other" minus note D&A is moved to Other operating activities, with net investment gains/losses and LT asset/liability changes. CFO totals match the 10-K.',
    'CFS FY2024A-FY2026A use the FY2026 10-K presentation; FY2023A uses the FY2025 10-K presentation (the FY2026 10-K reclassified equity-method losses within operating adjustments; CFO total unaffected).',
]
for i, n in enumerate(notes):
    cv[f'A{19 + i}'] = n

# ---------------- Historical data ----------------
hist = {
    'rev': [211915, 245122, 281724, 331839],
    'cogs_rep': [65863, 74114, 87831, 106374],
    'rd': [27195, 29510, 32488, 35562],
    'sm': [22759, 24456, 25654, 26710],
    'ga': [7575, 7609, 7223, 7956],
    'oi': [88523, 109433, 128528, 155237],
    'oie': [788, -1646, -4901, 10697],
    'int_inc': [2994, 3157, 2647, 3301],
    'int_exp': [1968, 2935, 2385, 3051],
    'tax': [16950, 19651, 21795, 32185],
    'ni': [72361, 88136, 101832, 133749],
    'sh_b': [7446, 7431, 7433, 7429],
    'sh_d': [7472, 7469, 7465, 7453],
    'eps_b': [9.72, 11.86, 13.70, 18.00],
    'eps_d': [9.68, 11.80, 13.64, 17.95],
    'dps': [2.72, 3.00, 3.32, 3.64],
    'dep': [11000, 15200, 22000, 34300],
    'amort': [2500, 4800, 6000, 4700],
    # BS
    'cash': [34704, 18315, 30242, 20935],
    'sti': [76558, 57228, 64323, 55908],
    'ar': [48688, 56924, 69905, 80876],
    'inv': [2500, 1246, 938, 1397],
    'oca': [21807, 26021, 25723, 48594],
    'ppe': [95641, 135591, 204966, 313076],
    'gw': [67886, 119220, 119509, 119651],
    'intang': [9366, 27597, 22604, 18609],
    'olta': [14346 + 9879 + 30601, 18961 + 14600 + 36460, 24823 + 15405 + 40565, 24177 + 36348 + 38805],
    'ap': [18095, 21996, 27724, 42416],
    'accr': [11009 + 4152 + 50901 + 14745, 12564 + 5017 + 57582 + 19185, 13709 + 7211 + 64555 + 25020,
             14945 + 2534 + 72965 + 26738],
    'std': [0 + 5247, 6693 + 2249, 0 + 2999, 9227],
    'ltd': [41990, 42688, 40152, 31067],
    'oltl': [25560 + 2912 + 433 + 12728 + 17981, 27931 + 2602 + 2618 + 15497 + 27064,
             25986 + 2710 + 2835 + 17437 + 45186, 28647 + 2747 + 3054 + 16532 + 65117],
    'apic': [93718, 100923, 109095, 117406],
    're': [118848, 173144, 237731, 328265],
    'aoci': [-6343, -5590, -3347, -3284],
    # CFS
    'cfs_da_other': [13861, 20958, 29433, 38534],
    'sbc': [9611, 10734, 11974, 12405],
    'gains': [196, 1634, 5329, -11047],
    'dtax': [-6059, -4738, -7056, 14189],
    'c_ar': [-4087, -7191, -10581, -12737],
    'c_inv': [1242, 1284, 309, -461],
    'c_oca': [-1991, -1648, -3044, -2627],
    'c_olta': [-2833, -6817, -2950, -3964],
    'c_ap': [-2721, 3545, 569, 5268],
    'c_unearned': [5535, 5348, 5438, 9361],
    'c_itax': [-358, 1687, -38, -1875],
    'c_ocl': [2272, 4867, 5922, 6847],
    'c_oltl': [553, 749, -975, -4707],
    'cfo': [87582, 118548, 136162, 182935],
    'capex': [-28107, -44477, -64551, -115948],
    'acq': [-1670, -69132, -5978, -1743],
    'inv_net': [-37651 + 33510 + 14354, -17732 + 24775 + 10894, -29775 + 16079 + 9309, -58351 + 34605 + 21798],
    'oth_inv': [-3116, -1298, 2317, -19861],
    'debt_iss': [0, 24395 + 5250, 0, 0],
    'debt_rep': [-2750, -29070, -3216 - 5746, -3000],
    'divs': [-19800, -21771, -24082, -26445],
    'buyb': [-22245, -17254, -18420, -22271],
    'stk_iss': [1866, 2002, 2056, 2009],
    'oth_fin': [-1006, -1309, -2291, -2839],
    'fx': [-194, -210, 63, -196],
}

# ---------------- Assumptions ----------------
A = wb['Assumptions']
A['A2'] = 'Units: USD millions unless stated. Historical columns show implied ratios (formulas); projection columns are inputs (blue).'
proj_in = {
    7: [0.160, 0.140, 0.120, 0.110, 0.100],
    8: [0.205, 0.205, 0.205, 0.205, 0.205],
    9: [0.102, 0.100, 0.098, 0.097, 0.096],
    10: [0.105, 0.105, 0.105, 0.105, 0.105],
    11: [0.0, 0.0, 0.0, 0.0, 0.0],
    12: [0.070, 0.070, 0.070, 0.070, 0.070],
    13: [0.035, 0.035, 0.035, 0.035, 0.035],
    14: [0.190, 0.190, 0.190, 0.190, 0.190],
    15: [7440, 7425, 7410, 7395, 7380],
    18: [89, 89, 89, 89, 89],
    19: [7.5, 7.5, 7.5, 7.5, 7.5],
    20: [220, 210, 200, 190, 180],
    21: [0.146, 0.146, 0.146, 0.146, 0.146],
    22: [0.353, 0.353, 0.353, 0.353, 0.353],
    25: [0.350, 0.330, 0.300, 0.280, 0.260],
    26: [0.170, 0.170, 0.170, 0.170, 0.170],
    27: [3097, 2141, 1944, 1477, 1128],
    30: [-9227, 0, -2001, 0, -500],
    31: [0.200, 0.200, 0.200, 0.200, 0.200],
    32: [24000, 25000, 26000, 27000, 28000],
    33: [15000, 15000, 15000, 15000, 15000],
}
pct_rows = {7, 8, 9, 10, 11, 12, 13, 14, 21, 22, 25, 26, 31}
for r, vals in proj_in.items():
    fmt = PCT if r in pct_rows else ('0.0' if r in (18, 19, 20) else NUM)
    for col, v in zip(P, vals):
        put(A, f'{col}{r}', v, 'in', fmt)

# extra assumption rows
A['A35'] = 'Other assumptions'
A['A35'].font = Font(bold=True)
extra = {
    36: ('Stock-based compensation (% of revenue)', [0.037] * 5, PCT),
    37: ('Proceeds from stock issuance', [2000] * 5, NUM),
    38: ('Other income / (expense), net excl. interest', [0] * 5, NUM),
    39: ('Current portion of debt at year end (next-year maturities)', [0, 2001, 0, 500, 0], NUM),
}
for r, (lab, vals, fmt) in extra.items():
    A[f'A{r}'] = lab
    for col, v in zip(P, vals):
        put(A, f'{col}{r}', v, 'in', fmt)

# historical implied ratios on Assumptions (formulas / links)
for i, c in enumerate(H):
    p = prev(c)
    if i > 0:
        put(A, f'{c}7', f'=IS!{c}6/IS!{p}6-1', fmt=PCT)
    put(A, f'{c}8', f'=IS!{c}7/IS!{c}6', fmt=PCT)
    put(A, f'{c}9', f'=IS!{c}9/IS!{c}6', fmt=PCT)
    put(A, f'{c}10', f'=IS!{c}10/IS!{c}6', fmt=PCT)
    put(A, f'{c}11', f'=IS!{c}11/IS!{c}6', fmt=PCT)
    if i > 0:
        put(A, f'{c}12', f'=IS!{c}15/AVERAGE(BS!{p}22+BS!{p}24,BS!{c}22+BS!{c}24)', fmt=PCT)
        put(A, f'{c}13', f'=IS!{c}16/(BS!{p}7+BS!{p}8)', fmt=PCT)
    put(A, f'{c}14', f'=IS!{c}19/IS!{c}18', fmt=PCT)
    put(A, f'{c}15', f'=IS!{c}24')
    put(A, f'{c}18', f'=BS!{c}9/IS!{c}6*365', fmt='0.0')
    put(A, f'{c}19', f'=BS!{c}10/IS!{c}7*365', fmt='0.0')
    put(A, f'{c}20', f'=BS!{c}20/IS!{c}7*365', fmt='0.0')
    put(A, f'{c}21', f'=BS!{c}11/IS!{c}6', fmt=PCT)
    put(A, f'{c}22', f'=BS!{c}21/IS!{c}6', fmt=PCT)
    put(A, f'{c}25', f'=-CFS!{c}20/IS!{c}6', fmt=PCT)
    if i > 0:
        put(A, f'{c}26', f'=Schedules!{c}9/Schedules!{c}7', fmt=PCT)
    put(A, f'{c}27', f'=Schedules!{c}12')
    put(A, f'{c}30', f'=CFS!{c}27+CFS!{c}28')
    put(A, f'{c}31', f'=-CFS!{c}29/IS!{c}20', fmt=PCT)
    put(A, f'{c}32', f'=-CFS!{c}30')
    put(A, f'{c}36', f'=CFS!{c}9/IS!{c}6', fmt=PCT)
    put(A, f'{c}37', f'=CFS!{c}31')
    put(A, f'{c}38', f'=IS!{c}17')
    put(A, f'{c}39', f'=BS!{c}22')
A['B12'].comment = Comment('Interest rate on average total debt; interest expense includes finance lease interest, so the implied rate exceeds bond coupons.', 'Claude')
A['B13'].comment = Comment('Interest & dividend income / prior-year (cash + short-term investments). Beginning balance used to avoid circularity.', 'Claude')
A['B20'].comment = Comment('DPO is high because accounts payable includes unpaid capex ($26.7bn at FY2026). Tapered toward 180 days.', 'Claude')

# ---------------- Income statement ----------------
I = wb['IS']
for i, c in enumerate(H):
    put(I, f'{c}6', hist['rev'][i])
    put(I, f'{c}7', f'=IS!{c}38-Schedules!{c}14', 'link')
    put(I, f'{c}9', hist['sm'][i] + hist['ga'][i])
    put(I, f'{c}10', hist['rd'][i])
    put(I, f'{c}11', 0)
    put(I, f'{c}15', hist['int_exp'][i])
    put(I, f'{c}16', hist['int_inc'][i])
    put(I, f'{c}17', hist['oie'][i] - hist['int_inc'][i] + hist['int_exp'][i])
    put(I, f'{c}19', hist['tax'][i])
    put(I, f'{c}23', hist['sh_b'][i])
    put(I, f'{c}24', hist['sh_d'][i])
    put(I, f'{c}27', hist['dps'][i], fmt=EPS)
    put(I, f'{c}38', hist['cogs_rep'][i])
    put(I, f'{c}40', hist['oi'][i])
    put(I, f'{c}41', hist['ni'][i])
    put(I, f'{c}42', hist['eps_d'][i], fmt=EPS)
for c in P:
    p = prev(c)
    put(I, f'{c}6', f'={p}6*(1+Assumptions!{c}7)', 'link')
    put(I, f'{c}7', f'={c}6*Assumptions!{c}8', 'link')
    put(I, f'{c}9', f'={c}6*Assumptions!{c}9', 'link')
    put(I, f'{c}10', f'={c}6*Assumptions!{c}10', 'link')
    put(I, f'{c}11', f'={c}6*Assumptions!{c}11', 'link')
    put(I, f'{c}15', f'=Schedules!{c}35')
    put(I, f'{c}16', f'=Schedules!{c}38')
    put(I, f'{c}17', f'=Assumptions!{c}38')
    put(I, f'{c}19', f'={c}18*Assumptions!{c}14', 'link')
    put(I, f'{c}24', f'=Assumptions!{c}15')
    put(I, f'{c}23', f'={c}24-({p}24-{p}23)')
    put(I, f'{c}27', f'=-CFS!{c}29/{c}23', fmt=EPS)
for c in ALL:
    p = prev(c)
    put(I, f'{c}8', f'={c}6-{c}7')
    put(I, f'{c}12', f'={c}8-{c}9-{c}10-{c}11')
    put(I, f'{c}13', f'=Schedules!{c}14')
    put(I, f'{c}14', f'={c}12-{c}13')
    put(I, f'{c}18', f'={c}14-{c}15+{c}16+{c}17')
    put(I, f'{c}20', f'={c}18-{c}19')
    put(I, f'{c}25', f'={c}20/{c}23', fmt=EPS)
    put(I, f'{c}26', f'={c}20/{c}24', fmt=EPS)
    if c != 'B':
        put(I, f'{c}30', f'={c}6/{p}6-1', fmt=PCT)
    put(I, f'{c}31', f'={c}8/{c}6', fmt=PCT)
    put(I, f'{c}32', f'={c}12/{c}6', fmt=PCT)
    put(I, f'{c}33', f'={c}14/{c}6', fmt=PCT)
    put(I, f'{c}34', f'={c}20/{c}6', fmt=PCT)
    put(I, f'{c}35', f'={c}19/{c}18', fmt=PCT)
I['A37'] = 'Memo: reported figures (10-K) and reconciliation'
I['A37'].font = Font(bold=True)
I['A38'] = 'Cost of revenue as reported (incl. D&A)'
I['A39'] = 'Gross margin as reported'
I['A40'] = 'Operating income as reported'
I['A41'] = 'Net income as reported'
I['A42'] = 'Diluted EPS as reported'
I['A43'] = 'Check: model EBIT less reported operating income'
I['A44'] = 'Check: model net income less reported net income'
for c in H:
    put(I, f'{c}39', f'={c}6-{c}38')
    put(I, f'{c}43', f'={c}14-{c}40')
    put(I, f'{c}44', f'={c}20-{c}41')
I['A7'].comment = Comment('Cost of revenue excluding D&A (D&A shown separately on row 13). Historical = reported cost of revenue less note D&A.', 'Claude')
I['B27'].comment = Comment('Historical: cash dividends declared per share (10-K equity statement). Projected: dividends paid / basic shares.', 'Claude')

# ---------------- Balance sheet ----------------
B = wb['BS']
bs_map = {7: 'cash', 8: 'sti', 9: 'ar', 10: 'inv', 11: 'oca', 13: 'ppe', 14: 'gw', 15: 'intang', 16: 'olta',
          20: 'ap', 21: 'accr', 22: 'std', 24: 'ltd', 25: 'oltl', 29: 'apic', 30: 're', 31: 'aoci'}
for i, c in enumerate(H):
    for r, k in bs_map.items():
        put(B, f'{c}{r}', hist[k][i])
    put(B, f'{c}32', 0)
for c in P:
    p = prev(c)
    put(B, f'{c}7', f'=CFS!{c}38')
    put(B, f'{c}8', f'={p}8-CFS!{c}22', 'link')
    put(B, f'{c}9', f'=Schedules!{c}17')
    put(B, f'{c}10', f'=Schedules!{c}18')
    put(B, f'{c}11', f'=Schedules!{c}19')
    put(B, f'{c}13', f'=Schedules!{c}10')
    put(B, f'{c}14', f'={p}14')
    put(B, f'{c}15', f'=Schedules!{c}13')
    put(B, f'{c}16', f'={p}16')
    put(B, f'{c}20', f'=Schedules!{c}20')
    put(B, f'{c}21', f'=Schedules!{c}21')
    put(B, f'{c}22', f'=MIN(Assumptions!{c}39,Schedules!{c}32)')
    put(B, f'{c}24', f'=Schedules!{c}32-{c}22', 'link')
    put(B, f'{c}25', f'={p}25')
    put(B, f'{c}29', f'={p}29+CFS!{c}9+CFS!{c}31', 'link')
    put(B, f'{c}30', f'={p}30+IS!{c}20+CFS!{c}29+CFS!{c}30', 'link')
    put(B, f'{c}31', f'={p}31')
    put(B, f'{c}32', f'={p}32')
for c in ALL:
    put(B, f'{c}12', f'=SUM({c}7:{c}11)')
    put(B, f'{c}17', f'={c}12+SUM({c}13:{c}16)')
    put(B, f'{c}23', f'=SUM({c}20:{c}22)')
    put(B, f'{c}26', f'={c}23+{c}24+{c}25')
    put(B, f'{c}33', f'=SUM({c}29:{c}32)')
    put(B, f'{c}34', f'={c}26+{c}33')
B['A16'].comment = Comment('Historical: operating lease ROU assets + equity and other investments + other long-term assets.', 'Claude')
B['A21'].comment = Comment('Accrued compensation + short-term income taxes + short-term unearned revenue + other current liabilities.', 'Claude')
B['A25'].comment = Comment('LT income taxes + LT unearned revenue + deferred income taxes + operating lease liabilities + other LT liabilities.', 'Claude')

# ---------------- Cash flow ----------------
C = wb['CFS']
for i, c in enumerate(H):
    other_op = (hist['cfs_da_other'][i] - hist['dep'][i] - hist['amort'][i] + hist['gains'][i]
                + hist['c_olta'][i] + hist['c_oltl'][i])
    put(C, f'{c}9', hist['sbc'][i])
    put(C, f'{c}10', hist['dtax'][i])
    put(C, f'{c}11', hist['c_ar'][i])
    put(C, f'{c}12', hist['c_inv'][i])
    put(C, f'{c}13', hist['c_oca'][i])
    put(C, f'{c}14', hist['c_ap'][i])
    put(C, f'{c}15', hist['c_unearned'][i] + hist['c_itax'][i] + hist['c_ocl'][i])
    put(C, f'{c}16', other_op)
    put(C, f'{c}20', hist['capex'][i])
    put(C, f'{c}21', hist['acq'][i])
    put(C, f'{c}22', hist['inv_net'][i])
    put(C, f'{c}23', hist['oth_inv'][i])
    put(C, f'{c}27', hist['debt_iss'][i])
    put(C, f'{c}28', hist['debt_rep'][i])
    put(C, f'{c}29', hist['divs'][i])
    put(C, f'{c}30', hist['buyb'][i])
    put(C, f'{c}31', hist['stk_iss'][i])
    put(C, f'{c}32', hist['oth_fin'][i])
    put(C, f'{c}35', hist['fx'][i])
    put(C, f'{c}41', hist['cfo'][i])
    put(C, f'{c}42', f'={c}17-{c}41')
for c in P:
    p = prev(c)
    put(C, f'{c}9', f'=IS!{c}6*Assumptions!{c}36', 'link')
    put(C, f'{c}10', 0, 'in')
    put(C, f'{c}11', f'=BS!{p}9-BS!{c}9')
    put(C, f'{c}12', f'=BS!{p}10-BS!{c}10')
    put(C, f'{c}13', f'=BS!{p}11-BS!{c}11')
    put(C, f'{c}14', f'=BS!{c}20-BS!{p}20')
    put(C, f'{c}15', f'=BS!{c}21-BS!{p}21')
    put(C, f'{c}16', 0, 'in')
    put(C, f'{c}20', f'=-Schedules!{c}8')
    put(C, f'{c}21', 0, 'in')
    put(C, f'{c}22', 0, 'in')
    put(C, f'{c}23', 0, 'in')
    put(C, f'{c}27', f'=Schedules!{c}30')
    put(C, f'{c}28', f'=Schedules!{c}31')
    put(C, f'{c}29', f'=-IS!{c}20*Assumptions!{c}31', 'link')
    put(C, f'{c}30', f'=-Assumptions!{c}32', 'link')
    put(C, f'{c}31', f'=Assumptions!{c}37', 'link')
    put(C, f'{c}32', 0, 'in')
    put(C, f'{c}35', 0, 'in')
for c in ALL:
    p = prev(c)
    put(C, f'{c}7', f'=IS!{c}20')
    put(C, f'{c}8', f'=IS!{c}13')
    put(C, f'{c}17', f'=SUM({c}7:{c}16)')
    put(C, f'{c}24', f'=SUM({c}20:{c}23)')
    put(C, f'{c}33', f'=SUM({c}27:{c}32)')
    put(C, f'{c}36', f'={c}17+{c}24+{c}33+{c}35')
    if c == 'B':
        put(C, 'B37', 13931, 'in')
    else:
        put(C, f'{c}37', f'={p}38')
    put(C, f'{c}38', f'={c}37+{c}36')
C['A40'] = 'Memo: reported figures (10-K)'
C['A40'].font = Font(bold=True)
C['A41'] = 'Net cash from operations as reported'
C['A42'] = 'Check: model CFO less reported CFO'
C['A16'].comment = Comment('Historical: reported "Depreciation, amortization, and other" less note D&A, plus net recognized (gains)/losses on investments and derivatives, plus changes in other LT assets and other LT liabilities.', 'Claude')
C['A15'].comment = Comment('Historical: changes in unearned revenue + income taxes + other current liabilities.', 'Claude')
C['B37'].comment = Comment('FY2022 ending cash and cash equivalents per 10-K.', 'Claude')

# ---------------- Schedules ----------------
S = wb['Schedules']
S['A36'] = 'Cash + ST investments, beginning of year (basis for interest income)'
for i, c in enumerate(H):
    p = prev(c)
    if c == 'B':
        put(S, 'B7', 74398, 'in')
        put(S, 'B11', 11298, 'in')
        put(S, 'B29', 2749 + 47032, 'in')
        put(S, 'B36', 13931 + 90826, 'in')
    else:
        put(S, f'{c}7', f'={p}10')
        put(S, f'{c}11', f'={p}13')
        put(S, f'{c}29', f'={p}32')
        put(S, f'{c}36', f'=BS!{p}7+BS!{p}8')
    put(S, f'{c}8', f'=-CFS!{c}20')
    put(S, f'{c}9', hist['dep'][i])
    put(S, f'{c}10', f'=BS!{c}13')
    put(S, f'{c}12', hist['amort'][i])
    put(S, f'{c}13', f'=BS!{c}15')
    put(S, f'{c}30', f'=CFS!{c}27')
    put(S, f'{c}31', f'=CFS!{c}28')
    put(S, f'{c}32', f'=BS!{c}22+BS!{c}24')
    put(S, f'{c}35', f'=IS!{c}15')
    put(S, f'{c}38', f'=IS!{c}16')
    put(S, f'{c}34', f'=IFERROR({c}35/{c}33,0)', fmt=PCT)
    put(S, f'{c}37', f'=IFERROR({c}38/{c}36,0)', fmt=PCT)
for c in P:
    p = prev(c)
    put(S, f'{c}7', f'={p}10')
    put(S, f'{c}8', f'=IS!{c}6*Assumptions!{c}25', 'link')
    put(S, f'{c}9', f'={c}7*Assumptions!{c}26', 'link')
    put(S, f'{c}10', f'={c}7+{c}8-{c}9')
    put(S, f'{c}11', f'={p}13')
    put(S, f'{c}12', f'=MIN({c}11,Assumptions!{c}27)', 'link')
    put(S, f'{c}13', f'={c}11-{c}12')
    put(S, f'{c}17', f'=IS!{c}6*Assumptions!{c}18/365', 'link')
    put(S, f'{c}18', f'=IS!{c}7*Assumptions!{c}19/365', 'link')
    put(S, f'{c}19', f'=IS!{c}6*Assumptions!{c}21', 'link')
    put(S, f'{c}20', f'=IS!{c}7*Assumptions!{c}20/365', 'link')
    put(S, f'{c}21', f'=IS!{c}6*Assumptions!{c}22', 'link')
    put(S, f'{c}29', f'={p}32')
    put(S, f'{c}30', f'=MAX(Assumptions!{c}30,0)', 'link')
    put(S, f'{c}31', f'=MIN(Assumptions!{c}30,0)', 'link')
    put(S, f'{c}32', f'={c}29+{c}30+{c}31')
    put(S, f'{c}34', f'=Assumptions!{c}12', 'link', PCT)
    put(S, f'{c}35', f'={c}33*{c}34')
    put(S, f'{c}36', f'=BS!{p}7+BS!{p}8')
    put(S, f'{c}37', f'=Assumptions!{c}13', 'link', PCT)
    put(S, f'{c}38', f'={c}36*{c}37')
for c in ALL:
    put(S, f'{c}14', f'={c}9+{c}12')
    if c in H:
        for r, br in zip((17, 18, 19, 20, 21), (9, 10, 11, 20, 21)):
            put(S, f'{c}{r}', f'=BS!{c}{br}')
    put(S, f'{c}22', f'={c}17+{c}18+{c}19-{c}20-{c}21')
    if c != 'B':
        put(S, f'{c}23', f'={prev(c)}22-{c}22')
    put(S, f'{c}24', f'={c}17/IS!{c}6*365', fmt='0.0')
    put(S, f'{c}25', f'={c}18/IS!{c}7*365', fmt='0.0')
    put(S, f'{c}26', f'={c}20/IS!{c}7*365', fmt='0.0')
    put(S, f'{c}33', f'=AVERAGE({c}29,{c}32)')
S['A9'].comment = Comment('Historical depreciation expense per 10-K PP&E note (reported in $bn, converted to $m). FY2023 from model knowledge.', 'Claude')
S['A12'].comment = Comment('Historical intangible amortization per 10-K note ($bn). Projected = 10-K estimated future amortization schedule.', 'Claude')

# ---------------- Checks ----------------
K = wb['Checks']
K['A8'] = 'Ending cash at or above minimum cash balance (1 = yes)'
K['A9'] = 'Net income: IS vs CFS (should be 0)'
K['A11'] = 'All checks pass (1 = yes)'
for c in ALL:
    put(K, f'{c}6', f'=ROUND(BS!{c}17-BS!{c}34,3)')
    put(K, f'{c}7', f'=ROUND(CFS!{c}38-BS!{c}7,3)')
    put(K, f'{c}8', f'=IF(BS!{c}7>=IF(ISNUMBER(Assumptions!{c}33),Assumptions!{c}33,0),1,0)', fmt='0')
    put(K, f'{c}9', f'=ROUND(IS!{c}20-CFS!{c}7,3)')
    put(K, f'{c}11', f'=IF(AND({c}6=0,{c}7=0,{c}8=1,{c}9=0),1,0)', fmt='0')
K['A12'] = 'Model-wide: all periods pass (1 = yes)'
put(K, 'B12', '=MIN(B11:J11)', fmt='0')

wb.save(OUT)
print('saved', OUT)
