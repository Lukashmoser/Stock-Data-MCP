import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = "/home/claude/results/2026-10-13/runs/FA-03/B/MSFT_Trading_Comps_2026-10-14.xlsx"
FONT = "Times New Roman"
NAVY = PatternFill("solid", fgColor="1F4E79")
LBLUE = PatternFill("solid", fgColor="D9E1F2")
GREY = PatternFill("solid", fgColor="F2F2F2")
BLUE_TXT = "0000FF"
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
AUTH = "Comps build"

QSRC = "stock-data get_financials (SEC XBRL, quarterly)"
F10Q = {
    "MSFT": "10-K filed 2026-07-29 (FY ended 2026-06-30)",
    "GOOGL": "10-Q filed 2026-07-23 (Q ended 2026-06-30)",
    "AMZN": "10-Q filed 2026-07-31 (Q ended 2026-06-30)",
    "META": "10-Q filed 2026-07-30 (Q ended 2026-06-30)",
    "AAPL": "10-Q filed 2026-07-31 (Q ended 2026-06-27)",
    "ORCL": "10-Q filed 2026-09-11 (Q ended 2026-08-31)",
}

cos = [
    # ticker, name, qend, fiscal label, price, shares, shares_src, cash, sti, debt_cur, debt_nc, pref, nci
    ("MSFT", "Microsoft", "2026-06-30", "FY26 Q4", 533.69, 7427.0, "CommonStockSharesOutstanding at 2026-06-30 (stock-data balance)", 20935, 55908, 9227, 31067, 0, 0),
    ("GOOGL", "Alphabet", "2026-06-30", "Q2 CY26", 347.64, 12230.0, "CommonStockSharesOutstanding (all classes) at 2026-06-30 (stock-data balance)", 55911, 186563, 1999, 98165, 0, 0),
    ("AMZN", "Amazon", "2026-06-30", "Q2 CY26", 254.08, 10783.0, "CommonStockSharesOutstanding at 2026-06-30 (stock-data balance)", 78213, 44775, 3330, 128894, 0, 0),
    ("META", "Meta Platforms", "2026-06-30", "Q2 CY26", 737.54, 2547.506, "10-Q cover (filed 2026-07-30): 2,205.1M Class A + 342.4M Class B as of 2026-07-24", 15462, 74798, 0, 83664, 0, 0),
    ("AAPL", "Apple", "2026-06-27", "FY26 Q3", 342.96, 14608.963, "CommonStockSharesOutstanding at 2026-06-27 (stock-data balance)", 39544, 22855, 13004, 71340, 0, 0),
    ("ORCL", "Oracle", "2026-08-31", "FY27 Q1", 130.55, 3024.0, "CommonStockSharesOutstanding at 2026-08-31 (stock-data balance)", 36369, 708, 7625, 117712, 4954, 424),
]
debt_note = {
    "MSFT": "Current portion of LT debt 9,227 / LT debt 31,067 at 2026-06-30 (10-K FY26; earnings release Ex.99.1 8-K 2026-07-29). Excludes finance/operating leases.",
    "GOOGL": "DebtCurrent 1,999 / LT debt 98,165 at 2026-06-30 (stock-data balance; 10-Q filed 2026-07-23). Excludes leases.",
    "AMZN": "LongTermDebtCurrent 3,330 / LT debt 128,894 at 2026-06-30 (stock-data balance). Excludes leases and other short-term borrowings in accrued liabilities.",
    "META": "No current debt tagged; LT debt 83,664 at 2026-06-30 (stock-data balance). Excludes leases.",
    "AAPL": "Commercial paper 1,997 + term debt current 11,007 = 13,004; term debt non-current 71,340 at 2026-06-27 (10-Q filed 2026-07-31, balance sheet R4).",
    "ORCL": "Notes payable & other borrowings current 7,625 / non-current 117,712 at 2026-08-31 (10-Q filed 2026-09-11, balance sheet R2).",
}

rev8 = {
    "MSFT": [65585, 69632, 70066, 76441, 77673, 81273, 82886, 90007],
    "GOOGL": [88268, 96469, 90234, 96428, 102346, 113829, 109896, 119796],
    "AMZN": [158877, 187792, 155667, 167702, 180169, 213386, 181519, 200606],
    "META": [40589, 48385, 42314, 47516, 51242, 59893, 56311, 60801],
    "AAPL": [94930, 124300, 95359, 94036, 102466, 143756, 111184, 109417],
    "ORCL": [14059, 14130, 15903, 14926, 16058, 17190, 19184, 19345],
}
qlabels = {
    "MSFT": ["2024-09-30", "2024-12-31", "2025-03-31", "2025-06-30", "2025-09-30", "2025-12-31", "2026-03-31", "2026-06-30"],
    "GOOGL": ["2024-09-30", "2024-12-31", "2025-03-31", "2025-06-30", "2025-09-30", "2025-12-31", "2026-03-31", "2026-06-30"],
    "AMZN": ["2024-09-30", "2024-12-31", "2025-03-31", "2025-06-30", "2025-09-30", "2025-12-31", "2026-03-31", "2026-06-30"],
    "META": ["2024-09-30", "2024-12-31", "2025-03-31", "2025-06-30", "2025-09-30", "2025-12-31", "2026-03-31", "2026-06-30"],
    "AAPL": ["2024-09-28", "2024-12-28", "2025-03-29", "2025-06-28", "2025-09-27", "2025-12-27", "2026-03-28", "2026-06-27"],
    "ORCL": ["2024-11-30", "2025-02-28", "2025-05-31", "2025-08-31", "2025-11-30", "2026-02-28", "2026-05-31", "2026-08-31"],
}
derivedQ = {  # quarters derived as FY - 9M by data tool
    "MSFT": ["2025-06-30", "2026-06-30"], "GOOGL": ["2024-12-31", "2025-12-31"], "AMZN": ["2024-12-31", "2025-12-31"],
    "META": ["2024-12-31", "2025-12-31"], "AAPL": ["2024-09-28", "2025-09-27"], "ORCL": ["2025-05-31", "2026-05-31"],
}
ebit4 = {"MSFT": [37961, 38275, 38398, 40603], "GOOGL": [31228, 35934, 39696, 40770], "AMZN": [17422, 24977, 23852, 27461],
         "META": [20535, 24745, 22872, 18775], "AAPL": [32427, 50852, 35885, 35695], "ORCL": [4731, 5464, 6133, 6728]}
ni4 = {"MSFT": [27747, 38458, 31778, 35766], "GOOGL": [34979, 34455, 62578, 112193], "AMZN": [21187, 21192, 30255, 62647],
       "META": [2709, 22768, 26773, 15848], "AAPL": [27466, 42097, 29578, 29789], "ORCL": [6135, 3721, 4304, 4760]}
tax4 = {"MSFT": [6554, 9788, 7562, 8281], "GOOGL": [9008, 4662, 14834, 26560], "AMZN": [6910, 4946, 9560, 18199],
        "META": [18954, 2586, -5021, 2908], "AAPL": [5338, 8905, 6255, 6478], "ORCL": [207, 695, 1065, 847]}
pti4 = {"MSFT": [34301, 48246, 39340, 44047], "GOOGL": [43987, 39117, 77412, 138753], "AMZN": [28170, 26605, 39834, 80857],
        "META": [21663, 25354, 21752, 18756], "AAPL": [32804, 51002, 35833, 36267], "ORCL": None}
# D&A: (values[4], formula pattern over cols, method text, comment per value)
da = {
    "MSFT": ([38534, None, None, None], "={a}", "FY26 (Jul-25 to Jun-26) = LTM",
             ["'Depreciation, amortization, and other', twelve months ended 2026-06-30, cash flow statement, MSFT earnings release Ex.99.1 (8-K filed 2026-07-29). Includes 'other' non-cash items.", None, None, None]),
    "GOOGL": ([21136, 13586, 9485, None], "={a}+{b}-{c}", "FY25 + 6M26 - 6M25",
              ["Depreciation of property & equipment FY2025, 10-K filed 2026-02-05 (cash flow statement, R10).",
               "Depreciation of P&E, six months ended 2026-06-30, 10-Q filed 2026-07-23.",
               "Depreciation of P&E, six months ended 2025-06-30, 10-Q filed 2026-07-23 (comparative).", None]),
    "AMZN": ([16796, 19471, 18945, 19988], "=SUM({a}:{d})", "Sum of last 4 quarters", [QSRC + " depreciation_amortization, Q ended 2025-09-30",
             QSRC + ", Q ended 2025-12-31 (derived FY-9M)", QSRC + ", Q ended 2026-03-31", QSRC + ", Q ended 2026-06-30 (derived 6M-3M)"]),
    "META": ([4963, 5411, 5999, 6356], "=SUM({a}:{d})", "Sum of last 4 quarters", [QSRC + " depreciation_amortization, Q ended 2025-09-30 (derived)",
             QSRC + ", Q ended 2025-12-31 (derived)", QSRC + ", Q ended 2026-03-31", QSRC + ", Q ended 2026-06-30 (derived)"]),
    "AAPL": ([3127, 3214, 3439, 3320], "=SUM({a}:{d})", "Sum of last 4 quarters", [QSRC + " depreciation_amortization, Q ended 2025-09-27 (derived)",
             QSRC + ", Q ended 2025-12-27", QSRC + ", Q ended 2026-03-28 (derived)", QSRC + ", Q ended 2026-06-27 (derived)"]),
    "ORCL": (["=7623+1671", "=3156+202", "=1351+420", None], "={a}+{b}-{c}", "FY26 + Q1 FY27 - Q1 FY26 (depreciation + intangible amortization)",
             ["FY ended 2026-05-31: Depreciation 7,623 + Amortization of intangibles 1,671; 10-K filed 2026-06-22, cash flow statement (R8).",
              "Q ended 2026-08-31: Depreciation 3,156 + Amortization 202; 10-Q filed 2026-09-11.",
              "Q ended 2025-08-31: Depreciation 1,351 + Amortization 420; 10-Q filed 2026-09-11 (comparative).", None]),
}

wb = openpyxl.Workbook()
ws_c = wb.active
ws_c.title = "Comps"
ws_i = wb.create_sheet("Inputs")
ws_n = wb.create_sheet("Notes")


def style(c, bold=False, fill=None, color="000000", align=CENTER, size=11, fmt=None):
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        c.fill = fill
    c.alignment = align
    if fmt:
        c.number_format = fmt


def inp(ws, ref, val, comment, fmt="#,##0"):
    c = ws[ref]
    c.value = val
    style(c, color=BLUE_TXT, fmt=fmt)
    if comment:
        c.comment = Comment(comment, AUTH, width=380, height=110)
    return c


def frm(ws, ref, f, fmt="#,##0", bold=False, fill=None):
    c = ws[ref]
    c.value = f
    style(c, fmt=fmt, bold=bold, fill=fill)
    return c


def section(ws, row, text, ncols):
    for col in range(1, ncols + 1):
        style(ws.cell(row=row, column=col), fill=NAVY)
    c = ws.cell(row=row, column=1, value=text)
    style(c, bold=True, fill=NAVY, color="FFFFFF", align=LEFT, size=12)


def headers(ws, row, hdrs):
    for j, h in enumerate(hdrs, start=1):
        c = ws.cell(row=row, column=j, value=h)
        style(c, bold=True, fill=LBLUE)


PRICE_SRC = "stock-data get_quote (Yahoo Finance), last price for 2026-10-14 session, fetched 2026-10-15T01:58Z"

# ---------------------------------------------------------------- INPUTS
wi = ws_i
wi["A1"] = "INPUTS - MARKET DATA, BALANCE SHEET AND QUARTERLY FINANCIALS (USD millions except per-share)"
style(wi["A1"], bold=True, align=LEFT, size=12)
wi["A2"] = "Blue = hard-coded input (each has a cell comment with its source); black = formula."
style(wi["A2"], align=LEFT)

# Section A: market & balance sheet, rows 5..10
section(wi, 3, "A. MARKET DATA & BALANCE SHEET (latest reported balance sheet)", 14)
headers(wi, 4, ["Ticker", "Company", "Latest Q end (LTM)", "Fiscal period", "Share price ($)", "Shares out (M)",
                "Cash & equiv.", "ST investments", "Debt - current", "Debt - non-current", "Preferred stock", "Noncontrolling int.", "Total debt", "Net debt (cash)"])
A0 = 5
rowA = {}
for k, co in enumerate(cos):
    r = A0 + k
    t, name, qend, fl, px, sh, shsrc, cash, sti, dc, dn, pref, nci = co
    rowA[t] = r
    style(wi.cell(row=r, column=1, value=t), bold=True)
    style(wi.cell(row=r, column=2, value=name), bold=True)
    inp(wi, f"C{r}", qend, f"Latest reported quarter end. Source: {F10Q[t]}; stock-data get_filings.", fmt="@")
    inp(wi, f"D{r}", fl, "Company fiscal label of latest quarter.", fmt="@")
    inp(wi, f"E{r}", px, PRICE_SRC, fmt="#,##0.00")
    inp(wi, f"F{r}", sh, shsrc + ". Basic shares outstanding; dilution not added.", fmt="#,##0.0")
    inp(wi, f"G{r}", cash, f"CashAndCashEquivalentsAtCarryingValue at {qend}; stock-data get_financials balance ({F10Q[t]}).")
    inp(wi, f"H{r}", sti, f"Current short-term investments / marketable securities at {qend}; stock-data get_financials balance. Non-current securities and equity stakes excluded.")
    inp(wi, f"I{r}", dc, debt_note[t])
    inp(wi, f"J{r}", dn, debt_note[t])
    inp(wi, f"K{r}", pref, ("Series D 6.50% mandatory convertible preferred, carrying value at 2026-08-31; 10-Q filed 2026-09-11 (R2)." if t == "ORCL" else "Assumed zero: no preferred stock on balance sheet."))
    inp(wi, f"L{r}", nci, ("Noncontrolling interests at 2026-08-31; 10-Q filed 2026-09-11 (R2)." if t == "ORCL" else "Assumed zero / immaterial: no NCI line captured by data source."))
    frm(wi, f"M{r}", f"=I{r}+J{r}")
    frm(wi, f"N{r}", f"=M{r}+K{r}+L{r}-G{r}-H{r}", fmt='#,##0;(#,##0)')

# Section B: quarterly revenue
B_hdr = 13
section(wi, B_hdr - 1, "B. QUARTERLY REVENUE (oldest to newest; Q-7 .. Q0 where Q0 = latest reported quarter)", 14)
headers(wi, B_hdr, ["Ticker", "Q-7", "Q-6", "Q-5", "Q-4", "Q-3", "Q-2", "Q-1", "Q0", "LTM revenue", "Prior LTM revenue", "LTM growth %", "Q0 end", ""])
rowB = {}
for k, co in enumerate(cos):
    t = co[0]
    r = B_hdr + 1 + k
    rowB[t] = r
    style(wi.cell(row=r, column=1, value=t), bold=True)
    for j, v in enumerate(rev8[t]):
        q = qlabels[t][j]
        note = f"{QSRC} revenue, quarter ended {q}." + (" Derived by tool as FY less 9M YTD." if q in derivedQ[t] else "")
        inp(wi, f"{L(2 + j)}{r}", v, note)
    frm(wi, f"J{r}", f"=SUM(F{r}:I{r})")
    frm(wi, f"K{r}", f"=SUM(B{r}:E{r})")
    frm(wi, f"L{r}", f"=J{r}/K{r}-1", fmt="0.0%")
    frm(wi, f"M{r}", f"=Inputs!C{rowA[t]}", fmt="@")


def block4(title, hdr_row, data, label, rowmap, extra=None):
    section(wi, hdr_row - 1, title, 14)
    headers(wi, hdr_row, ["Ticker", "Q-3", "Q-2", "Q-1", "Q0", f"LTM {label}", "", "", "", "", "", "", "", ""])
    for k, co in enumerate(cos):
        t = co[0]
        r = hdr_row + 1 + k
        rowmap[t] = r
        style(wi.cell(row=r, column=1, value=t), bold=True)
        vals = data[t]
        for j in range(4):
            q = qlabels[t][4 + j]
            col = L(2 + j)
            if vals is None:
                # ORCL pretax = NI + tax
                frm(wi, f"{col}{r}", f"={col}{rowNI[t]}+{col}{rowTX[t]}")
            else:
                note = f"{QSRC} {label}, quarter ended {q}." + (" Derived by tool as FY less 9M YTD." if q in derivedQ[t] else "")
                inp(wi, f"{col}{r}", vals[j], note, fmt='#,##0;(#,##0)')
        frm(wi, f"F{r}", f"=SUM(B{r}:E{r})", fmt='#,##0;(#,##0)')


rowEBIT, rowNI, rowTX, rowPT, rowDA = {}, {}, {}, {}, {}
block4("C. OPERATING INCOME (GAAP, last 4 quarters)", 22, ebit4, "operating income", rowEBIT)
block4("D. NET INCOME ATTRIBUTABLE TO COMMON (GAAP, last 4 quarters)", 31, ni4, "net income", rowNI)
block4("E. INCOME TAX EXPENSE (last 4 quarters)", 40, tax4, "income tax", rowTX)
block4("F. PRE-TAX INCOME (last 4 quarters; ORCL = net income + tax, pre-tax not tagged)", 49, pti4, "pre-tax income", rowPT)
# pretax ORCL refs rowNI and rowTX defined already

# Section G: D&A
G_hdr = 58
section(wi, G_hdr - 1, "G. DEPRECIATION & AMORTIZATION BUILD (cash flow statement)", 14)
headers(wi, G_hdr, ["Ticker", "Comp. 1", "Comp. 2", "Comp. 3", "Comp. 4", "LTM D&A", "Build method", "", "", "", "", "", "", ""])
for k, co in enumerate(cos):
    t = co[0]
    r = G_hdr + 1 + k
    rowDA[t] = r
    style(wi.cell(row=r, column=1, value=t), bold=True)
    vals, pat, method, notes = da[t]
    for j in range(4):
        if vals[j] is not None:
            inp(wi, f"{L(2 + j)}{r}", vals[j], notes[j])
    frm(wi, f"F{r}", pat.format(a=f"B{r}", b=f"C{r}", c=f"D{r}", d=f"E{r}"))
    c = wi.cell(row=r, column=7, value=method)
    style(c, align=LEFT)

for col in range(1, 15):
    wi.column_dimensions[L(col)].width = 15
wi.column_dimensions["B"].width = 16
wi.freeze_panes = "B5"

# ---------------------------------------------------------------- COMPS
wc = ws_c
NC = 12
wc["A1"] = "LARGE-CAP TECHNOLOGY - COMPARABLE COMPANY ANALYSIS: MICROSOFT (MSFT)"
style(wc["A1"], bold=True, align=LEFT, size=14)
wc["A2"] = "Target: Microsoft (MSFT)  |  Peers: Alphabet (GOOGL) • Amazon (AMZN) • Meta Platforms (META) • Apple (AAPL) • Oracle (ORCL)"
style(wc["A2"], align=LEFT)
wc["A3"] = "Prices as of 2026-10-14 close | LTM to each company's latest reported quarter | USD millions except per-share data and ratios | Peer statistics exclude MSFT"
style(wc["A3"], align=LEFT)

order = ["MSFT", "GOOGL", "AMZN", "META", "AAPL", "ORCL"]
names = {c[0]: c[1] for c in cos}

# Operating section
section(wc, 5, "OPERATING STATISTICS & FINANCIAL METRICS (LTM)", NC)
headers(wc, 6, ["Company", "Ticker", "LTM through", "LTM revenue", "Revenue growth (YoY LTM)", "LTM EBIT", "EBIT margin",
                "LTM D&A", "LTM EBITDA", "EBITDA margin", "LTM net income", "Net margin"])
op_rows = {}
r = 7
for t in order:
    if t == "GOOGL":
        r += 1  # spacer between target and peers
    op_rows[t] = r
    style(wc.cell(row=r, column=1, value=names[t]), bold=True, align=LEFT)
    style(wc.cell(row=r, column=2, value=t))
    frm(wc, f"C{r}", f"=Inputs!C{rowA[t]}", fmt="@")
    frm(wc, f"D{r}", f"=Inputs!J{rowB[t]}")
    frm(wc, f"E{r}", f"=Inputs!L{rowB[t]}", fmt="0.0%")
    frm(wc, f"F{r}", f"=Inputs!F{rowEBIT[t]}")
    frm(wc, f"G{r}", f"=F{r}/D{r}", fmt="0.0%")
    frm(wc, f"H{r}", f"=Inputs!F{rowDA[t]}")
    frm(wc, f"I{r}", f"=F{r}+H{r}")
    frm(wc, f"J{r}", f"=I{r}/D{r}", fmt="0.0%")
    frm(wc, f"K{r}", f"=Inputs!F{rowNI[t]}")
    frm(wc, f"L{r}", f"=K{r}/D{r}", fmt="0.0%")
    r += 1
p0, p1 = op_rows["GOOGL"], op_rows["ORCL"]
stats = [("Maximum", "MAX({r})"), ("75th Percentile", "QUARTILE({r},3)"), ("Median", "MEDIAN({r})"),
         ("25th Percentile", "QUARTILE({r},1)"), ("Minimum", "MIN({r})")]
r = p1 + 2
op_stat = {}
for lab, f in stats:
    op_stat[lab] = r
    for col in range(1, NC + 1):
        style(wc.cell(row=r, column=col), fill=GREY)
    style(wc.cell(row=r, column=1, value=f"Peer {lab}"), bold=True, fill=GREY, align=LEFT)
    for col in ["E", "G", "J", "L"]:
        frm(wc, f"{col}{r}", "=" + f.format(r=f"{col}{p0}:{col}{p1}"), fmt="0.0%", fill=GREY)
    r += 1
mr = op_rows["MSFT"]
style(wc.cell(row=r, column=1, value="MSFT vs peer median (pp)"), bold=True, align=LEFT)
for col in ["E", "G", "J", "L"]:
    frm(wc, f"{col}{r}", f"={col}{mr}-{col}{op_stat['Median']}", fmt='+0.0%;-0.0%')
op_end = r

# Valuation section
v0 = op_end + 2
section(wc, v0, "VALUATION MULTIPLES", NC)
headers(wc, v0 + 1, ["Company", "Ticker", "Share price ($)", "Shares out (M)", "Market cap", "Net debt (cash) incl. pref & NCI",
                     "Enterprise value", "EV / LTM revenue", "EV / LTM EBITDA", "P/E (LTM GAAP)", "LTM effective tax rate", "P/E (LTM EBIT after tax)"])
val_rows = {}
r = v0 + 2
for t in order:
    if t == "GOOGL":
        r += 1
    val_rows[t] = r
    o = op_rows[t]
    style(wc.cell(row=r, column=1, value=names[t]), bold=True, align=LEFT)
    style(wc.cell(row=r, column=2, value=t))
    frm(wc, f"C{r}", f"=Inputs!E{rowA[t]}", fmt="#,##0.00")
    frm(wc, f"D{r}", f"=Inputs!F{rowA[t]}", fmt="#,##0.0")
    frm(wc, f"E{r}", f"=C{r}*D{r}")
    frm(wc, f"F{r}", f"=Inputs!N{rowA[t]}", fmt='#,##0;(#,##0)')
    frm(wc, f"G{r}", f"=E{r}+F{r}")
    frm(wc, f"H{r}", f"=G{r}/D{o}", fmt='0.0"x"')
    frm(wc, f"I{r}", f"=G{r}/I{o}", fmt='0.0"x"')
    frm(wc, f"J{r}", f"=E{r}/K{o}", fmt='0.0"x"')
    frm(wc, f"K{r}", f"=Inputs!F{rowTX[t]}/Inputs!F{rowPT[t]}", fmt="0.0%")
    frm(wc, f"L{r}", f"=E{r}/(F{o}*(1-K{r}))", fmt='0.0"x"')
    r += 1
q0, q1 = val_rows["GOOGL"], val_rows["ORCL"]
r = q1 + 2
val_stat = {}
for lab, f in stats:
    val_stat[lab] = r
    for col in range(1, NC + 1):
        style(wc.cell(row=r, column=col), fill=GREY)
    style(wc.cell(row=r, column=1, value=f"Peer {lab}"), bold=True, fill=GREY, align=LEFT)
    for col in ["H", "I", "J", "L"]:
        frm(wc, f"{col}{r}", "=" + f.format(r=f"{col}{q0}:{col}{q1}"), fmt='0.0"x"', fill=GREY)
    frm(wc, f"K{r}", "=" + f.format(r=f"K{q0}:K{q1}"), fmt="0.0%", fill=GREY)
    r += 1
vm = val_rows["MSFT"]
style(wc.cell(row=r, column=1, value="MSFT premium / (discount) to peer median"), bold=True, align=LEFT)
for col in ["H", "I", "J", "L"]:
    frm(wc, f"{col}{r}", f"={col}{vm}/{col}{val_stat['Median']}-1", fmt='+0.0%;-0.0%')
val_end = r

# Implied valuation
i0 = val_end + 2
section(wc, i0, "IMPLIED MSFT VALUE PER SHARE AT PEER MULTIPLES (USD)", NC)
headers(wc, i0 + 1, ["Multiple", "MSFT LTM metric", "Peer 25th pct", "Peer median", "Peer 75th pct",
                     "Implied price @25th", "Implied price @median", "Implied price @75th", "Current price", "Upside/(down) @median", "", ""])
mo = op_rows["MSFT"]
specs = [("EV / LTM revenue", f"D{mo}", "H", True), ("EV / LTM EBITDA", f"I{mo}", "I", True),
         ("P/E (LTM GAAP)", f"K{mo}", "J", False), ("P/E (LTM EBIT after tax)", f"F{mo}*(1-K{vm})", "L", False)]
r = i0 + 2
for lab, metric, mcol, is_ev in specs:
    style(wc.cell(row=r, column=1, value=lab), bold=True, align=LEFT)
    frm(wc, f"B{r}", "=" + metric)
    frm(wc, f"C{r}", f"={mcol}{val_stat['25th Percentile']}", fmt='0.0"x"')
    frm(wc, f"D{r}", f"={mcol}{val_stat['Median']}", fmt='0.0"x"')
    frm(wc, f"E{r}", f"={mcol}{val_stat['75th Percentile']}", fmt='0.0"x"')
    for src, dst in [("C", "F"), ("D", "G"), ("E", "H")]:
        if is_ev:
            f = f"=($B{r}*{src}{r}-$F${vm})/$D${vm}"
        else:
            f = f"=$B{r}*{src}{r}/$D${vm}"
        frm(wc, f"{dst}{r}", f, fmt="#,##0.00")
    frm(wc, f"I{r}", f"=$C${vm}", fmt="#,##0.00")
    frm(wc, f"J{r}", f"=G{r}/I{r}-1", fmt='+0.0%;-0.0%')
    r += 1
imp_end = r

for col in range(1, NC + 1):
    wc.column_dimensions[L(col)].width = 16
wc.column_dimensions["A"].width = 34
for rr in range(5, imp_end + 1):
    wc.row_dimensions[rr].height = 22
for rr in (6, v0 + 1, i0 + 1):
    wc.row_dimensions[rr].height = 44
wc.freeze_panes = "C7"

# ---------------------------------------------------------------- NOTES
notes = [
    ("NOTES & METHODOLOGY", True),
    ("Purpose: trading comparables for an investment-committee valuation of Microsoft (MSFT). Peer set as specified: GOOGL, AMZN, META, AAPL, ORCL. Statistics (max / 75th / median / 25th / min) are computed on the five peers only; MSFT is shown as the target and benchmarked against the peer median.", False),
    ("", False),
    ("Data sources", True),
    ("Share prices and quote market caps: stock-data get_quote (Yahoo Finance), 2026-10-14 session, fetched 2026-10-15T01:58Z.", False),
    ("Income statement, balance sheet and cash-flow lines: stock-data get_financials (standardized SEC XBRL, quarterly). Fourth fiscal quarters are derived by the tool as full year less 9M YTD.", False),
    ("Items not available in the data tool were taken from SEC EDGAR filings via WebFetch: MSFT D&A (earnings release Ex.99.1, 8-K filed 2026-07-29); GOOGL depreciation (10-K filed 2026-02-05; 10-Q filed 2026-07-23); ORCL D&A, debt, preferred and NCI (10-K filed 2026-06-22; 10-Q filed 2026-09-11); AAPL commercial paper (10-Q filed 2026-07-31); META shares outstanding (10-Q cover, filed 2026-07-30).", False),
    ("", False),
    ("Definitions", True),
    ("LTM = sum of the four most recent reported fiscal quarters for each company. Latest quarter ends: MSFT 2026-06-30 (FY26 Q4), GOOGL/AMZN/META 2026-06-30, AAPL 2026-06-27 (FY26 Q3), ORCL 2026-08-31 (FY27 Q1). Fiscal calendars are not aligned; ORCL is two months later than the others.", False),
    ("Revenue growth = LTM revenue / prior-year LTM revenue (the four quarters before) - 1.", False),
    ("EBITDA = GAAP operating income + depreciation & amortization from the cash flow statement. Stock-based compensation is NOT added back; no one-off adjustments. MSFT D&A line is 'Depreciation, amortization, and other' and includes minor other non-cash items. GOOGL uses 'Depreciation of property and equipment' only (intangible amortization not separately disclosed in the cash flow statement).", False),
    ("Market cap = share price x basic shares outstanding from the latest balance sheet / 10-Q cover (no treasury-method dilution).", False),
    ("Enterprise value = market cap + total debt (current + non-current, excl. lease liabilities) + preferred stock + noncontrolling interests - cash & equivalents - current short-term investments / marketable securities. Non-current marketable securities and strategic equity stakes (e.g., AAPL non-current securities of 84,118; GOOGL/AMZN/MSFT AI-lab stakes) are excluded from cash.", False),
    ("P/E (LTM GAAP) = market cap / LTM GAAP net income. P/E (LTM EBIT after tax) = market cap / [LTM operating income x (1 - LTM effective tax rate)], a supplementary measure that strips non-operating gains (large equity-security mark-to-market gains at GOOGL and AMZN in 2026, and the MSFT OpenAI-related gain in the Dec-25 quarter) and the META Q3-25 one-off tax charge from the earnings base.", False),
    ("", False),
    ("Interpretation caveats", True),
    ("GOOGL LTM net income (244,205) and AMZN LTM net income (135,281) are inflated by non-operating gains on equity securities (GOOGL 6M-26 'other income' of 135,699). Their GAAP P/Es (~17x / ~20x) understate the operating multiple; use the EBIT-after-tax P/E column for like-for-like comparison.", False),
    ("ORCL EBITDA and multiples reflect heavy AI infrastructure build-out; ORCL's share price fell 8.8% on 2026-10-14, so its multiples moved materially on the pricing date.", False),
    ("Formatting: blue font = hard-coded inputs (each with a cell comment documenting its source); black = formulas. All derived values are live formulas.", False),
]
for k, (txt, bold) in enumerate(notes, start=1):
    c = ws_n.cell(row=k, column=1, value=txt)
    style(c, bold=bold, align=Alignment(horizontal="left", vertical="top", wrap_text=True), size=12 if bold else 11)
    if bold and txt:
        c.fill = NAVY
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=12)
ws_n.column_dimensions["A"].width = 150
for k in range(1, len(notes) + 1):
    if len(notes[k - 1][0]) > 150:
        ws_n.row_dimensions[k].height = 15 * (len(notes[k - 1][0]) // 140 + 1)

wb.save(OUT)
print("saved", OUT, op_rows, val_rows, op_stat, val_stat, i0)
