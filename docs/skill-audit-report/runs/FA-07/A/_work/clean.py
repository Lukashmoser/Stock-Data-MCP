import csv, re, datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
D="/home/claude/results/2026-10-13/runs/FA-07/A/"
rows=list(csv.reader(open(D+"quarterly_export.csv",newline="",encoding="utf-8")))
hdr,raw=rows[0],rows[1:]
log=[]; stats=dict(ticker=0,date=0,num=0,unit=0,wrong=0,dup=0,near=0)
def L(r,col,old,new,cat,why,src=""): log.append([r,col,old,new,cat,why,src])
UNIT={("META","FY2024Q4"),("META","FY2025Q3")}
SRC_GQ2="Alphabet Q2 2026 earnings release (8-K Ex.99.1, 2026-07-22), sec.gov/Archives/edgar/data/0001652044/000165204426000066/googexhibit991q22026.htm: 6M-2025 net income 62,736 less Q2-2025 28,196"
SRC_GQ4="Alphabet Q4 2025 earnings release (2026-02-04), s206.q4cdn.com/479360582/files/doc_financials/2025/q4/2025q4-alphabet-earnings-release.pdf"
SRC_AQ4="Amazon Q4 2025 earnings release (Business Wire, 2026-02-05), businesswire.com/news/home/20260204777163: FY2025 net sales 716,924 less Q1 155,667, Q2 167,702, Q4 213,386; also Amazon Q2 2026 release TTM check"
SRC_MQ4="Meta Q4 2025 earnings release (8-K Ex.99.1, 2026-01-28), sec.gov/Archives/edgar/data/1326801/000162828026003832/meta-12312025xexhibit991.htm"
WRONG={("GOOGL","FY2025Q1","net_income"):(34540,SRC_GQ2+" = 34,540 (also FY2025 NI 132,170 reconciles only with 34,540)"),
 ("AMZN","FY2025Q3","revenue"):(180169,SRC_AQ4+" = 180,169"),
 ("GOOGL","FY2025Q4","revenue"):(113828,SRC_GQ4+": Q4 2025 revenues 113,828"),
 ("META","FY2024Q4","operating_income"):(23365,SRC_MQ4+": Q4 2024 income from operations 23,365")}
VERIFIED_OUTLIER={("GOOGL","FY2026Q1","net_income"):"Large vs revenue but verified: Alphabet Q2 2026 release 6M-2026 NI 174,771 less Q2 112,193 = 62,578 (non-operating equity gains)",
 ("GOOGL","FY2026Q2","net_income"):"Exceeds 90% of revenue but verified in Alphabet Q2 2026 release (112,193; non-operating gains)",
 ("AMZN","FY2026Q2","net_income"):"Large vs operating income but verified in Amazon Q2 2026 release (62,647; non-operating gains)",
 ("META","FY2025Q3","net_income"):"Very low vs operating income but verified: Meta FY2025 NI 60,458 reconciles to 2,709 (one-time tax charge)",
 ("MSFT","FY2026Q2","net_income"):"NI > operating income but verified: MSFT FY2026 NI 133,749 less Q1/Q3/Q4 = 38,458 (GAAP incl. OpenAI-related gain)"}
def pdate(s):
    s=s.strip()
    for f in ("%Y-%m-%d","%m/%d/%Y","%d-%b-%y"):
        try: return datetime.datetime.strptime(s,f).date().isoformat(),f
        except ValueError: pass
    raise ValueError(s)
out=[]
for i,r in enumerate(raw, start=2):
    t,pe,fl,m,v,n=r
    T=t.strip().upper()
    if T!=t: stats["ticker"]+=1; L(i,"ticker",repr(t),T,"ticker","trim whitespace / uppercase")
    iso,f=pdate(pe)
    if f!="%Y-%m-%d": stats["date"]+=1; L(i,"period_end",pe,iso,"date","normalize to YYYY-MM-DD")
    vs=v.strip(); num=vs.replace("$","").replace(",","")
    val=float(num)
    if vs!=num: stats["num"]+=1; L(i,"value",v,num,"text-to-number","strip $ and thousands separators")
    if val==int(val): val=int(val)
    flag=[]; note_out=n.strip()
    key=(T,fl.strip())
    if key in UNIT and val>1_000_000 or (key in UNIT and m=="net_income"):
        nv=val/1000; nv=int(nv) if nv==int(nv) else nv
        stats["unit"]+=1; L(i,"value",val,nv,"unit error","reported in USD thousands; /1,000 to USD millions; reconciles to "+("Meta Q4 2025 release" ),SRC_MQ4); val=nv; flag.append("UNIT_FIXED")
    k3=(T,fl.strip(),m.strip())
    if k3 in WRONG:
        cv,src=WRONG[k3]; stats["wrong"]+=1; L(i,"value",val,cv,"wrong value","does not match company-reported figure",src); val=cv; flag.append("VALUE_CORRECTED")
    if k3 in VERIFIED_OUTLIER: flag.append("OUTLIER_VERIFIED"); L(i,"value",val,val,"outlier check (no change)",VERIFIED_OUTLIER[k3])
    if isinstance(val,float): flag.append("NONINTEGER")
    nt="blank" if note_out=="" else ("numeric" if re.fullmatch(r"[\d.]+",note_out) else "text")
    out.append(dict(src_row=i,ticker=T,period_end=iso,fiscal_label=fl.strip(),metric=m.strip(),value_usd_m=val,notes=note_out,notes_type=nt,flags=";".join(flag)))
# dedupe
seen={}; final=[]
for o in out:
    k=(o["ticker"],o["period_end"],o["metric"])
    if k in seen:
        p=seen[k]
        if p["value_usd_m"]==o["value_usd_m"] and p["notes"]==o["notes"]:
            stats["dup"]+=1; L(o["src_row"],"(row)","duplicate of row %d"%p["src_row"],"removed","dedupe","exact duplicate after normalization")
        else:
            stats["near"]+=1
            keep,drop=(p,o) if isinstance(o["value_usd_m"],float) else (o,p)
            L(drop["src_row"],"(row)","value %s conflicts with row %d value %s"%(drop["value_usd_m"],keep["src_row"],keep["value_usd_m"]),"removed","dedupe (near-duplicate)","kept integer figure matching MSFT FY2025 Q2 reported operating income 31,653; 31,652.6 is an unrounded variant","Microsoft FY2025 10-K / FY2025 Q2 earnings release (model knowledge); FY2025 OI 128,528 per MSFT FY26 Q4 release reconciles with 31,653")
            if keep is o: final[final.index(p)]=o; seen[k]=o
        continue
    seen[k]=o; final.append(o)
final.sort(key=lambda o:(o["ticker"],o["period_end"],["revenue","operating_income","net_income"].index(o["metric"])))
cols=["src_row","ticker","period_end","fiscal_label","metric","value_usd_m","notes","notes_type","flags"]
with open(D+"quarterly_export_clean.csv","w",newline="") as f:
    w=csv.writer(f); w.writerow(cols); [w.writerow([o[c] for c in cols]) for o in final]
lc=["src_row","column","original","new","category","reason","source"]
with open(D+"change_log.csv","w",newline="") as f:
    w=csv.writer(f); w.writerow(lc); w.writerows(log)
# workbook
wb=Workbook(); ws=wb.active; ws.title="Cleaned"; b=Font(bold=True)
ws.append(cols); [setattr(c,"font",b) for c in ws[1]]
hl=PatternFill("solid",fgColor="FFF2CC")
for o in final:
    ws.append([o[c] for c in cols])
    if o["flags"]:
        for c in ws[ws.max_row]: c.fill=hl
    ws.cell(ws.max_row,6).number_format="#,##0"
l=wb.create_sheet("Change_Log"); l.append(lc); [setattr(c,"font",b) for c in l[1]]
for x in log: l.append([str(y) if not isinstance(y,(int,float)) else y for y in x])
r=wb.create_sheet("Raw_with_helpers")
r.append(hdr+["ticker_clean","value_clean"]); [setattr(c,"font",b) for c in r[1]]
for i,x in enumerate(raw,start=2):
    r.append(x+[f"=UPPER(TRIM(A{i}))",f'=VALUE(SUBSTITUTE(SUBSTITUTE(E{i},"$",""),",",""))'])
s=wb.create_sheet("Summary")
s.append(["metric","count"])
for k,v in [("rows in",len(raw)),("rows out",len(final)),("exact duplicates removed",stats["dup"]),("near-duplicates removed",stats["near"]),("ticker cells fixed",stats["ticker"]),("dates normalized",stats["date"]),("text-to-number conversions",stats["num"]),("unit errors fixed",stats["unit"]),("wrong values corrected",stats["wrong"]),("notes column: numeric entries",sum(o['notes_type']=='numeric' for o in out)),("notes column: text entries",sum(o['notes_type']=='text' for o in out)),("notes column: blank",sum(o['notes_type']=='blank' for o in out))]:
    s.append([k,v])
for wsx in wb: 
    for col in wsx.columns: wsx.column_dimensions[col[0].column_letter].width=min(60,max(10,max(len(str(c.value or "")) for c in col)+2))
wb.save(D+"quarterly_export_clean.xlsx")
for row in s.iter_rows(values_only=True): print(row)
